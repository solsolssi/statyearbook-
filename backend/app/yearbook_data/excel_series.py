from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterable

from app.core.numeric_text import parse_numeric_value


PERIOD_LABELS = {
    "기준년도",
    "기준연도",
    "자료시점",
    "연도",
    "년도",
    "year",
    "period",
}
DIMENSION_LABELS = {
    "구분",
    "분류",
    "지역",
    "시도",
    "시도별",
    "성별",
    "유형",
    "종류",
    "등급",
    "직급",
    "기관",
    "부문",
    "대상",
    "항목",
}
SUPPORTED_SERIES_SUFFIXES = {".xls", ".xlsx"}


@dataclass(frozen=True)
class SeriesColumn:
    index: int
    path: tuple[str, ...]
    key: str
    role: str


@dataclass(frozen=True)
class SeriesRow:
    period: str
    dimensions: tuple[str, ...]
    values: tuple[str, ...]
    source_row_number: int

    @property
    def dimension_key(self) -> str:
        return json.dumps(
            [canonical_key(value) for value in self.dimensions],
            ensure_ascii=False,
            separators=(",", ":"),
        )


@dataclass(frozen=True)
class SeriesDataset:
    source_key: str
    title: str
    source_sheet: str
    header_row_count: int
    period_column: int
    dimension_columns: tuple[int, ...]
    columns: tuple[SeriesColumn, ...]
    rows: tuple[SeriesRow, ...]

    @property
    def structure_signature(self) -> str:
        payload = {
            "period_column": self.period_column,
            "dimension_columns": self.dimension_columns,
            "columns": [
                {"path": column.path, "role": column.role}
                for column in self.columns
            ],
        }
        serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True)
        return hashlib.sha256(serialized.encode("utf-8")).hexdigest()[:24]

    @property
    def value_columns(self) -> tuple[SeriesColumn, ...]:
        return tuple(column for column in self.columns if column.role == "value")


class SeriesWorkbookError(ValueError):
    pass


def parse_series_workbook(source_path: Path) -> list[SeriesDataset]:
    path = source_path.expanduser().resolve()
    if path.suffix.lower() not in SUPPORTED_SERIES_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SERIES_SUFFIXES))
        raise SeriesWorkbookError(f"시계열 원장은 {supported} 형식만 지원합니다.")

    matrices = (
        xlsx_matrices(path)
        if path.suffix.lower() == ".xlsx"
        else xls_matrices(path)
    )
    datasets: list[SeriesDataset] = []
    errors: list[str] = []
    for sheet_index, (sheet_name, matrix) in enumerate(matrices):
        if not matrix or not any(any(value.strip() for value in row) for row in matrix):
            continue
        if not looks_like_series_sheet(matrix):
            continue
        try:
            datasets.append(
                parse_series_matrix(
                    matrix,
                    title=sheet_name,
                    source_sheet=sheet_name,
                    sheet_index=sheet_index,
                )
            )
        except SeriesWorkbookError as error:
            errors.append(f"{sheet_name}: {error}")

    if not datasets:
        detail = "; ".join(errors[:5])
        suffix = f" ({detail})" if detail else ""
        raise SeriesWorkbookError(f"시계열 데이터가 있는 시트를 찾지 못했습니다.{suffix}")
    if errors:
        raise SeriesWorkbookError(
            "일부 시트의 구조를 해석할 수 없습니다: " + "; ".join(errors[:10])
        )
    return datasets


def looks_like_series_sheet(matrix: list[list[str]]) -> bool:
    return any(
        normalize_label(value) in PERIOD_LABELS
        for row in matrix[:10]
        for value in row
    )


def parse_series_matrix(
    matrix: list[list[str]],
    *,
    title: str,
    source_sheet: str = "",
    sheet_index: int = 0,
) -> SeriesDataset:
    rows = trim_matrix(matrix)
    period_column, data_start = locate_period_axis(rows)
    header_rows = rows[:data_start]
    if not header_rows:
        raise SeriesWorkbookError("헤더 행을 찾지 못했습니다.")

    width = max(len(row) for row in rows)
    padded_headers = [pad_row(row, width) for row in header_rows]
    paths = tuple(
        header_path(padded_headers, column_index)
        for column_index in range(width)
    )
    dimension_columns = tuple(
        column_index
        for column_index, path in enumerate(paths)
        if column_index != period_column and is_dimension_path(path)
    )
    value_columns = tuple(
        column_index
        for column_index in range(width)
        if column_index != period_column and column_index not in dimension_columns
    )
    if not value_columns:
        raise SeriesWorkbookError("측정값 열을 찾지 못했습니다.")

    columns = tuple(
        SeriesColumn(
            index=column_index,
            path=paths[column_index],
            key=column_key(paths[column_index], column_index),
            role=(
                "period"
                if column_index == period_column
                else ("dimension" if column_index in dimension_columns else "value")
            ),
        )
        for column_index in range(width)
    )
    parsed_rows: list[SeriesRow] = []
    seen_keys: set[tuple[str, tuple[str, ...]]] = set()
    for row_index, source_row in enumerate(rows[data_start:], start=data_start):
        row = pad_row(source_row, width)
        period = canonical_period(row[period_column])
        if not period:
            if any(value.strip() for value in row):
                raise SeriesWorkbookError(
                    f"{row_index + 1}행의 기준연도를 해석할 수 없습니다: "
                    f"{row[period_column]!r}"
                )
            continue
        dimensions = tuple(row[column_index].strip() for column_index in dimension_columns)
        identity = (period, tuple(canonical_key(value) for value in dimensions))
        if identity in seen_keys:
            raise SeriesWorkbookError(
                f"{row_index + 1}행의 기준연도·차원 조합이 중복되었습니다: "
                f"{period} / {' / '.join(dimensions)}"
            )
        seen_keys.add(identity)
        parsed_rows.append(
            SeriesRow(
                period=period,
                dimensions=dimensions,
                values=tuple(row[column_index] for column_index in value_columns),
                source_row_number=row_index + 1,
            )
        )

    if not parsed_rows:
        raise SeriesWorkbookError("저장할 시계열 행을 찾지 못했습니다.")
    source_key = canonical_key(title) or f"sheet-{sheet_index + 1}"
    return SeriesDataset(
        source_key=source_key,
        title=title.strip() or source_sheet.strip() or f"Sheet {sheet_index + 1}",
        source_sheet=source_sheet or title,
        header_row_count=data_start,
        period_column=period_column,
        dimension_columns=dimension_columns,
        columns=columns,
        rows=tuple(parsed_rows),
    )


def locate_period_axis(matrix: list[list[str]]) -> tuple[int, int]:
    if not matrix:
        raise SeriesWorkbookError("빈 시트입니다.")
    width = max(len(row) for row in matrix)
    candidates: list[tuple[int, int, int]] = []
    for header_row in range(min(10, len(matrix))):
        row = pad_row(matrix[header_row], width)
        for column_index, value in enumerate(row):
            if normalize_label(value) not in PERIOD_LABELS:
                continue
            data_start = next(
                (
                    row_index
                    for row_index in range(header_row + 1, len(matrix))
                    if canonical_period(pad_row(matrix[row_index], width)[column_index])
                ),
                -1,
            )
            if data_start >= 0:
                candidates.append((data_start, header_row, column_index))
    if not candidates:
        raise SeriesWorkbookError(
            "기준년도/연도 열을 찾지 못했습니다. 첫 10행에 기준년도 또는 연도 헤더가 필요합니다."
        )
    data_start, _, period_column = min(candidates)
    return period_column, data_start


def header_path(header_rows: list[list[str]], column_index: int) -> tuple[str, ...]:
    values: list[str] = []
    for row in header_rows:
        value = row[column_index].strip()
        if value and (not values or canonical_key(value) != canonical_key(values[-1])):
            values.append(value)
    return tuple(values) or (f"열{column_index + 1}",)


def is_dimension_path(path: tuple[str, ...]) -> bool:
    return any(normalize_label(part) in DIMENSION_LABELS for part in path)


def column_key(path: tuple[str, ...], column_index: int) -> str:
    normalized = [canonical_key(part) for part in path if canonical_key(part)]
    return ">".join(normalized) or f"column-{column_index + 1}"


def canonical_period(value: str) -> str:
    text = str(value or "").strip()
    match = re.fullmatch(r"((?:19|20)\d{2})(?:\\.0+)?", text)
    if match:
        return match.group(1)
    match = re.fullmatch(r"((?:19|20)\d{2})(?:00|년)", text)
    if match:
        return match.group(1)
    return ""


def canonical_key(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(value or "")).lower()
    return re.sub(r"[^0-9a-z가-힣]", "", normalized)


def normalize_label(value: str) -> str:
    lines = str(value).splitlines()
    return canonical_key(lines[0] if lines else "")


def xlsx_matrices(path: Path) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=False)
    result: list[tuple[str, list[list[str]]]] = []
    try:
        for worksheet in workbook.worksheets:
            values = [
                [
                    spreadsheet_cell_text(
                        worksheet.cell(row=row, column=column).value,
                        str(worksheet.cell(row=row, column=column).number_format or ""),
                    )
                    for column in range(1, worksheet.max_column + 1)
                ]
                for row in range(1, worksheet.max_row + 1)
            ]
            for merged in worksheet.merged_cells.ranges:
                value = values[merged.min_row - 1][merged.min_col - 1]
                for row in range(merged.min_row - 1, merged.max_row):
                    for column in range(merged.min_col - 1, merged.max_col):
                        values[row][column] = value
            result.append((worksheet.title, trim_matrix(values)))
    finally:
        workbook.close()
    return result


def xls_matrices(path: Path) -> list[tuple[str, list[list[str]]]]:
    try:
        import xlrd
    except ImportError as error:  # pragma: no cover - deployment dependency guard
        raise SeriesWorkbookError(
            "구형 .xls 파일을 읽으려면 xlrd 패키지가 필요합니다."
        ) from error

    workbook = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
    result: list[tuple[str, list[list[str]]]] = []
    try:
        for sheet_name in workbook.sheet_names():
            worksheet = workbook.sheet_by_name(sheet_name)
            matrix = [
                [
                    xls_cell_text(workbook, worksheet.cell(row, column))
                    for column in range(worksheet.ncols)
                ]
                for row in range(worksheet.nrows)
            ]
            result.append((sheet_name, trim_matrix(matrix)))
    finally:
        workbook.release_resources()
    return result


def scalar_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def spreadsheet_cell_text(value: Any, number_format: str) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return formatted_numeric_text(float(value), number_format)
    return scalar_text(value)


def xls_cell_text(workbook: Any, cell: Any) -> str:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, workbook.datemode).strftime("%Y-%m-%d")
    if cell.ctype != xlrd.XL_CELL_NUMBER:
        return scalar_text(cell.value)
    number_format = ""
    if 0 <= cell.xf_index < len(workbook.xf_list):
        format_key = workbook.xf_list[cell.xf_index].format_key
        format_info = workbook.format_map.get(format_key)
        number_format = str(format_info.format_str if format_info is not None else "")
    return formatted_numeric_text(float(cell.value), number_format)


def formatted_numeric_text(value: float, number_format: str) -> str:
    percent = "%" in number_format
    numeric = value * 100 if percent else value
    format_without_literals = re.sub(r'"[^"]*"', "", number_format)
    decimal_match = re.search(r"\.([0#]+)", format_without_literals)
    decimals = len(decimal_match.group(1)) if decimal_match else 0
    use_grouping = "," in format_without_literals.split(".", 1)[0]
    if decimals:
        rendered = f"{numeric:,.{decimals}f}" if use_grouping else f"{numeric:.{decimals}f}"
    elif float(numeric).is_integer():
        rendered = f"{int(numeric):,}" if use_grouping else str(int(numeric))
    else:
        rendered = str(numeric)
    return f"{rendered}%" if percent else rendered


def trim_matrix(matrix: Iterable[Iterable[str]]) -> list[list[str]]:
    rows = [[str(value or "").strip() for value in row] for row in matrix]
    while rows and not any(rows[0]):
        rows.pop(0)
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        return []
    width = max(
        max((index for index, value in enumerate(row) if value), default=-1) + 1
        for row in rows
    )
    return [pad_row(row, width) for row in rows]


def pad_row(row: Iterable[str], width: int) -> list[str]:
    values = list(row)
    return [*values[:width], *([""] * max(0, width - len(values)))]


def numeric_or_none(value: str) -> float | None:
    return parse_numeric_value(value)

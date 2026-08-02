from __future__ import annotations

from collections.abc import Iterable, Iterator
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
from typing import Any

from app.db.connection import DB_PATH, DatabaseConnection, connect
from app.db.schema import init_db


FULL_INFO_COLUMNS = {"표ID", "표이름", "출처구분", "저장테이블명"}
FULL_ITEM_COLUMNS = {
    "표ID",
    "itmId",
    "레벨",
    "자료번호",
    "항목명",
    "데이터유형",
    "전체항목명",
    "표시순번",
}
FULL_VALUE_COLUMNS = {
    "seqce_no",
    "statbl_id",
    "dtacycle_cd",
    "wrttime_idtfr_id",
    "dtadvs_cd",
    "itm_datano",
    "cls_datano",
    "dta_val",
}
DELTA_INFO_COLUMNS = {
    "작성관리번호",
    "표id",
    "표이름",
    "입력상태",
    "자료시점",
    "wrttimeIdtfrId",
}
DELTA_ITEM_COLUMNS = {
    "STATBL_ID",
    "itmId",
    "레벨",
    "자료번호",
    "항목명",
    "데이터유형",
    "전체항목명",
    "표시순번",
}
DELTA_VALUE_COLUMNS = {
    "STATBL_ID",
    "wrttimeIdtfrId",
    "자료번호",
    "데이터행번호",
    "데이터값",
    "검증여부",
}
VALUE_UPDATE_FIELDS = (
    "data_value",
    "string_value",
    "comment_identifier",
    "comment_text",
    "english_comment_text",
    "write_state",
    "check_code",
    "verified",
    "verification_result",
)


@dataclass(frozen=True)
class SourceDataImportResult:
    batch_id: int
    datasets: int
    items: int
    values: int
    inserted_values: int
    updated_values: int
    unchanged_values: int
    reused: bool = False


class RelationalSourceDataImporter:
    """Import the source system's normalized table/item/value Excel exports."""

    def __init__(self, db_path: Path = DB_PATH) -> None:
        self._db_path = db_path

    def import_full(
        self,
        *,
        info_path: Path,
        item_path: Path,
        value_path: Path,
        schema_path: Path | None = None,
    ) -> SourceDataImportResult:
        paths = [info_path, item_path, value_path]
        if schema_path is not None:
            paths.append(schema_path)
        source_hash = combined_digest(paths)
        with connect(self._db_path) as connection:
            init_db(connection)
            with connection:
                reused = self._existing_result(
                    connection,
                    source_hash=source_hash,
                    import_kind="relational_full",
                    target_year=0,
                )
                if reused is not None:
                    return reused
                batch_id = self._create_batch(
                    connection,
                    paths=paths,
                    source_hash=source_hash,
                    import_kind="relational_full",
                    target_year=0,
                )
                dataset_ids = self._import_full_info(connection, batch_id, info_path)
                item_count = self._import_items(
                    connection,
                    batch_id,
                    item_path,
                    dataset_ids,
                    full=True,
                )
                before_count = scalar_count(connection, "statistical_values")
                value_count = self._import_full_values(
                    connection,
                    batch_id,
                    value_path,
                    dataset_ids,
                )
                after_count = scalar_count(connection, "statistical_values")
                inserted = max(after_count - before_count, 0)
                result = SourceDataImportResult(
                    batch_id=batch_id,
                    datasets=len(dataset_ids),
                    items=item_count,
                    values=value_count,
                    inserted_values=inserted,
                    updated_values=max(value_count - inserted, 0),
                    unchanged_values=0,
                )
                self._complete_batch(connection, result)
                return result

    def import_delta(
        self,
        *,
        info_path: Path,
        item_path: Path,
        value_path: Path,
        target_year: int,
    ) -> SourceDataImportResult:
        paths = [info_path, item_path, value_path]
        source_hash = combined_digest(paths)
        with connect(self._db_path) as connection:
            init_db(connection)
            with connection:
                reused = self._existing_result(
                    connection,
                    source_hash=source_hash,
                    import_kind="relational_incremental",
                    target_year=target_year,
                )
                if reused is not None:
                    return reused
                batch_id = self._create_batch(
                    connection,
                    paths=paths,
                    source_hash=source_hash,
                    import_kind="relational_incremental",
                    target_year=target_year,
                )
                dataset_ids, metadata_count = self._import_delta_info(
                    connection,
                    batch_id,
                    info_path,
                )
                item_count = self._import_items(
                    connection,
                    batch_id,
                    item_path,
                    dataset_ids,
                    full=False,
                )
                values = list(read_delta_values(value_path, dataset_ids, batch_id))
                inserted, updated, unchanged = self._import_delta_values(
                    connection,
                    batch_id,
                    values,
                )
                result = SourceDataImportResult(
                    batch_id=batch_id,
                    datasets=len(dataset_ids),
                    items=item_count,
                    values=len(values),
                    inserted_values=inserted,
                    updated_values=updated,
                    unchanged_values=unchanged,
                )
                self._complete_batch(
                    connection,
                    result,
                    extra={"period_metadata": metadata_count},
                )
                return result

    def _import_full_info(
        self,
        connection: DatabaseConnection,
        batch_id: int,
        path: Path,
    ) -> dict[str, int]:
        headers, rows = tabular_rows(path)
        require_columns(headers, FULL_INFO_COLUMNS, path)
        for row in rows:
            external_id = text(row.get("표ID"))
            if not external_id:
                continue
            self._upsert_dataset(
                connection,
                external_id=external_id,
                title=text(row.get("표이름")) or external_id,
                source_kind=text(row.get("출처구분")),
                storage_table=text(row.get("저장테이블명")),
                introduced_batch_id=None,
            )
        return dataset_id_map(connection)

    def _import_delta_info(
        self,
        connection: DatabaseConnection,
        batch_id: int,
        path: Path,
    ) -> tuple[dict[str, int], int]:
        headers, rows = tabular_rows(path)
        require_columns(headers, DELTA_INFO_COLUMNS, path)
        materialized_rows = list(rows)
        touched_external_ids: set[str] = set()
        for row in materialized_rows:
            external_id = text(row.get("표id"))
            if not external_id:
                continue
            touched_external_ids.add(external_id)
            self._upsert_dataset(
                connection,
                external_id=external_id,
                title=text(row.get("표이름")) or external_id,
                introduced_batch_id=batch_id,
            )
        known = dataset_id_map(connection)
        touched = {
            external_id: known[external_id]
            for external_id in touched_external_ids
            if external_id in known
        }
        metadata_rows: list[tuple[object, ...]] = []
        for row in materialized_rows:
            external_id = text(row.get("표id"))
            dataset_id = touched.get(external_id)
            period = period_text(row.get("wrttimeIdtfrId"))
            if dataset_id is None or not period:
                continue
            metadata_rows.append(
                (
                    dataset_id,
                    period,
                    integer_text(row.get("작성관리번호")),
                    text(row.get("하위디렉토리")),
                    text(row.get("디렉토리")),
                    text(row.get("입력상태")),
                    text(row.get("담당부서")),
                    text(row.get("담당자명")),
                    text(row.get("입력기간")),
                    text(row.get("자료시점")),
                    batch_id,
                )
            )
        connection.executemany(
            """
            INSERT INTO statistical_period_metadata (
                dataset_id, period, management_number, subdirectory,
                directory, input_status, department, officer_name,
                input_period, data_reference_time, source_batch_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(dataset_id, period) DO UPDATE SET
                management_number = excluded.management_number,
                subdirectory = excluded.subdirectory,
                directory = excluded.directory,
                input_status = excluded.input_status,
                department = excluded.department,
                officer_name = excluded.officer_name,
                input_period = excluded.input_period,
                data_reference_time = excluded.data_reference_time,
                source_batch_id = excluded.source_batch_id,
                updated_at = CURRENT_TIMESTAMP
            """,
            metadata_rows,
        )
        return touched, len(metadata_rows)

    def _upsert_dataset(
        self,
        connection: DatabaseConnection,
        *,
        external_id: str,
        title: str,
        source_kind: str = "",
        storage_table: str = "",
        introduced_batch_id: int | None = None,
    ) -> None:
        source_key = external_id.lower()
        connection.execute(
            """
            INSERT INTO statistical_datasets (
                source_key, external_id, title, source_kind, storage_table,
                introduced_batch_id
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(source_key) DO UPDATE SET
                external_id = excluded.external_id,
                title = excluded.title,
                source_kind = CASE
                    WHEN excluded.source_kind <> '' THEN excluded.source_kind
                    ELSE statistical_datasets.source_kind
                END,
                storage_table = CASE
                    WHEN excluded.storage_table <> '' THEN excluded.storage_table
                    ELSE statistical_datasets.storage_table
                END,
                status = 'active',
                updated_at = CURRENT_TIMESTAMP
            """,
            (
                source_key,
                external_id,
                title,
                source_kind,
                storage_table,
                introduced_batch_id,
            ),
        )

    def _import_items(
        self,
        connection: DatabaseConnection,
        batch_id: int,
        path: Path,
        dataset_ids: dict[str, int],
        *,
        full: bool,
    ) -> int:
        headers, rows = tabular_rows(path)
        require_columns(headers, FULL_ITEM_COLUMNS if full else DELTA_ITEM_COLUMNS, path)
        id_column = "표ID" if full else "STATBL_ID"
        values: list[tuple[object, ...]] = []
        for row in rows:
            external_id = text(row.get(id_column))
            dataset_id = dataset_ids.get(external_id)
            if dataset_id is None:
                continue
            values.append(
                (
                    dataset_id,
                    integer_text(row.get("itmId")),
                    integer(row.get("레벨")),
                    integer(row.get("자료번호")),
                    text(row.get("항목명")),
                    text(row.get("데이터유형")),
                    yes_no(row.get("dummy")),
                    yes_no(row.get("inputNeedYn")),
                    integer(row.get("maxLevel")),
                    yes_no(row.get("leaf")),
                    text(row.get("전체항목명")),
                    integer(row.get("표시순번")),
                    batch_id,
                )
            )
        connection.executemany(
            """
            INSERT INTO statistical_dataset_items (
                dataset_id, item_id, level, item_data_no, item_name,
                data_type, is_dummy, input_required, max_level, is_leaf,
                full_name, display_order, source_batch_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(dataset_id, item_data_no) DO UPDATE SET
                item_id = excluded.item_id,
                level = excluded.level,
                item_name = excluded.item_name,
                data_type = excluded.data_type,
                is_dummy = excluded.is_dummy,
                input_required = excluded.input_required,
                max_level = excluded.max_level,
                is_leaf = excluded.is_leaf,
                full_name = excluded.full_name,
                display_order = excluded.display_order,
                source_batch_id = excluded.source_batch_id,
                updated_at = CURRENT_TIMESTAMP
            """,
            values,
        )
        return len(values)

    def _import_full_values(
        self,
        connection: DatabaseConnection,
        batch_id: int,
        path: Path,
        dataset_ids: dict[str, int],
    ) -> int:
        headers, rows = tabular_rows(path)
        require_columns(headers, FULL_VALUE_COLUMNS, path)
        count = 0
        chunk: list[tuple[object, ...]] = []
        for row in rows:
            external_id = text(row.get("statbl_id"))
            dataset_id = dataset_ids.get(external_id)
            if dataset_id is None:
                continue
            chunk.append(full_value_tuple(row, dataset_id, batch_id))
            count += 1
            if len(chunk) >= 5000:
                upsert_values(connection, chunk)
                chunk = []
        if chunk:
            upsert_values(connection, chunk)
        return count

    def _import_delta_values(
        self,
        connection: DatabaseConnection,
        batch_id: int,
        values: list[tuple[object, ...]],
    ) -> tuple[int, int, int]:
        if not values:
            return 0, 0, 0
        dataset_ids = sorted({int(value[0]) for value in values})
        periods = sorted({str(value[2]) for value in values})
        existing_rows = connection.execute(
            f"""
            SELECT *
            FROM statistical_values
            WHERE dataset_id IN ({", ".join("?" for _ in dataset_ids)})
              AND period IN ({", ".join("?" for _ in periods)})
            """,
            (*dataset_ids, *periods),
        ).fetchall()
        existing = {
            value_identity_from_row(row): row
            for row in existing_rows
        }
        inserted = 0
        updated = 0
        unchanged = 0
        revisions: list[tuple[object, ...]] = []
        changed_values: list[tuple[object, ...]] = []
        for value in values:
            identity = value_identity(value)
            current = existing.get(identity)
            if current is None:
                inserted += 1
                changed_values.append(value)
                continue
            merged_value = merge_delta_value(value, current)
            if value_payload(merged_value) == value_payload_from_row(current):
                unchanged += 1
                continue
            updated += 1
            revisions.append(
                (
                    int(current["id"]),
                    batch_id,
                    *(current[field] for field in VALUE_UPDATE_FIELDS),
                    int(current["source_batch_id"]),
                )
            )
            changed_values.append(merged_value)
        if revisions:
            connection.executemany(
                """
                INSERT INTO statistical_value_revisions (
                    statistical_value_id, replaced_by_batch_id,
                    data_value, string_value, comment_identifier,
                    comment_text, english_comment_text, write_state,
                    check_code, verified, verification_result, source_batch_id
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                revisions,
            )
        for start in range(0, len(changed_values), 5000):
            upsert_values(connection, changed_values[start : start + 5000])
        return inserted, updated, unchanged

    def _existing_result(
        self,
        connection: DatabaseConnection,
        *,
        source_hash: str,
        import_kind: str,
        target_year: int,
    ) -> SourceDataImportResult | None:
        row = connection.execute(
            """
            SELECT id, summary_json
            FROM series_import_batches
            WHERE file_hash = ? AND import_kind = ? AND target_year = ?
            """,
            (source_hash, import_kind, target_year),
        ).fetchone()
        if row is None:
            return None
        summary = json.loads(str(row["summary_json"] or "{}"))
        return SourceDataImportResult(
            batch_id=int(row["id"]),
            datasets=int(summary.get("datasets") or 0),
            items=int(summary.get("items") or 0),
            values=int(summary.get("values") or 0),
            inserted_values=int(summary.get("inserted_values") or 0),
            updated_values=int(summary.get("updated_values") or 0),
            unchanged_values=int(summary.get("unchanged_values") or 0),
            reused=True,
        )

    def _create_batch(
        self,
        connection: DatabaseConnection,
        *,
        paths: list[Path],
        source_hash: str,
        import_kind: str,
        target_year: int,
    ) -> int:
        cursor = connection.execute(
            """
            INSERT INTO series_import_batches (
                source_file_name, source_file_path, file_hash,
                import_kind, target_year, status
            ) VALUES (?, ?, ?, ?, ?, 'importing')
            """,
            (
                ", ".join(path.name for path in paths),
                json.dumps([str(path.resolve()) for path in paths], ensure_ascii=False),
                source_hash,
                import_kind,
                target_year,
            ),
        )
        return int(cursor.lastrowid)

    def _complete_batch(
        self,
        connection: DatabaseConnection,
        result: SourceDataImportResult,
        *,
        extra: dict[str, object] | None = None,
    ) -> None:
        summary = {
            "datasets": result.datasets,
            "items": result.items,
            "values": result.values,
            "inserted_values": result.inserted_values,
            "updated_values": result.updated_values,
            "unchanged_values": result.unchanged_values,
            **(extra or {}),
        }
        connection.execute(
            """
            UPDATE series_import_batches
            SET status = 'imported', summary_json = ?
            WHERE id = ?
            """,
            (json.dumps(summary, ensure_ascii=False), result.batch_id),
        )


def upsert_values(
    connection: DatabaseConnection,
    values: list[tuple[object, ...]],
) -> None:
    connection.executemany(
        """
        INSERT INTO statistical_values (
            dataset_id, data_cycle, period, data_division,
            item_data_no, class_data_no, data_value, string_value,
            comment_identifier, comment_text, english_comment_text,
            write_state, check_code, verified, verification_result,
            source_sequence, source_registered_at, source_batch_id
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(
            dataset_id, data_cycle, period, data_division,
            item_data_no, class_data_no
        ) DO UPDATE SET
            data_value = excluded.data_value,
            string_value = excluded.string_value,
            comment_identifier = excluded.comment_identifier,
            comment_text = excluded.comment_text,
            english_comment_text = excluded.english_comment_text,
            write_state = excluded.write_state,
            check_code = excluded.check_code,
            verified = excluded.verified,
            verification_result = excluded.verification_result,
            source_sequence = CASE
                WHEN excluded.source_sequence <> '' THEN excluded.source_sequence
                ELSE statistical_values.source_sequence
            END,
            source_registered_at = CASE
                WHEN excluded.source_registered_at <> '' THEN excluded.source_registered_at
                ELSE statistical_values.source_registered_at
            END,
            source_batch_id = excluded.source_batch_id,
            updated_at = CURRENT_TIMESTAMP
        """,
        values,
    )


def full_value_tuple(
    row: dict[str, Any],
    dataset_id: int,
    batch_id: int,
) -> tuple[object, ...]:
    return (
        dataset_id,
        text(row.get("dtacycle_cd")),
        period_text(row.get("wrttime_idtfr_id")),
        text(row.get("dtadvs_cd")),
        integer(row.get("itm_datano")),
        nullable_integer(row.get("cls_datano"), default=-1),
        text(row.get("dta_val")),
        text(row.get("dta_sval")),
        text(row.get("cmmt_idtfr")),
        text(row.get("cmmt_cont")),
        text(row.get("eng_cmmt_cont")),
        text(row.get("wrtstate_cd")),
        text(row.get("chck_cd")),
        yes_no(row.get("verify_yn")),
        text(row.get("verify_rslt")),
        integer_text(row.get("seqce_no")),
        text(row.get("reg_dtts")),
        batch_id,
    )


def read_delta_values(
    path: Path,
    dataset_ids: dict[str, int],
    batch_id: int,
) -> Iterator[tuple[object, ...]]:
    headers, rows = tabular_rows(path)
    require_columns(headers, DELTA_VALUE_COLUMNS, path)
    for row in rows:
        dataset_id = dataset_ids.get(text(row.get("STATBL_ID")))
        if dataset_id is None:
            continue
        yield (
            dataset_id,
            "YY",
            period_text(row.get("wrttimeIdtfrId")),
            "OD",
            integer(row.get("자료번호")),
            nullable_integer(row.get("데이터행번호"), default=-1),
            text(row.get("데이터값")),
            "",
            "",
            "",
            "",
            "",
            "",
            yes_no(row.get("검증여부")),
            "",
            "",
            "",
            batch_id,
        )


def value_identity(value: tuple[object, ...]) -> tuple[object, ...]:
    return value[:6]


def value_identity_from_row(row: Any) -> tuple[object, ...]:
    return (
        int(row["dataset_id"]),
        str(row["data_cycle"]),
        str(row["period"]),
        str(row["data_division"]),
        int(row["item_data_no"]),
        int(row["class_data_no"]),
    )


def value_payload(value: tuple[object, ...]) -> tuple[object, ...]:
    return value[6:15]


def value_payload_from_row(row: Any) -> tuple[object, ...]:
    return tuple(row[field] for field in VALUE_UPDATE_FIELDS)


def merge_delta_value(
    value: tuple[object, ...],
    current: Any,
) -> tuple[object, ...]:
    merged = list(value)
    # The incremental export contains only the displayed value and verification
    # flag. Preserve source-system fields that are absent from that export.
    for index, field in (
        (7, "string_value"),
        (8, "comment_identifier"),
        (9, "comment_text"),
        (10, "english_comment_text"),
        (11, "write_state"),
        (12, "check_code"),
        (14, "verification_result"),
        (15, "source_sequence"),
        (16, "source_registered_at"),
    ):
        merged[index] = current[field]
    return tuple(merged)


def tabular_rows(path: Path) -> tuple[list[str], Iterator[dict[str, Any]]]:
    source = path.expanduser().resolve()
    if source.suffix.lower() == ".xlsx":
        return xlsx_rows(source)
    if source.suffix.lower() == ".xls":
        return xls_rows(source)
    raise ValueError(f"지원하지 않는 원장 파일 형식입니다: {source.suffix}")


def xlsx_rows(path: Path) -> tuple[list[str], Iterator[dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    header_row = next(iterator, ())
    headers = [text(value) for value in header_row]

    def rows() -> Iterator[dict[str, Any]]:
        try:
            for values in iterator:
                yield {
                    header: values[index] if index < len(values) else None
                    for index, header in enumerate(headers)
                }
        finally:
            workbook.close()

    return headers, rows()


def xls_rows(path: Path) -> tuple[list[str], Iterator[dict[str, Any]]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    worksheet = workbook.sheet_by_index(0)
    headers = [text(value) for value in worksheet.row_values(0)]

    def rows() -> Iterator[dict[str, Any]]:
        try:
            for row_index in range(1, worksheet.nrows):
                values = worksheet.row_values(row_index)
                yield {
                    header: values[index] if index < len(values) else None
                    for index, header in enumerate(headers)
                }
        finally:
            workbook.release_resources()

    return headers, rows()


def require_columns(headers: list[str], required: set[str], path: Path) -> None:
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"{path.name}에 필수 열이 없습니다: {', '.join(missing)}")


def dataset_id_map(connection: DatabaseConnection) -> dict[str, int]:
    return {
        str(row["external_id"]): int(row["id"])
        for row in connection.execute(
            "SELECT id, external_id FROM statistical_datasets WHERE external_id <> ''"
        ).fetchall()
    }


def scalar_count(connection: DatabaseConnection, table_name: str) -> int:
    row = connection.execute(f"SELECT COUNT(*) AS count FROM {table_name}").fetchone()
    return int(row["count"]) if row is not None else 0


def combined_digest(paths: Iterable[Path]) -> str:
    digest = hashlib.sha256()
    for path in paths:
        source = path.expanduser().resolve()
        digest.update(source.name.encode("utf-8"))
        with source.open("rb") as source_file:
            for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
                digest.update(chunk)
    return digest.hexdigest()


def text(value: Any) -> str:
    if value is None:
        return ""
    rendered = str(value).strip()
    if rendered.upper() in {"NULL", "(NULL)", "NONE", "NAN"}:
        return ""
    return rendered


def integer(value: Any) -> int:
    if value is None or text(value) == "":
        return 0
    return int(float(value))


def nullable_integer(value: Any, *, default: int) -> int:
    if value is None or text(value) == "":
        return default
    return int(float(value))


def integer_text(value: Any) -> str:
    if value is None or text(value) == "":
        return ""
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return text(value)


def period_text(value: Any) -> str:
    return integer_text(value)


def yes_no(value: Any) -> int:
    return 1 if text(value).upper() in {"Y", "YES", "TRUE", "1"} else 0

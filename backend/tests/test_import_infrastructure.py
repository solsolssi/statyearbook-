from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from app.core.config import Settings
from app.domain.parsed import ParsedReport, ParsedTable
from app.ingest.parser_registry import (
    ReportParseError,
    parse_report,
    validate_parsed_report,
)
from app.ingest.repository import ReportImportRepository
from app.services.report_import_service import ReportImportRequest, ReportImportService
from app.storage.source_storage import LocalSourceStorage, ObjectSourceStorage
from app.workflows.job_executor import InProcessJobExecutor, QueuedJobExecutor
from app.workflows.report_job_repository import ReportProcessingJobRepository
from app.workflows.report_processing import ReportProcessingWorkflow


async def chunks(*values: bytes):  # type: ignore[no-untyped-def]
    for value in values:
        yield value


class FakeObjectClient:
    def __init__(self) -> None:
        self.objects: dict[tuple[str, str], bytes] = {}

    def upload_file(self, path: str, bucket: str, key: str) -> None:
        self.objects[(bucket, key)] = Path(path).read_bytes()

    def download_file(self, bucket: str, key: str, path: str) -> None:
        Path(path).write_bytes(self.objects[(bucket, key)])

    def delete_object(self, *, Bucket: str, Key: str) -> None:
        self.objects.pop((Bucket, Key), None)


class SourceStorageTest(unittest.IsolatedAsyncioTestCase):
    async def test_local_storage_hashes_materializes_and_deletes_source(self) -> None:
        with TemporaryDirectory() as directory:
            storage = LocalSourceStorage(Path(directory))
            stored = await storage.save("report.xlsx", chunks(b"abc", b"123"), max_bytes=20)

            self.assertEqual(stored.size, 6)
            self.assertEqual(
                stored.sha256,
                "6ca13d52ca70c883e0f0bb101e425a89e8624de51db2d2392593af6a84118090",
            )
            with storage.materialize(stored.object_key, stored.filename) as path:
                self.assertEqual(path.read_bytes(), b"abc123")
            storage.delete(stored.object_key)
            self.assertFalse(Path(stored.object_key).exists())

    async def test_object_storage_round_trip_uses_content_addressed_key(self) -> None:
        with TemporaryDirectory() as directory:
            client = FakeObjectClient()
            storage = ObjectSourceStorage(
                bucket="reports",
                temp_dir=Path(directory),
                client=client,
            )
            stored = await storage.save("report.hwpx", chunks(b"document"), max_bytes=20)

            self.assertTrue(stored.object_key.startswith(f"annual-reports/{stored.sha256}/"))
            with storage.materialize(stored.object_key, stored.filename) as path:
                self.assertEqual(path.read_bytes(), b"document")
            storage.delete(stored.object_key)
            self.assertEqual(client.objects, {})


class ParserRegistryTest(unittest.TestCase):
    def test_excel_parser_returns_validated_report_with_cell_coordinates(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "2027.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "1-1-1-1 테스트"
            sheet.append(["표번호", "1-1-1-1"])
            sheet.append(["통계명", "테스트"])
            sheet.append(["구분", "합계"])
            sheet.append(["계", 3])
            workbook.save(source)
            workbook.close()

            report = parse_report(
                source,
                source_type="xlsx",
                year=2027,
                title="2027 연보",
            )

            self.assertEqual(report.year, 2027)
            self.assertEqual(report.tables[0].code, "1-1-1-1")
            self.assertEqual(report.tables[0].cells[-1].source_coordinate, "B2")

    def test_structurally_invalid_report_is_rejected_before_persistence(self) -> None:
        report = ParsedReport(
            year=2027,
            title="오류 연보",
            source_path=Path("invalid.xlsx"),
            tables=(
                ParsedTable(code="1-1", title="첫 표", matrix=[["구분", "값"]]),
                ParsedTable(code="1-1", title="둘째 표", matrix=[["구분", "값"]]),
            ),
        )
        with self.assertRaisesRegex(ReportParseError, "중복"):
            validate_parsed_report(report)

    def test_parser_warnings_and_source_coordinates_are_persisted(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            source.write_bytes(b"source")
            with patch(
                "app.ingest.parser_registry.parse_tables",
                return_value=[
                    ParsedTable(
                        code="1-1",
                        title="경고 표",
                        matrix=[["구분", "값"], ["계"]],
                    )
                ],
            ):
                report = parse_report(
                    source,
                    source_type="xlsx",
                    year=2027,
                    title="2027 연보",
                )
            result = ReportImportRepository(root / "report.sqlite").replace_parsed_report(report)

            from app.db.connection import connect

            with connect(root / "report.sqlite") as connection:
                table = connection.execute(
                    "SELECT parse_warnings_json FROM stat_tables WHERE report_id = ?",
                    (result.report_id,),
                ).fetchone()
                cell = connection.execute(
                    """
                    SELECT source_coordinate
                    FROM stat_table_cells
                    WHERE table_id = (
                        SELECT id FROM stat_tables WHERE report_id = ? LIMIT 1
                    )
                    ORDER BY row_index DESC, col_index DESC
                    LIMIT 1
                    """,
                    (result.report_id,),
                ).fetchone()
            self.assertIn("정규화", str(table["parse_warnings_json"]))
            self.assertEqual(cell["source_coordinate"], "B2")


class JobInfrastructureTest(unittest.IsolatedAsyncioTestCase):
    async def test_duplicate_upload_reuses_the_existing_job(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            settings = Settings(
                upload_dir=root / "uploads",
                max_upload_bytes=1024,
            )
            workflow = ReportProcessingWorkflow(
                root / "report.sqlite",
                settings=settings,
            )
            service = ReportImportService(settings=settings, workflow=workflow)
            request = ReportImportRequest(filename="2027 연보.xlsx")

            first = await service.create_job(request, chunks(b"same workbook"))
            second = await service.create_job(request, chunks(b"same workbook"))

            self.assertEqual(first["id"], second["id"])
            self.assertEqual(len(list(settings.upload_dir.iterdir())), 1)

    async def test_job_lease_blocks_another_worker_and_tracks_progress(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            source.write_bytes(b"source")
            repository = ReportProcessingJobRepository(root / "report.sqlite")
            job_id = repository.create(
                source_file_name=source.name,
                source_file_path=str(source),
                source_type="xlsx",
                report_year=2027,
                report_title="2027 연보",
                options={},
            )

            self.assertTrue(repository.acquire_lease(job_id, owner="worker-a", lease_seconds=60))
            self.assertFalse(repository.acquire_lease(job_id, owner="worker-b", lease_seconds=60))
            self.assertTrue(
                repository.heartbeat(
                    job_id,
                    owner="worker-a",
                    lease_seconds=60,
                    progress_percent=42,
                )
            )
            self.assertEqual(repository.get(job_id)["progress_percent"], 42)
            repository.release_lease(job_id, owner="worker-a")
            self.assertTrue(repository.acquire_lease(job_id, owner="worker-b", lease_seconds=60))

    async def test_job_executor_adapters_share_the_submit_contract(self) -> None:
        scheduled: list[tuple[object, int, bool]] = []

        def schedule(runner, job_id, retry):  # type: ignore[no-untyped-def]
            scheduled.append((runner, job_id, retry))

        def runner(job_id: int, retry: bool) -> None:
            del job_id, retry

        InProcessJobExecutor(schedule, runner).submit(3, retry=True)
        QueuedJobExecutor().submit(4)

        self.assertEqual(scheduled, [(runner, 3, True)])


if __name__ == "__main__":
    unittest.main()

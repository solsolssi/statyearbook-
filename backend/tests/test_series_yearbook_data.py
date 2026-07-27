from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from app.domain.parsed import ParsedTable
from app.ingest.repository import ReportImportRepository
from app.validation.models import ValidationCell, ValidationTable
from app.validation.profile_repository import ValidationProfileRepository
from app.validation.profiles import (
    HeuristicProfileDraftProvider,
    compatible_structure_signature,
    structure_signature,
)
from app.validation.repository import ValidationRepository
from app.yearbook_data.excel_series import canonical_key, parse_series_workbook
from app.yearbook_data.materializer import (
    SeriesMaterializationError,
    materialize_series_table,
)
from app.yearbook_data.repository import SeriesImportRepository


def save_series_workbook(
    path: Path,
    *,
    title: str = "연도별 안전교육",
    latest_total: int = 30,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(["기준년도", "합계", "학교", "학교"])
    sheet.append(["기준년도", "합계", "학생", "교사"])
    sheet.append([2023, 10, 6, 4])
    sheet.append([2024, 20, 12, 8])
    sheet.append([2025, latest_total, 18, 12])
    workbook.save(path)
    workbook.close()


def validation_table(matrix: list[list[str]], *, header_count: int) -> ValidationTable:
    cells = [
        ValidationCell(
            row_index=row_index,
            col_index=col_index,
            text_value=value,
            numeric_value=None,
            is_header=row_index < header_count,
        )
        for row_index, row in enumerate(matrix)
        for col_index, value in enumerate(row)
    ]
    return ValidationTable(
        id=1,
        report_id=1,
        code="1-1",
        title="테스트",
        unit="명",
        base_date="2025. 12. 31.",
        source="테스트과",
        note="",
        cells=cells,
    )


class ExcelSeriesParserTest(unittest.TestCase):
    def test_multirow_headers_are_stored_as_stable_column_paths(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "series.xlsx"
            save_series_workbook(path)

            dataset = parse_series_workbook(path)[0]

            self.assertEqual(dataset.header_row_count, 2)
            self.assertEqual(dataset.period_column, 0)
            self.assertEqual(dataset.dimension_columns, ())
            self.assertEqual(dataset.columns[2].path, ("학교", "학생"))
            self.assertEqual(dataset.rows[-1].period, "2025")
            self.assertEqual(dataset.rows[-1].values, ("30", "18", "12"))

    def test_region_column_is_recognized_as_a_dimension(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "regions.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "지역별 시설"
            sheet.append(["기준년도", "지역", "합계"])
            sheet.append([2025, "합계", 3])
            sheet.append([2025, "서울", 1])
            sheet.append([2025, "부산", 2])
            workbook.save(path)
            workbook.close()

            dataset = parse_series_workbook(path)[0]

            self.assertEqual(dataset.dimension_columns, (1,))
            self.assertEqual(dataset.rows[1].dimensions, ("서울",))

    def test_non_data_cover_sheet_is_ignored(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "series.xlsx"
            workbook = Workbook()
            cover = workbook.active
            cover.title = "안내"
            cover.append(["통계 원장 입력 안내"])
            sheet = workbook.create_sheet("연도별 시설")
            sheet.append(["기준년도", "합계"])
            sheet.append([2025, 3])
            workbook.save(path)
            workbook.close()

            datasets = parse_series_workbook(path)

            self.assertEqual([dataset.title for dataset in datasets], ["연도별 시설"])


class SeriesImportRepositoryTest(unittest.TestCase):
    def test_imports_are_idempotent_and_new_batches_override_observations(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "series.xlsx"
            repository = SeriesImportRepository(root / "report.sqlite")
            save_series_workbook(path, latest_total=30)

            first = repository.import_workbook(path, import_kind="full")
            repeated = repository.import_workbook(path, import_kind="full")
            save_series_workbook(path, latest_total=31)
            corrected = repository.import_workbook(
                path,
                import_kind="incremental",
                target_year=2026,
            )
            latest = repository.latest_rows(canonical_key("연도별 안전교육"))

            self.assertFalse(first.reused)
            self.assertTrue(repeated.reused)
            self.assertNotEqual(first.batch_id, corrected.batch_id)
            row_2025 = next(row for row in latest if row["period"] == "2025")
            self.assertEqual(row_2025["values"][0], "31")


class SeriesMaterializerTest(unittest.TestCase):
    def test_rolling_year_rows_keep_the_approved_layout(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "series.xlsx"
            save_series_workbook(path)
            dataset = parse_series_workbook(path)[0]
            template = ParsedTable(
                code="6-1-1",
                title="연도별 안전교육",
                title_en="Safety Education by Year",
                matrix=[
                    ["연도\nYear", "합계\nTotal", "학교\nSchool", ""],
                    ["연도\nYear", "합계\nTotal", "학생\nStudent", "교사\nTeacher"],
                    ["2023", "10", "6", "4"],
                    ["2024", "20", "12", "8"],
                ],
                header_count=2,
                base_date="2024. 12. 31.",
            )

            materialized = materialize_series_table(
                template,
                dataset,
                edition_year=2026,
            )

            self.assertEqual(materialized.strategy, "period_rows")
            self.assertEqual(materialized.source_periods, ("2024", "2025"))
            self.assertEqual(
                materialized.table.matrix,
                [
                    ["연도\nYear", "합계\nTotal", "학교\nSchool", ""],
                    ["연도\nYear", "합계\nTotal", "학생\nStudent", "교사\nTeacher"],
                    ["2024", "20", "12", "8"],
                    ["2025", "30", "18", "12"],
                ],
            )
            self.assertEqual(materialized.table.base_date, "2025. 12. 31.")

    def test_dimension_changes_require_a_recipe_review(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "regions.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "지역별 시설"
            sheet.append(["기준년도", "지역", "합계"])
            sheet.append([2025, "합계", 3])
            sheet.append([2025, "서울", 1])
            sheet.append([2025, "신규지역", 2])
            workbook.save(path)
            workbook.close()
            dataset = parse_series_workbook(path)[0]
            template = ParsedTable(
                code="7-1-1",
                title="지역별 시설",
                matrix=[
                    ["지역\nRegion", "합계\nTotal"],
                    ["합계\nTotal", "2"],
                    ["서울\nSeoul", "1"],
                    ["부산\nBusan", "1"],
                ],
                header_count=1,
            )

            with self.assertRaisesRegex(
                SeriesMaterializationError,
                "삭제·명칭변경 후보.*신규 행 후보",
            ):
                materialize_series_table(template, dataset, edition_year=2026)


class RollingProfileCompatibilityTest(unittest.TestCase):
    def test_calendar_year_labels_do_not_change_compatible_structure(self) -> None:
        previous = validation_table(
            [["연도", "합계"], ["2023", "10"], ["2024", "20"]],
            header_count=1,
        )
        current = validation_table(
            [["연도", "합계"], ["2024", "20"], ["2025", "30"]],
            header_count=1,
        )

        self.assertNotEqual(structure_signature(previous), structure_signature(current))
        self.assertEqual(
            compatible_structure_signature(previous),
            compatible_structure_signature(current),
        )

    def test_repository_reuses_profile_when_only_the_year_window_rolls(self) -> None:
        class CountingProvider(HeuristicProfileDraftProvider):
            def __init__(self) -> None:
                self.calls = 0

            def draft(self, *args, **kwargs):  # type: ignore[no-untyped-def]
                self.calls += 1
                return super().draft(*args, **kwargs)

        with TemporaryDirectory() as directory:
            root = Path(directory)
            db_path = root / "report.sqlite"
            first_source = root / "first.xlsx"
            second_source = root / "second.xlsx"
            first_source.write_bytes(b"first")
            second_source.write_bytes(b"second")
            importer = ReportImportRepository(db_path)
            provider = CountingProvider()
            profiles = ValidationProfileRepository(db_path)

            first = importer.replace_report(
                source_path=first_source,
                year=2026,
                title="2026 연보",
                tables=[
                    ParsedTable(
                        code="1-1",
                        title="연도별 테스트",
                        matrix=[["연도", "합계"], ["2023", "10"], ["2024", "20"]],
                        header_count=1,
                    )
                ],
            )
            profiles.ensure_profiles(
                report_id=first.report_id,
                tables=ValidationRepository(db_path).load_tables(first.report_id),
                provider=provider,
            )
            second = importer.replace_report(
                source_path=second_source,
                year=2027,
                title="2027 연보",
                tables=[
                    ParsedTable(
                        code="1-1",
                        title="연도별 테스트",
                        matrix=[["연도", "합계"], ["2024", "20"], ["2025", "30"]],
                        header_count=1,
                    )
                ],
            )
            inherited = profiles.ensure_profiles(
                report_id=second.report_id,
                tables=ValidationRepository(db_path).load_tables(second.report_id),
                provider=provider,
            )

            self.assertEqual(provider.calls, 1)
            self.assertEqual(inherited["1-1"].source, "inherited")


if __name__ == "__main__":
    unittest.main()

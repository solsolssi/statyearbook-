"""Apply the user-edited slide 5's open, minimal visual style to the deck."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from build_yearbook_validation_presentation import force_font_family


FONT_NAME = "Noto Sans KR"

NAVY = "0C2F61"
BLUE = "1F5A91"
TEAL = "2F7D68"
TEXT = "172331"
MUTED = "667085"
LIGHT = "F7F9FB"
LIGHT_LINE = "D9E0E6"
WHITE = "FFFFFF"

DARK_ACCENTS = {NAVY, BLUE, TEAL, "153E70", "172B59"}
SOFT_FILLS = {"EEF6FB", "E8F5F1", "F7F9FB", "F8FAFC", "FAFBFC"}
STATUS_FILLS = {"FDECEF", "FFF4DE"}
HEADER_NAMES = {"상단 분류", "슬라이드 제목", "제목 구분선", "슬라이드 부제목", "하단 문구", "페이지 번호"}


def color_hex(color_format) -> str | None:
    try:
        rgb = color_format.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    return str(rgb) if rgb is not None else None


def fill_hex(shape) -> str | None:
    try:
        if shape.fill.type is None:
            return None
        return color_hex(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def line_hex(shape) -> str | None:
    try:
        sp_pr = shape._element.spPr
        line = sp_pr.ln
        if line is None:
            return None
        solid_fill = line.find(qn("a:solidFill"))
        if solid_fill is None:
            return None
        srgb = solid_fill.find(qn("a:srgbClr"))
        return srgb.get("val") if srgb is not None else None
    except (AttributeError, TypeError, ValueError):
        return None


def set_fill(shape, value: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(value)


def remove_line(shape) -> None:
    try:
        shape.line.fill.background()
    except (AttributeError, TypeError, ValueError):
        pass


def set_line(shape, value: str, width: float = 0.8) -> None:
    try:
        shape.line.color.rgb = RGBColor.from_string(value)
        shape.line.width = Pt(width)
    except (AttributeError, TypeError, ValueError):
        pass


def set_text_color(shape, value: str, *, bold: bool | None = None) -> None:
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(value)
            if bold is not None:
                run.font.bold = bold


def intersects_header_text(text_shape, header_shape) -> bool:
    cx = text_shape.left + text_shape.width / 2
    cy = text_shape.top + text_shape.height / 2
    return (
        header_shape.left <= cx <= header_shape.left + header_shape.width
        and header_shape.top <= cy <= header_shape.top + header_shape.height
    )


def add_accent_bar(slide, shape, accent: str, *, width: float = 0.045) -> None:
    bar = slide.shapes.add_shape(
        1,
        shape.left,
        shape.top,
        Inches(width),
        shape.height,
    )
    bar.name = f"클린 스타일 강조선_{shape.name}"
    set_fill(bar, accent)
    remove_line(bar)


def restyle_dark_header(slide, header_shape, accent: str) -> None:
    overlaps = [
        shape
        for shape in slide.shapes
        if shape is not header_shape
        and shape.has_text_frame
        and shape.text.strip()
        and intersects_header_text(shape, header_shape)
    ]
    number_shapes = [shape for shape in overlaps if re.fullmatch(r"\d{1,2}", shape.text.strip())]

    set_fill(header_shape, WHITE)
    remove_line(header_shape)
    add_accent_bar(slide, header_shape, accent, width=0.055)

    for shape in overlaps:
        is_number = shape in number_shapes
        set_text_color(shape, accent if is_number else TEXT, bold=True)
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
        if is_number:
            shape.left = header_shape.left + Inches(0.14)
            shape.width = Inches(0.38)
        else:
            offset = 0.57 if number_shapes else 0.20
            shape.left = header_shape.left + Inches(offset)
            shape.width = max(Inches(0.4), header_shape.width - Inches(offset + 0.14))


def soften_table(table_shape) -> None:
    table = table_shape.table
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if row_index == 0:
                continue
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(LIGHT if row_index % 2 else WHITE)
            tc_pr = cell._tc.get_or_add_tcPr()
            for child in list(tc_pr):
                if child.tag.rsplit("}", 1)[-1].startswith("ln"):
                    tc_pr.remove(child)


def restyle_slide(slide, slide_number: int) -> None:
    # Slide 5 is the user's visual reference and is intentionally preserved.
    if slide_number in {1, 5}:
        return

    dark_headers = []
    row_cards = []

    for shape in list(slide.shapes):
        if shape.name in HEADER_NAMES or shape.has_table:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            # Explicit no-line styling prevents theme fallbacks from rendering
            # text-box bounds as black rectangles in Quick Look or PowerPoint.
            remove_line(shape)
        top = shape.top / 914400
        height = shape.height / 914400
        width = shape.width / 914400
        fill = fill_hex(shape)
        line = line_hex(shape)
        is_image_slot = "이미지 삽입" in shape.name

        if fill in DARK_ACCENTS and 1.70 <= top < 4.45 and 0.28 <= height <= 0.68 and width >= 1.20:
            dark_headers.append((shape, fill))
            continue

        if is_image_slot:
            if width >= 1.3 and height >= 0.9:
                set_fill(shape, WHITE)
                set_line(shape, LIGHT_LINE, 0.9)
            continue

        if fill in SOFT_FILLS:
            set_fill(shape, LIGHT)

        if fill in {WHITE, *SOFT_FILLS} and line and width >= 1.4:
            if 0.34 <= height <= 0.90:
                row_cards.append((shape, line if line in DARK_ACCENTS else BLUE))
            remove_line(shape)
        elif fill == WHITE and line and height > 0.90:
            remove_line(shape)
        elif fill in SOFT_FILLS and line:
            remove_line(shape)

        if fill not in STATUS_FILLS and line == LIGHT_LINE:
            remove_line(shape)

    for header_shape, accent in dark_headers:
        restyle_dark_header(slide, header_shape, accent)

    for row_shape, accent in row_cards:
        # Open list rows use one slim vertical mark instead of a full border.
        add_accent_bar(slide, row_shape, accent, width=0.035)

    for shape in slide.shapes:
        if shape.has_table:
            soften_table(shape)


def clear_all_notes(prs: Presentation) -> None:
    for slide in prs.slides:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            frame.clear()
            frame.text = ""


def fix_last_slide_spacing(slide) -> None:
    table_shapes = [shape for shape in slide.shapes if shape.has_table]
    if table_shapes:
        table_shape = table_shapes[0]
        target_heights = [0.46, 0.47, 0.47, 0.47, 0.47, 0.47]
        for row, height in zip(table_shape.table.rows, target_heights):
            row.height = Inches(height)
        table_shape.height = Inches(sum(target_heights))

    for shape in slide.shapes:
        text = shape.text.strip() if shape.has_text_frame else ""
        if text.startswith("확실한 규칙부터 적용하고"):
            shape.top = Inches(4.86)
            shape.height = Inches(0.18)
            continue
        if shape.top > Inches(4.65) and shape.width > Inches(5.0) and fill_hex(shape) in DARK_ACCENTS:
            shape.top = Inches(4.82)
            shape.height = Inches(0.24)


def build(source: Path, output: Path) -> None:
    prs = Presentation(source)
    if len(prs.slides) != 18:
        raise RuntimeError(f"예상한 18쪽이 아니라 {len(prs.slides)}쪽입니다.")

    for index, slide in enumerate(prs.slides, start=1):
        restyle_slide(slide, index)

    fix_last_slide_spacing(prs.slides[17])

    clear_all_notes(prs)
    force_font_family(prs, FONT_NAME)
    prs.core_properties.title = "통계연보 검수 보조 시스템"
    prs.core_properties.subject = "통계연보 검수 보조 시스템 발표자료 - 클린 디자인"

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()

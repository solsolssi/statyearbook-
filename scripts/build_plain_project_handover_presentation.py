"""Build a plain, large-text project handover presentation.

The presentation intentionally uses only a white background and black text.
It assumes that the recipient already has the complete project directory.

Run from the repository root:
    ./.venv/bin/python scripts/build_plain_project_handover_presentation.py
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "exports" / "통계연보_검수_보조_시스템_인수인계_흰배경_큰글씨_상세설명본.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "맑은 고딕"
MONO = "Consolas"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)


@dataclass(frozen=True)
class SlideSpec:
    title: str
    items: tuple[str, ...] = ()
    intro: str = ""
    commands: str = ""
    kind: str = "body"


def set_typeface(run, *, size: float, bold: bool = False, font: str = FONT) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    properties = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = properties.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            properties.append(node)
        node.set("typeface", font)


def set_white_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def add_text_box(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    bold: bool = False,
    font: str = FONT,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0)
    frame.margin_right = Inches(0)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    set_typeface(run, size=size, bold=bold, font=font)
    return shape


def add_page_number(slide, page: int) -> None:
    add_text_box(
        slide,
        str(page),
        12.15,
        7.08,
        0.45,
        0.18,
        size=9.5,
        align=PP_ALIGN.RIGHT,
    )


def add_cover(slide, spec: SlideSpec, page: int) -> None:
    set_white_background(slide)
    add_text_box(
        slide,
        spec.title,
        0.85,
        1.68,
        11.65,
        1.45,
        size=34,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text_box(
        slide,
        spec.intro,
        1.05,
        3.55,
        11.25,
        1.40,
        size=21,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text_box(
        slide,
        "기준일 2026-08-28  |  상세 문서 docs/PROJECT_HANDOVER_GUIDE.md",
        1.0,
        6.60,
        11.30,
        0.30,
        size=12,
        align=PP_ALIGN.CENTER,
    )
    add_page_number(slide, page)


def add_section(slide, spec: SlideSpec, page: int) -> None:
    set_white_background(slide)
    add_text_box(
        slide,
        spec.title,
        0.85,
        2.05,
        11.65,
        0.85,
        size=34,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text_box(
        slide,
        spec.intro,
        1.05,
        3.32,
        11.25,
        1.35,
        size=21,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_page_number(slide, page)


def body_font_size(items: Sequence[str]) -> float:
    total = sum(len(item) for item in items)
    if len(items) <= 4 and total < 450:
        return 19.0
    if total > 850:
        return 15.5
    if total > 650:
        return 16.5
    return 17.5


def add_body(slide, spec: SlideSpec, page: int) -> None:
    set_white_background(slide)
    add_text_box(slide, spec.title, 0.70, 0.43, 11.95, 0.55, size=28, bold=True)
    shape = slide.shapes.add_textbox(Inches(0.82), Inches(1.30), Inches(11.72), Inches(5.55))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0)
    frame.margin_right = Inches(0)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    size = body_font_size(spec.items)
    for index, item in enumerate(spec.items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(9)
        paragraph.line_spacing = 1.13
        run = paragraph.add_run()
        run.text = item
        set_typeface(run, size=size)
    add_page_number(slide, page)


def add_command(slide, spec: SlideSpec, page: int) -> None:
    set_white_background(slide)
    add_text_box(slide, spec.title, 0.70, 0.43, 11.95, 0.55, size=28, bold=True)
    add_text_box(slide, spec.intro, 0.82, 1.30, 11.72, 1.03, size=18)
    add_text_box(slide, spec.commands, 0.92, 2.50, 11.50, 3.90, size=16.2, font=MONO)
    add_text_box(slide, "명령은 한 줄을 입력한 뒤 Enter를 누르고, 오류가 없을 때 다음 줄로 진행하세요.", 0.82, 6.52, 11.72, 0.32, size=15.5)
    add_page_number(slide, page)


def build_specs() -> list[SlideSpec]:
    return [
        SlideSpec(
            "통계연보 검수 보조 시스템\n프로젝트 인수인계 설명서",
            intro="프로젝트 파일 전체를 이미 전달받았다는 전제에서, 빈 컴퓨터에 필요한 프로그램을 설치하고 서비스를 실행한 뒤 기능을 사용하고 유지보수하는 순서입니다.",
            kind="cover",
        ),
        SlideSpec(
            "이 PPT에서 가정하는 시작 상태",
            items=(
                "1. 기존 담당자에게 프로젝트 폴더 전체를 받았다고 가정하세요. 이 폴더에는 backend, frontend, docs, Dockerfile, .env.example 같은 파일이 있어야 합니다.",
                "2. Windows에서는 프로젝트 폴더 이름을 statyearbook으로 정하고 C:\\work\\statyearbook에 두는 것을 기준으로 설명합니다. OneDrive, 바탕 화면, 네트워크 드라이브는 피하세요.",
                "3. 프로젝트 코드와 함께 실제 SQLite DB와 backend/uploads 원본을 받았다면 원본 복사본을 별도 보관하세요. 서버 첫 실행이 DB 구조를 갱신할 수 있으므로 백업 없이 실제 DB를 먼저 열면 안 됩니다.",
                "4. 처음 실행할 때는 외부 LLM 기능을 끄고 빈 테스트 SQLite를 사용하세요. 화면과 서버가 정상적으로 실행되는 것을 확인한 뒤 실제 DB와 외부 서비스를 연결하세요.",
            ),
        ),
        SlideSpec(
            "전체 진행 순서",
            items=(
                "1. Python 3.12, Node.js 24 LTS, 선택적으로 Git과 VS Code를 설치하세요.",
                "2. 프로젝트 폴더를 C:\\work\\statyearbook에 두고, 필요한 파일이 모두 있는지 확인하세요.",
                "3. 루트 .env를 만들고 안전한 SQLite와 LLM 비활성 설정을 확인하세요.",
                "4. 백엔드 Python 가상환경과 프런트엔드 npm 패키지를 각각 설치하세요.",
                "5. 백엔드와 프런트엔드를 서로 다른 PowerShell 창에서 실행하고 /health, /ready, 브라우저 화면을 확인하세요.",
                "6. 설치가 끝나면 기능 사용법을 익히고, 마지막에 프로젝트 구조·테스트·백업·배포 절차를 확인하세요.",
            ),
        ),
        SlideSpec(
            "1부. 아무것도 없는 컴퓨터에 설치하기",
            intro="Windows를 본문 기준으로 설명합니다. 각 단계가 정상인지 확인한 뒤 다음 단계로 이동하세요.",
            kind="section",
        ),
        SlideSpec(
            "1-1. 설치 전에 컴퓨터와 폴더를 준비하세요",
            items=(
                "1. Windows 시작 버튼을 누르고 PowerShell을 검색한 뒤 일반 권한으로 여세요. 관리자 권한은 조직 정책이 요구할 때만 사용하세요.",
                "2. C 드라이브에 work 폴더를 만드세요. PowerShell에서 New-Item -ItemType Directory -Force C:\\work를 입력하면 됩니다.",
                "3. 전달받은 statyearbook 폴더를 C:\\work 아래에 복사하여 최종 경로가 C:\\work\\statyearbook이 되게 하세요.",
                "4. 프로젝트를 OneDrive 동기화 폴더나 네트워크 드라이브에 두지 마세요. node_modules와 대용량 SQLite가 동기화되면 설치가 매우 느려지거나 DB 잠금 오류가 날 수 있습니다.",
                "5. Windows 탐색기에서 파일 확장명 표시를 켜세요. .env.example, .xlsx, .hwpx 같은 확장자를 정확히 구분해야 합니다.",
            ),
        ),
        SlideSpec(
            "1-2. Python 3.12를 설치하세요",
            items=(
                "1. 브라우저에서 python.org의 Python 3.12 다운로드 페이지를 여세요. 이 프로젝트의 운영 Docker 이미지가 Python 3.12이므로 먼저 같은 계열을 사용하세요.",
                "2. Windows 64-bit installer를 내려받아 실행하세요. 첫 화면에서 Add python.exe to PATH를 반드시 체크하세요.",
                "3. Python Launcher인 py가 함께 설치되도록 기본 옵션을 유지하고 Install Now를 누르세요.",
                "4. 설치가 끝나면 열려 있던 PowerShell을 모두 닫고 새 PowerShell을 여세요. PATH는 새 창에서 반영됩니다.",
                "5. py -3.12 --version을 입력하세요. Python 3.12.x가 표시되면 정상입니다. 찾을 수 없다는 메시지가 나오면 다음 단계로 넘어가지 마세요.",
            ),
        ),
        SlideSpec(
            "1-3. Node.js 24 LTS와 npm을 설치하세요",
            items=(
                "1. 브라우저에서 nodejs.org/en/download를 열고 Current가 아니라 LTS로 표시된 24 계열 Windows Installer를 받으세요.",
                "2. 설치 프로그램을 실행하고 특별한 조직 지침이 없다면 기본 옵션으로 진행하세요. npm은 Node.js와 함께 설치되므로 별도 설치하지 마세요.",
                "3. 설치가 끝나면 PowerShell을 닫고 새로 여세요.",
                "4. node --version을 입력하여 v24로 시작하는 버전이 나오는지 확인하세요.",
                "5. npm --version을 입력하여 숫자가 표시되는지 확인하세요. node와 npm 중 하나라도 찾지 못하면 프런트엔드 설치를 시작하지 마세요.",
            ),
        ),
        SlideSpec(
            "1-4. Git과 VS Code는 유지보수를 위해 설치하세요",
            items=(
                "1. 서비스를 실행하는 데 Git이 반드시 필요한 것은 아니지만, 다음 담당자가 변경 이력을 확인하고 안전하게 수정하려면 설치하는 것이 좋습니다.",
                "2. git-scm.com/install/windows에서 64-bit Git for Windows를 받아 기본 옵션으로 설치하세요. 명령줄에서 Git을 사용할 수 있는 PATH 권장 옵션을 유지하세요.",
                "3. 새 PowerShell에서 git --version을 입력하여 버전이 표시되는지 확인하세요.",
                "4. 코드를 수정할 예정이면 code.visualstudio.com에서 VS Code를 설치하세요. 프로젝트 폴더를 열 때 C:\\work\\statyearbook 폴더 자체를 선택하세요.",
                "5. Git 저장소가 함께 전달되었다면 프로젝트 루트에서 git status와 git log -1 --oneline을 실행하여 기준 커밋과 미커밋 변경을 기록하세요.",
            ),
        ),
        SlideSpec(
            "1-5. PowerShell에서 폴더를 이동하는 방법",
            items=(
                "1. PowerShell 명령은 현재 위치에 따라 결과가 달라집니다. 먼저 Set-Location C:\\work\\statyearbook을 입력하여 프로젝트 루트로 이동하세요.",
                "2. Get-Location을 입력하세요. Path가 C:\\work\\statyearbook으로 나오면 프로젝트 루트에 있는 것입니다.",
                "3. Get-ChildItem을 입력하세요. backend, frontend, docs, Dockerfile, .env.example이 목록에 보여야 합니다.",
                "4. frontend 명령을 실행할 때는 Set-Location C:\\work\\statyearbook\\frontend로 이동하세요. 백엔드 명령을 실행할 때는 다시 프로젝트 루트로 돌아오세요.",
                "5. 폴더 위치가 확실하지 않으면 명령을 추측해서 실행하지 말고 Get-Location과 Get-ChildItem으로 먼저 확인하세요.",
            ),
        ),
        SlideSpec(
            "1-6. 프로젝트의 필수 파일을 확인하세요",
            intro="새 PowerShell을 열고 아래 명령을 프로젝트 루트에서 순서대로 실행하세요. 각 명령은 파일이나 폴더가 있으면 해당 정보를 보여줍니다.",
            commands="Set-Location C:\\work\\statyearbook\n\nGet-Item .\\.env.example\nGet-Item .\\backend\\requirements.txt\nGet-Item .\\backend\\app\\main.py\nGet-Item .\\frontend\\package.json\nGet-Item .\\frontend\\package-lock.json\nGet-Item .\\frontend\\src\\App.tsx",
            kind="command",
        ),
        SlideSpec(
            "1-7. 백엔드 환경설정 파일 .env를 만드세요",
            intro="프로젝트 루트 C:\\work\\statyearbook에서 아래 명령을 실행하세요. .env.example은 공개 가능한 예시이고 .env는 실제 실행값을 담는 로컬 파일입니다.",
            commands="Set-Location C:\\work\\statyearbook\nCopy-Item .env.example .env\nnotepad .env",
            kind="command",
        ),
        SlideSpec(
            "1-8. 첫 실행용 .env 값은 안전하게 유지하세요",
            items=(
                "1. DATABASE_BACKEND=sqlite로 두세요. 처음부터 운영 PostgreSQL을 연결하면 실수로 운영 데이터를 변경할 수 있습니다.",
                "2. ANNUAL_STATISTICS_DB_PATH=backend/app/db/annual_statistics_test.sqlite로 두세요. 파일이 없으면 백엔드가 새 테스트 DB를 준비합니다.",
                "3. LLM_REVIEW_ENABLED=0, PROFILE_LLM_ENABLED=0, AUTO_IMPORT_INCLUDE_LLM=0으로 두세요. API 키가 컴퓨터 환경에 있더라도 첫 실행에서는 외부 호출을 막습니다.",
                "4. SOURCE_STORAGE_BACKEND=local과 JOB_EXECUTOR_BACKEND=inprocess로 두세요. 로컬에서는 API 프로세스가 업로드 작업을 처리하므로 별도 worker가 필요 없습니다.",
                "5. ALLOWED_ORIGINS에 http://localhost:5173과 http://127.0.0.1:5173이 포함되어 있는지 확인하고 메모장을 저장한 뒤 닫으세요.",
            ),
        ),
        SlideSpec(
            "1-9. 실제 DB와 업로드 원본을 받은 경우 먼저 백업하세요",
            items=(
                "1. backend/app/db/annual_statistics.sqlite가 있다면 실제 연보와 검수 기록이 들어 있을 수 있습니다. 백엔드를 실행하기 전에 다른 안전한 폴더로 복사하세요.",
                "2. backend/uploads 폴더도 같은 시점에 함께 복사하세요. DB에는 원본 파일 경로와 작업 기록이 있으므로 DB와 uploads 중 하나만 복원하면 과거 작업 재시도가 실패할 수 있습니다.",
                "3. 복사한 DB에서 Get-FileHash 파일경로 -Algorithm SHA256을 실행하고 해시 값을 인계 기록에 남기세요.",
                "4. 첫 설치 확인은 실제 DB가 아니라 annual_statistics_test.sqlite로 진행하세요. 화면과 서버가 정상인 것을 확인한 뒤 실제 DB 경로로 바꾸세요.",
                "5. 실제 DB로 바꾸기 전 백엔드와 worker가 모두 종료되어 있는지 확인하세요. 실행 중인 SQLite 파일을 단순 복사하면 일관성이 깨질 수 있습니다.",
            ),
        ),
        SlideSpec(
            "1-10. Python 가상환경을 만드세요",
            intro="프로젝트 루트 C:\\work\\statyearbook에서 실행하세요. .venv는 이 프로젝트만 사용하는 Python 패키지 보관 폴더입니다.",
            commands="Set-Location C:\\work\\statyearbook\npy -3.12 -m venv .venv\n& .\\.venv\\Scripts\\python.exe --version",
            kind="command",
        ),
        SlideSpec(
            "1-11. 백엔드 패키지를 설치하세요",
            intro="가상환경을 활성화하지 않아도 됩니다. 아래처럼 .venv 안의 python.exe를 직접 실행하면 PowerShell 실행 정책 문제를 피할 수 있습니다.",
            commands="Set-Location C:\\work\\statyearbook\n& .\\.venv\\Scripts\\python.exe -m pip install --upgrade pip\n& .\\.venv\\Scripts\\python.exe -m pip install -r .\\backend\\requirements.txt\n\n& .\\.venv\\Scripts\\python.exe -c \"import fastapi, openpyxl, psycopg; print('backend dependencies ok')\"",
            kind="command",
        ),
        SlideSpec(
            "1-12. 백엔드 설치 결과를 판단하세요",
            items=(
                "1. pip install 명령의 마지막 부분에 Successfully installed 또는 이미 설치되었다는 메시지가 나오면 정상입니다.",
                "2. backend dependencies ok가 출력되면 FastAPI, Excel 처리, PostgreSQL 드라이버를 가상환경에서 읽을 수 있다는 뜻입니다.",
                "3. No module named fastapi가 나오면 시스템 Python을 사용했을 가능성이 큽니다. 반드시 .venv\\Scripts\\python.exe 경로가 포함된 명령을 사용하세요.",
                "4. 다운로드가 실패하면 조직 네트워크의 프록시, 방화벽, 인증서 정책을 확인하세요. 비공식 사이트에서 패키지 파일을 따로 받지 마세요.",
                "5. .venv 폴더는 다른 컴퓨터에서 복사해 사용하지 마세요. 운영체제와 Python 경로가 달라지면 실행되지 않으므로 새 컴퓨터에서 다시 만드세요.",
            ),
        ),
        SlideSpec(
            "1-13. 프런트엔드 패키지를 설치하고 빌드하세요",
            intro="이번에는 frontend 폴더로 이동하세요. npm ci는 package-lock.json에 기록된 정확한 버전으로 패키지를 설치합니다.",
            commands="Set-Location C:\\work\\statyearbook\\frontend\nnpm ci\nnpm run build\nSet-Location C:\\work\\statyearbook",
            kind="command",
        ),
        SlideSpec(
            "1-14. 프런트엔드 설치 결과를 판단하세요",
            items=(
                "1. npm ci가 끝나면 frontend/node_modules 폴더가 생깁니다. 이 폴더는 자동 생성되므로 Git이나 인계 압축 파일에 넣을 필요가 없습니다.",
                "2. npm run build가 성공하면 TypeScript 검사와 Vite 운영 빌드가 끝나고 frontend/dist 폴더가 생깁니다.",
                "3. EBADENGINE 오류가 나오면 node --version을 확인하세요. 권장 버전은 Node.js 24 LTS입니다.",
                "4. package-lock.json과 package.json이 서로 맞지 않는 오류가 나면 임의로 lock 파일을 지우지 말고 인계자가 사용한 최종 파일인지 확인하세요.",
                "5. node_modules를 macOS나 다른 Windows 컴퓨터에서 복사했다면 삭제하고 현재 컴퓨터에서 npm ci로 다시 설치하세요.",
            ),
        ),
        SlideSpec(
            "1-15. 첫 번째 PowerShell에서 백엔드를 실행하세요",
            intro="PowerShell 창 1을 열고 반드시 프로젝트 루트에서 실행하세요. 이 창은 백엔드 로그 창이므로 서비스를 사용하는 동안 닫지 마세요.",
            commands="Set-Location C:\\work\\statyearbook\n& .\\.venv\\Scripts\\python.exe -m uvicorn app.main:app --app-dir backend --reload --host 127.0.0.1 --port 8000",
            kind="command",
        ),
        SlideSpec(
            "1-16. 백엔드가 정상인지 주소로 확인하세요",
            items=(
                "1. 백엔드 창에 Uvicorn running on http://127.0.0.1:8000이 보이면 프로세스가 시작된 것입니다.",
                "2. 브라우저에서 http://127.0.0.1:8000/health를 여세요. {\"status\":\"ok\"}가 보이면 서버 프로세스가 살아 있습니다.",
                "3. http://127.0.0.1:8000/ready를 여세요. status가 ready이고 database가 sqlite이면 DB 연결과 초기화가 정상입니다.",
                "4. http://127.0.0.1:8000/docs를 열면 FastAPI API 문서가 보여야 합니다. 이 화면은 개발 및 장애 확인용입니다.",
                "5. /health는 되는데 /ready가 실패하면 DB 경로, 쓰기 권한, 디스크 여유 공간, migration 오류를 백엔드 창에서 확인하세요.",
            ),
        ),
        SlideSpec(
            "1-17. 두 번째 PowerShell에서 프런트엔드를 실행하세요",
            intro="백엔드를 실행 중인 창은 그대로 두고 새 PowerShell 창 2를 여세요. 이번 명령은 frontend 폴더에서 실행합니다.",
            commands="Set-Location C:\\work\\statyearbook\\frontend\nnpm run dev",
            kind="command",
        ),
        SlideSpec(
            "1-18. 브라우저에서 전체 서비스를 확인하세요",
            items=(
                "1. 프런트엔드 창에 Local: http://127.0.0.1:5173과 비슷한 주소가 보이면 Chrome 또는 Edge에서 그 주소를 여세요.",
                "2. 상단에 통계연보 검수 보조라는 이름과 연보 등록, 연보 비교, 통계 연보, 통계 원장, 수정 권한 메뉴가 보이는지 확인하세요.",
                "3. 새 테스트 DB를 사용했다면 연보 목록이 비어 있는 것이 정상입니다. 데이터가 없다는 사실과 서버 오류를 구분하세요.",
                "4. 화면이 열리지만 데이터를 가져오지 못하면 먼저 /health와 /ready를 다시 확인하고 브라우저 개발자 도구의 Network 요청 주소를 확인하세요.",
                "5. 여기까지 성공하면 설치가 완료된 것입니다. 작은 테스트 HWPX 또는 Markdown을 등록해 업로드부터 상세 화면까지 확인하세요.",
            ),
        ),
        SlideSpec(
            "1-19. 실제 인계 DB로 전환하는 방법",
            items=(
                "1. 테스트가 끝나면 백엔드와 프런트엔드 창에서 Ctrl+C를 눌러 두 서버를 종료하세요.",
                "2. 실제 annual_statistics.sqlite와 backend/uploads의 백업이 존재하는지 다시 확인하세요.",
                "3. 루트 .env를 열고 ANNUAL_STATISTICS_DB_PATH=backend/app/db/annual_statistics.sqlite로 변경하세요.",
                "4. 백엔드를 다시 실행하고 /ready에서 sqlite를 확인한 다음 프런트엔드를 실행하세요.",
                "5. 화면의 연보 제목, 연도, 표 개수를 기존 담당자가 알려준 기준과 비교하세요. 예상 데이터가 없으면 빈 DB에 잘못 연결했을 가능성을 먼저 확인하세요.",
            ),
        ),
        SlideSpec(
            "1-20. 매일 서비스를 시작하는 순서",
            items=(
                "1. 프로젝트 루트에서 git status를 확인하고, .env가 테스트 DB인지 실제 DB인지 확인하세요.",
                "2. 실제 DB를 사용한다면 최근 백업 시각과 처리 중 작업이 없는지 확인하세요.",
                "3. PowerShell 창 1에서 프로젝트 루트로 이동한 뒤 .venv의 Python으로 Uvicorn을 실행하세요.",
                "4. /health와 /ready를 확인하세요. 두 주소가 정상인 뒤 PowerShell 창 2에서 frontend로 이동하여 npm run dev를 실행하세요.",
                "5. 브라우저에서 127.0.0.1:5173을 열고 대상 연보의 제목과 표 개수를 확인한 뒤 작업을 시작하세요.",
            ),
        ),
        SlideSpec(
            "1-21. 매일 서비스를 종료하는 순서",
            items=(
                "1. 화면에서 업로드나 검수가 진행 중인지 확인하세요. 진행 중인 작업을 종료하면 일부 단계가 중단될 수 있습니다.",
                "2. 프런트엔드 PowerShell에서 Ctrl+C를 누르고 종료 질문이 나오면 확인하세요.",
                "3. 백엔드 PowerShell에서 Ctrl+C를 누르고 Application shutdown complete가 나오는지 확인하세요.",
                "4. 실제 데이터를 변경했다면 DB와 uploads 백업 또는 변경 기록을 남기세요.",
                "5. 코드도 수정했다면 테스트 결과와 git status를 확인하고, 무엇을 바꿨는지 기록한 뒤 컴퓨터를 종료하세요.",
            ),
        ),
        SlideSpec(
            "1-22. 설치 중 자주 생기는 오류",
            items=(
                "1. py, node, npm, git을 찾지 못하면 프로그램 설치 후 PowerShell을 새로 열었는지 확인하세요. PATH가 반영되기 전 창에서는 계속 실패합니다.",
                "2. No module named app이 나오면 백엔드를 frontend나 backend 내부에서 잘못 실행했을 수 있습니다. 프로젝트 루트에서 --app-dir backend를 포함한 명령을 사용하세요.",
                "3. 포트 8000 또는 5173이 사용 중이라고 나오면 같은 서버를 이미 실행했는지 확인하세요. 작업 주체가 불분명한 프로세스를 무작정 종료하지 마세요.",
                "4. CORS 오류가 나오면 브라우저 주소와 .env의 ALLOWED_ORIGINS가 프로토콜, 호스트, 포트까지 정확히 같은지 확인하고 백엔드를 재시작하세요.",
                "5. SQLite database is locked가 나오면 같은 DB를 사용하는 백엔드나 worker가 여러 개 실행 중인지, DB가 OneDrive 또는 공유 드라이브에 있는지 확인하세요.",
            ),
        ),
        SlideSpec(
            "1-23. macOS에서 필요한 프로그램을 설치하세요",
            items=(
                "1. Command와 Space를 함께 누르고 Terminal을 검색하여 터미널을 여세요.",
                "2. Git이 없다면 xcode-select --install을 입력하고 Apple 명령줄 도구 설치를 완료하세요.",
                "3. python.org에서 Python 3.12 macOS universal2 installer를 설치하거나 기존 Homebrew의 python@3.12를 사용하세요.",
                "4. nodejs.org에서 Node.js 24 LTS macOS Installer를 설치하세요. 설치가 끝나면 터미널을 새로 여세요.",
                "5. git --version, python3.12 --version, node --version, npm --version을 차례로 실행하여 모두 버전이 나오는지 확인하세요.",
            ),
        ),
        SlideSpec(
            "1-24. macOS에서 패키지를 설치하세요",
            intro="프로젝트 폴더가 /Users/사용자이름/work/statyearbook에 있다고 가정합니다. 실제 사용자 이름과 경로에 맞게 첫 줄만 바꾸세요.",
            commands="cd /Users/사용자이름/work/statyearbook\ncp .env.example .env\npython3.12 -m venv .venv\n./.venv/bin/python -m pip install --upgrade pip\n./.venv/bin/python -m pip install -r backend/requirements.txt\ncd frontend\nnpm ci\nnpm run build\ncd ..",
            kind="command",
        ),
        SlideSpec(
            "1-25. macOS에서 두 서버를 실행하세요",
            items=(
                "1. 터미널 창 1에서 프로젝트 루트로 이동한 뒤 ./.venv/bin/python -m uvicorn app.main:app --app-dir backend --reload --host 127.0.0.1 --port 8000을 실행하세요.",
                "2. /health, /ready, /docs를 브라우저에서 확인하세요.",
                "3. 터미널 창 2를 새로 열고 프로젝트의 frontend 폴더로 이동한 뒤 npm run dev를 실행하세요.",
                "4. 브라우저에서 http://127.0.0.1:5173을 열고 Windows와 같은 화면이 나오는지 확인하세요.",
                "5. 종료할 때는 각 터미널에서 Ctrl+C를 누르세요. macOS에서 만든 .venv와 node_modules를 Windows로 복사하지 마세요.",
            ),
        ),
        SlideSpec(
            "2부. 시스템 기능을 사용하는 방법",
            intro="설치와 실행이 끝난 뒤 화면에서 연보를 등록하고 검수하고 수정하는 실제 순서입니다.",
            kind="section",
        ),
        SlideSpec(
            "2-1. 기본 사용 흐름",
            items=(
                "1. 상단 연보 등록에서 전체 연보 업로드 또는 증분 업로드를 선택하세요.",
                "2. 통계 연보 화면에서 등록한 연보를 선택하고 표 코드나 표명으로 검색하세요.",
                "3. 검수 진행에서 필요한 항목을 선택하고 기존 결과 재사용, 계산 profile LLM, 언어 LLM 사용 여부를 정하세요.",
                "4. 정상, 확인 필요, 오류 의심 상태를 보고 상세 화면에서 근거와 하이라이트 셀을 확인하세요.",
                "5. 필요하면 수정 비밀번호로 인증하여 표를 수정하고 다시 검수하세요. 마지막으로 표, 연보, 검수 리포트를 내려받으세요.",
            ),
        ),
        SlideSpec(
            "2-2. 전체 연보를 등록하세요",
            items=(
                "1. 화면 상단의 연보 등록을 누르고 전체 연보 업로드를 선택하세요.",
                "2. 대상 연도와 연보 이름을 입력하고 화면에서 지원하는 .hwpx 또는 .md 파일을 선택하세요. .hwp 파일은 직접 지원하지 않습니다.",
                "3. 첫 업로드에서는 LLM 옵션을 끄고 업로드를 시작하세요. 기본 최대 파일 크기는 100MB입니다.",
                "4. 화면은 작업 ID를 받아 상태를 확인합니다. 완료되면 생성된 연보가 자동 선택되는지 확인하세요.",
                "5. 실패하면 작업 ID, 파일명, 파일 SHA-256, 오류 메시지, 백엔드 로그, 사용 DB를 함께 기록하세요. 파일 이름만 바꾸어 반복 업로드하지 마세요.",
            ),
        ),
        SlideSpec(
            "2-3. 검수를 실행하세요",
            items=(
                "1. 대상 연보를 선택하고 검수 진행을 누르세요.",
                "2. 화면에서 합계, 비율, 오탈자, 번역, 메타정보, 이상치, 파란색 표기 중 필요한 항목을 선택하세요.",
                "3. 증감률은 비율 검수 범주에, 증감액은 합계 검수 범주에 포함되어 표시됩니다. 현재 화면에는 별도 증감률 선택 항목이 없습니다.",
                "4. 기존 검수 결과 재사용을 켜면 구조와 셀 값이 완전히 같은 표의 결과만 복사됩니다. 재사용할 이전 실행에 선택 항목이 포함되어 있는지 확인하세요.",
                "5. 검수 실행 후 현재 단계, 완료 단위, 전체 단위, 백분율을 확인하세요. 오래 걸려도 브라우저를 새로고침하거나 서버를 강제 종료하지 마세요.",
            ),
        ),
        SlideSpec(
            "2-4. LLM 옵션과 검수 중단을 이해하세요",
            items=(
                "1. 합계·비율 profile LLM 보조는 새 표나 구조가 바뀐 표의 계산식 초안을 보강할 때만 사용하세요.",
                "2. 오탈자·번역 LLM 검수는 대기 중인 언어 후보를 외부 모델로 판정할 때 사용하세요.",
                "3. 두 옵션은 비용과 외부 데이터 전송이 발생할 수 있습니다. 기관 정책과 비용 승인이 없으면 끄세요.",
                "4. 검수를 중단해야 하면 중단하고 저장을 누르세요. 현재 처리의 안전한 중단 지점에 도달한 뒤 멈추고 완료된 결과는 저장합니다.",
                "5. 중단 상태와 완료 상태를 구분하세요. 중단한 작업의 결과만 보고 전체 연보 검수가 끝났다고 판단하면 안 됩니다.",
            ),
        ),
        SlideSpec(
            "2-5. 검수 결과를 확인하세요",
            items=(
                "1. 오류 의심 표를 먼저 확인하세요. 계산 또는 번역 불일치 가능성이 크므로 원본과 산식 근거를 비교하세요.",
                "2. 확인 필요 표는 자동 판단이 애매한 항목입니다. 단위, 출처, 기준일, 이상치, 번역을 담당자가 확인하세요.",
                "3. 정상은 자동 규칙이 문제를 찾지 못했다는 뜻이지 원자료의 절대적인 정확성을 보증하지 않습니다. 중요한 표는 표본 대조하세요.",
                "4. 표 상세에서 검수 결과 한 건을 선택하면 대상 셀, 관련 계산 셀, 헤더가 강조되는지 확인하세요.",
                "5. 특정 검수 유형을 화면에서 숨겨도 DB 결과가 삭제되는 것은 아닙니다. 표시 필터와 실제 검수 데이터 삭제를 구분하세요.",
            ),
        ),
        SlideSpec(
            "2-6. 표를 수정하고 다시 검수하세요",
            items=(
                "1. 표 상세 화면에서 수정을 누르고 해당 통계의 수정 비밀번호를 입력하세요.",
                "2. 제목, 영문 제목, 단위, 기준일, 출처, 담당 정보, 표 셀을 수정할 수 있습니다. 표 번호를 바꿀 때는 다른 표와 중복되지 않는지 확인하세요.",
                "3. 표 구조를 복제하여 새 통계를 만들거나 허가된 통계를 삭제할 수 있으므로 대상 연보와 표를 다시 확인하세요.",
                "4. 저장 후 목록과 상세 화면에서 변경값이 다시 조회되는지 확인하세요.",
                "5. 수정 전 검수 결과는 오래된 결과가 될 수 있습니다. 수정한 항목에 해당하는 검수를 다시 실행하고 최신 실행 시각을 확인하세요.",
            ),
        ),
        SlideSpec(
            "2-7. 증분 업로드로 다음 연도 초안을 만드세요",
            items=(
                "1. 같은 추출 배치의 표정보, 표항목, 표데이터 Excel 세 파일과 기준 연보를 준비하세요.",
                "2. 연보 등록에서 증분 업로드를 선택하고 기준 연보, 대상 연도, 새 연보 제목을 정확히 선택하세요.",
                "3. 세 Excel을 각 역할에 맞는 칸에 올리고 업로드하세요. .xls와 .xlsx를 지원합니다.",
                "4. 생성된 초안에서 자동 반영 표와 확인 필요 표를 구분하고, 상세 화면에서 추천 M키와 적용 예정 셀을 확인하세요.",
                "5. 모든 확인 대상을 처리한 뒤 연보 확정을 누르고 최종 연보에서 검수를 실행하세요. 초안 생성 상태를 최종 검수 상태로 오해하지 마세요.",
            ),
        ),
        SlideSpec(
            "2-8. 조회·비교·다운로드·권한 기능을 사용하세요",
            items=(
                "1. 통계 원장에서는 통계명, 표 ID, 부서로 검색하고 기준시점을 바꾸어 표정보·표항목·표데이터의 원천 값을 확인하세요.",
                "2. 연보 비교에서는 서로 다른 두 연보를 선택하고 추가, 삭제, 변경 표와 셀을 확인하세요. 비교 결과는 오류 판정이 아니라 차이 목록입니다.",
                "3. 표 상세 다운로드에서 CSV 또는 XLSX를 선택하세요. 연보 다운에서는 전체 또는 일부 표를 골라 XLSX를 만드세요.",
                "4. 검수 리포트에서는 오류 의심과 확인 필요 상태, 검수 내역·하이라이트 표·메타데이터 파일을 선택하여 ZIP을 받으세요.",
                "5. 수정 권한 메뉴에서 통계별 비밀번호를 변경하세요. 초기값은 담당자 내선 끝 네 자리이고 번호가 없으면 0000이므로 운영 전에 모두 바꾸세요.",
            ),
        ),
        SlideSpec(
            "3부. 유지보수와 프로젝트 구조",
            intro="기능을 추가하거나 기존 기능을 수정할 때 파일 위치, 테스트, DB, 배포, 보안을 확인하는 순서입니다.",
            kind="section",
        ),
        SlideSpec(
            "3-1. 최상위 프로젝트 구조를 이해하세요",
            items=(
                "1. backend는 Python FastAPI 서버입니다. app 아래에 API, 업무 로직, 파서, 검수, DB, 작업 처리 코드가 있습니다.",
                "2. frontend는 React와 TypeScript 화면입니다. src 아래에 components, hooks, api, utils, styles, types가 있습니다.",
                "3. docs에는 기능, 설계, 배포, 인수인계 문서가 있습니다. 코드를 바꾸면 관련 문서도 함께 갱신하세요.",
                "4. Dockerfile, railway.json, railway.worker.json은 Railway 백엔드와 worker 배포에 사용합니다. vercel.json은 프런트 배포에 사용합니다.",
                "5. backend/uploads, SQLite DB, .env, .venv, node_modules, dist는 실행 과정에서 생기거나 비밀·대용량 데이터이므로 일반 소스와 구분하세요.",
            ),
        ),
        SlideSpec(
            "3-2. 요청이 처리되는 코드 흐름",
            items=(
                "1. 사용자가 화면 버튼을 누르면 frontend component가 입력을 받습니다.",
                "2. frontend hook이 loading, success, error, polling 상태를 관리하고 api/reportApi.ts가 백엔드 HTTP 요청을 보냅니다.",
                "3. backend/app/api/routes.py가 URL, 입력값, HTTP 오류를 처리합니다.",
                "4. application, services, workflows가 실제 업무 규칙과 여러 처리 단계를 실행합니다.",
                "5. repository와 db adapter가 SQLite 또는 PostgreSQL에 저장하고, storage가 로컬 폴더나 S3/R2의 원본 파일을 처리합니다.",
            ),
        ),
        SlideSpec(
            "3-3. 프런트 기능을 수정하는 순서",
            items=(
                "1. frontend/src/types.ts에서 화면이 사용하는 데이터 모양을 확인하세요.",
                "2. frontend/src/api/reportApi.ts에서 어떤 URL로 어떤 요청과 응답을 주고받는지 확인하세요.",
                "3. hooks에서 비동기 상태와 폴링, 오류 처리를 수정하세요. 컴포넌트 안에 모든 네트워크 상태를 한꺼번에 넣지 마세요.",
                "4. components에서 화면을 수정하고 styles/global.css에서 필요한 최소 스타일만 변경하세요.",
                "5. 같은 이름의 .test.tsx 또는 관련 유틸 테스트를 추가하고 npm run test와 npm run build를 모두 실행하세요.",
            ),
        ),
        SlideSpec(
            "3-4. 백엔드 기능을 추가하는 순서",
            items=(
                "1. 요구사항의 입력, 출력, 실패 조건, 권한, 저장값을 먼저 문장으로 작성하세요.",
                "2. backend/app/models에 Pydantic 요청·응답 계약을 만들고 frontend types.ts와 같은 의미인지 확인하세요.",
                "3. services 또는 application에 실제 업무 규칙을 구현하세요. routes.py 안에 긴 SQL과 계산 로직을 직접 넣지 마세요.",
                "4. api/routes.py에는 서비스 호출, 입력 검증, HTTP 상태 코드 변환을 추가하세요.",
                "5. 성공, 잘못된 입력, 없는 대상, 권한 실패, 재시도 조건에 대한 backend/tests 회귀 테스트를 추가하세요.",
            ),
        ),
        SlideSpec(
            "3-5. DB 구조를 변경하는 순서",
            items=(
                "1. backend/app/db/schema.py에서 새로 설치하는 DB의 최종 구조를 수정하세요.",
                "2. backend/app/db/migrations.py에 기존 DB가 새 구조로 이동하는 절차를 추가하세요.",
                "3. 같은 migration이 두 번 실행되어도 데이터가 중복되거나 손상되지 않는지 확인하세요.",
                "4. 새 빈 SQLite와 기존 DB 복사본 양쪽에서 테스트하고, 운영이 PostgreSQL이면 전용 테스트 PostgreSQL에서도 확인하세요.",
                "5. 앱 시작 시 migration이 자동 실행될 수 있습니다. 코드 롤백이 이미 적용된 DB 변경을 자동으로 되돌리지 않는다는 점을 배포 계획에 적으세요.",
            ),
        ),
        SlideSpec(
            "3-6. 주요 DB 테이블을 기능 묶음으로 보세요",
            items=(
                "1. annual_reports, stat_tables, stat_table_cells는 연보, 표, 셀의 본문과 메타정보를 저장합니다.",
                "2. validation_profiles, validation_runs, validation_checks, validation_jobs는 산식과 검수 실행, 결과, 진행 작업을 저장합니다.",
                "3. report_processing_jobs와 stages는 업로드 원본 위치, 처리 단계, 재시도, worker lease를 저장합니다.",
                "4. statistical_datasets, items, values와 revision 테이블은 원천 통계 원장과 변경 이력을 저장합니다.",
                "5. report_incremental 계열, edit_passwords/sessions/manual_edits, comparison 계열, glossary와 LLM usage 테이블은 각각 증분·수정·비교·언어 기능을 담당합니다.",
            ),
        ),
        SlideSpec(
            "3-7. 코드를 수정하기 전과 후의 표준 절차",
            items=(
                "1. 작업 전 요구사항과 재현 방법, 기대 결과를 적고 새 Git 브랜치를 만드세요.",
                "2. 수정 전 관련 테스트를 실행하여 기존 실패와 새 실패를 구분할 기준을 남기세요.",
                "3. 한 번에 작은 범위만 수정하고 목적과 관계없는 다른 담당자의 변경을 함께 정리하지 마세요.",
                "4. 관련 테스트, 백엔드 전체 테스트, 프런트 전체 테스트, production build 순으로 확인하세요.",
                "5. PR이나 변경 기록에 DB migration, 환경변수, 배포 순서, 롤백 방법, 실행한 테스트, 남은 문제를 적으세요.",
            ),
        ),
        SlideSpec(
            "3-8. 자동 테스트를 실행하세요",
            intro="백엔드 명령은 프로젝트 루트에서, 프런트 명령은 frontend 폴더에서 실행하세요.",
            commands="Set-Location C:\\work\\statyearbook\n$env:PYTHONPATH = \"backend\"\n& .\\.venv\\Scripts\\python.exe -m unittest discover -s backend\\tests -q\nRemove-Item Env:PYTHONPATH\n\nSet-Location C:\\work\\statyearbook\\frontend\nnpm run test\nnpm run build",
            kind="command",
        ),
        SlideSpec(
            "3-9. 현재 테스트 기준선을 사실대로 인계하세요",
            items=(
                "1. 2026-08-28 기준 프런트엔드 테스트는 9개 파일, 49개 테스트가 모두 통과했습니다.",
                "2. TypeScript 검사와 Vite production build도 성공했습니다.",
                "3. 백엔드는 397개 테스트를 실행했고 6개가 실패했으며 PostgreSQL 전용 2개는 테스트 DB가 없어 건너뛰었습니다.",
                "4. 실패 1개는 로컬 .env의 LLM 사용 가능 상태가 LLM 없음 가정에 영향을 준 문제이고, 실패 5개는 로컬 검수 결과와 승인 snapshot 불일치입니다.",
                "5. 전체 테스트가 통과한다고 보고하지 마세요. 검수 담당자의 승인 없이 baseline fixture를 통과 목적만으로 변경하지 마세요.",
            ),
        ),
        SlideSpec(
            "3-10. SQLite와 uploads를 백업하세요",
            items=(
                "1. 처리 중인 업로드와 검수가 없는지 확인한 뒤 백엔드와 worker를 종료하세요.",
                "2. annual_statistics.sqlite와 backend/uploads를 같은 시점 기준으로 승인된 백업 위치에 복사하세요.",
                "3. DB와 주요 압축 파일에 SHA-256 해시를 만들고 백업 시각, 코드 commit, 파일 크기와 함께 기록하세요.",
                "4. 복원 시험은 운영 파일을 덮어쓰지 말고 별도 복사본과 별도 서비스에서 수행하세요.",
                "5. DB만 복원하고 uploads를 누락하거나 반대로 uploads만 복원하면 과거 작업과 원본의 연결이 끊길 수 있습니다.",
            ),
        ),
        SlideSpec(
            "3-11. 운영 배포 구조를 이해하세요",
            items=(
                "1. 프런트엔드는 Vercel에서 빌드하고 VITE_API_BASE_URL에 Railway API 주소를 설정하세요. 주소 뒤에 /api를 붙이지 마세요.",
                "2. Railway API는 Dockerfile과 railway.json을 사용하고 배포 전 python -m app.db.migrate를 실행하며 /ready를 상태 확인 주소로 사용합니다.",
                "3. 별도 Railway worker는 railway.worker.json으로 실행하고 API와 같은 commit, DATABASE_URL, 객체 저장소 설정을 사용하세요.",
                "4. API와 worker를 분리하면 서로 로컬 파일을 볼 수 없으므로 S3/R2 또는 승인된 공유 저장소를 사용하세요.",
                "5. 배포 전 DB 백업, 테스트, 환경변수, CORS, 롤백 담당자를 확인하고 배포 후 health, ready, worker, 작은 업로드, 검수, 다운로드를 점검하세요.",
            ),
        ),
        SlideSpec(
            "3-12. 현재 인증 한계를 반드시 보완하세요",
            items=(
                "1. 현재 서비스에는 전체 사용자 로그인, 조직 계정 연동, 역할별 권한 관리가 없습니다.",
                "2. 읽기, 업로드, 연보 삭제, 검수, 다운로드, 운영 지표 API도 애플리케이션 공통 로그인을 요구하지 않습니다. CORS는 인증이나 방화벽이 아닙니다.",
                "3. 표 수정에만 통계별 비밀번호가 있고 초기값은 내선 끝 네 자리 또는 0000입니다. 운영 전에 모든 초기 비밀번호를 변경하세요.",
                "4. 수정 세션은 비밀번호 변경 시 폐기되지만 현재 명시적인 만료 검사가 없습니다. 보안 요구사항이면 별도 기능으로 개선하세요.",
                "5. 운영 Railway API는 VPN, SSO, 접근 제어 proxy 또는 API gateway 등 조직 승인 방식으로 보호하고 키, token, DB URL을 로그나 PPT에 넣지 마세요.",
            ),
        ),
        SlideSpec(
            "3-13. 장애가 발생하면 이 순서로 확인하세요",
            items=(
                "1. 추가 업로드, 검수, 배포를 잠시 중단하고 발생 시각, 코드 commit, 작업 ID, 화면 오류, 서버 오류를 기록하세요.",
                "2. /health로 프로세스, /ready로 DB, 프런트 Network로 API URL과 CORS를 순서대로 확인하세요.",
                "3. 어떤 .env, 어떤 DB 파일 또는 DATABASE_URL, 어떤 원본 SHA-256을 사용했는지 확인하세요.",
                "4. worker 작업이면 API와 worker의 commit, DB, storage, lease, LLM timeout을 비교하세요.",
                "5. 데이터를 삭제하거나 DB를 덮어쓰기 전에 백업과 복구 가능성을 확인하고, 가능하면 별도 복사본에서 재현하세요.",
            ),
        ),
        SlideSpec(
            "3-14. 다음 담당자의 첫날 완료 조건",
            items=(
                "1. 새 테스트 SQLite로 백엔드와 프런트엔드를 직접 실행하고 /health와 /ready를 확인하세요.",
                "2. 승인된 작은 HWPX 또는 Markdown을 등록하고 표 목록과 상세 화면을 확인하세요.",
                "3. 합계와 메타정보 검수만 먼저 실행하고 진행률, 결과, 하이라이트, 리포트 다운로드를 확인하세요.",
                "4. 작은 화면 문구를 새 브랜치에서 수정하고 프런트 테스트와 build를 실행하여 변경 절차를 연습하세요.",
                "5. 테스트 DB를 백업하고 복사본으로 복원하는 연습을 완료한 뒤 실제 DB와 운영 계정 권한을 받으세요.",
            ),
        ),
        SlideSpec(
            "최종 인계 체크리스트",
            items=(
                "1. 최종 Git branch, commit, 미커밋 변경 목록을 기존 담당자와 화면으로 확인하세요.",
                "2. 실제 DB, 테스트 DB, uploads의 크기, 백업 시각, SHA-256, 보관 위치를 기록하세요.",
                "3. GitHub, Vercel, Railway, PostgreSQL, S3/R2, LLM 계정의 소유자와 권한을 확인하세요. 비밀번호와 키 값은 PPT에 적지 마세요.",
                "4. 운영 프런트 URL, API URL, 대표 정상 연보, 승인된 검수 baseline, 알려진 예외, 장애 연락망을 기록하세요.",
                "5. 다음 담당자가 설치, 실행, 업로드, 검수, 수정, 다운로드, 백업, 롤백을 직접 재현했을 때 인수인계를 완료하세요.",
            ),
        ),
        SlideSpec(
            "문제가 생기면 다섯 가지를 먼저 기록하세요",
            intro="어떤 코드 commit인지, 어떤 환경설정인지, 어떤 DB인지, 어떤 원본 파일인지, 어떤 작업 ID인지 기록하면 대부분의 장애를 다시 재현할 수 있습니다.",
            kind="section",
        ),
    ]


def build_presentation() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    blank = presentation.slide_layouts[6]
    specs = build_specs()
    for page, spec in enumerate(specs, start=1):
        slide = presentation.slides.add_slide(blank)
        if spec.kind == "cover":
            add_cover(slide, spec, page)
        elif spec.kind == "section":
            add_section(slide, spec, page)
        elif spec.kind == "command":
            add_command(slide, spec, page)
        else:
            add_body(slide, spec, page)
    presentation.core_properties.title = "통계연보 검수 보조 시스템 인수인계 흰배경 상세 설명본"
    presentation.core_properties.subject = "빈 컴퓨터 설치부터 기능 사용과 유지보수까지"
    presentation.core_properties.author = "Project Handover"
    presentation.core_properties.keywords = "FastAPI, React, Windows, installation, handover"
    return presentation


def validate(presentation: Presentation) -> None:
    errors: list[str] = []
    if len(presentation.slides) != len(build_specs()):
        errors.append("slide count does not match source specs")
    for slide_number, slide in enumerate(presentation.slides, start=1):
        if not slide.shapes:
            errors.append(f"slide {slide_number}: no shapes")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                errors.append(f"slide {slide_number}: negative shape position")
            if shape.left + shape.width > presentation.slide_width + Inches(0.01):
                errors.append(f"slide {slide_number}: shape exceeds right edge")
            if shape.top + shape.height > presentation.slide_height + Inches(0.01):
                errors.append(f"slide {slide_number}: shape exceeds bottom edge")
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color.type is not None and run.font.color.rgb != BLACK:
                            errors.append(f"slide {slide_number}: non-black text found")
        fill = slide.background.fill
        if fill.fore_color.rgb != WHITE:
            errors.append(f"slide {slide_number}: background is not white")
    if errors:
        raise RuntimeError("presentation validation failed:\n" + "\n".join(errors))


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation = build_presentation()
    validate(presentation)
    presentation.save(OUTPUT)
    print(f"saved: {OUTPUT}")
    print(f"slides: {len(presentation.slides)}")


if __name__ == "__main__":
    main()

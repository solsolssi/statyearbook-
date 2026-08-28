"""Polish alignment across the full deck and rebuild the main flow with thin connectors."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from build_yearbook_validation_presentation import (
    C,
    MUTED,
    MUTED_2,
    TEXT,
    WHITE,
    add_bullets,
    add_pill,
    add_rect,
    add_shape,
    add_text,
    force_font_family,
)
from refine_yearbook_validation_final_report import (
    DIVIDER,
    ERROR,
    ERROR_BG,
    FONT_NAME,
    PRIMARY,
    PRIMARY_2,
    PRIMARY_LIGHT,
    PRIMARY_PALE,
    WARNING,
    WARNING_BG,
    add_image_frame,
    add_section_title,
    add_status_badge,
)
from adjust_user_modified_final_report import clear_body, clear_notes


def add_thin_connector(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = PRIMARY_2,
    width: float = 1.25,
    arrow: bool = True,
    name: str = "얇은 연결선",
):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    style = connector._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        connector._element.remove(style)
    connector.name = name
    connector.line.color.rgb = C(color)
    connector.line.width = Pt(width)
    if arrow:
        ln = connector._element.spPr.get_or_add_ln()
        tail_end = OxmlElement("a:tailEnd")
        tail_end.set("type", "triangle")
        tail_end.set("w", "sm")
        tail_end.set("len", "sm")
        ln.append(tail_end)
    return connector


def normalize_common_grid(prs: Presentation) -> None:
    """Fix every slide to the same reference-deck header and footer grid."""
    geometry = {
        "상단 분류": (0.48, 0.24, 3.91, 0.19),
        "슬라이드 제목": (0.48, 0.51, 7.70, 0.41),
        "상단 구분선": (0.48, 0.98, 9.03, 0.012),
        "슬라이드 부제목": (0.48, 1.12, 8.82, 0.30),
        "하단 문구": (0.48, 5.33, 3.55, 0.17),
        "페이지 번호": (9.23, 5.34, 0.28, 0.16),
    }
    for slide in list(prs.slides)[1:]:
        for shape in slide.shapes:
            if shape.name in geometry:
                x, y, w, h = geometry[shape.name]
                shape.left, shape.top = Inches(x), Inches(y)
                shape.width, shape.height = Inches(w), Inches(h)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.top / 914400 < 0.5:
                shape.left, shape.top = Inches(8.30), Inches(0.11)
                shape.width, shape.height = Inches(1.26), Inches(0.36)


def add_flow_node(slide, x: float, num: str, title: str, detail: str, *, active: bool = False) -> None:
    color = PRIMARY_2 if active else PRIMARY
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 1.88, 0.46, 0.46, fill=color, line=None, radius_name=f"이용 흐름_{num}")
    add_text(slide, num, x, 2.00, 0.46, 0.18, size=8.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, title, x - 0.43, 2.54, 1.32, 0.24, size=10.7, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, detail, x - 0.58, 2.86, 1.62, 0.38, size=8.3, color=MUTED, align=PP_ALIGN.CENTER)


def slide6_thin_flow(slide) -> None:
    clear_body(slide)
    nodes = [
        (1.06, "01", "연보 등록", "파일·연도 선택"),
        (3.30, "02", "검수 실행", "검수 항목 선택"),
        (5.54, "03", "결과 확인", "문제 위치·근거"),
        (7.78, "04", "연보 완성", "오류 없는 결과\n표·검수 리포트 내려받기"),
    ]
    for idx, (x, num, title, detail) in enumerate(nodes):
        add_flow_node(slide, x, num, title, detail, active=idx == 2)
        if idx < len(nodes) - 1:
            add_thin_connector(slide, x + 0.49, 2.11, nodes[idx + 1][0] - 0.03, 2.11, name=f"상단 흐름_{num}")
    add_text(slide, "오류 없음?  →  예", 6.68, 1.80, 1.20, 0.20, size=8.0, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)

    add_thin_connector(slide, 5.77, 3.22, 5.77, 3.58, color=ERROR, name="오류 분기")
    add_text(slide, "아니오  ·  오류 있음", 5.02, 3.55, 1.50, 0.20, size=8.4, color=ERROR, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 4.94, 3.86, 1.88, 0.72, fill=WARNING_BG, line=None)
    add_text(slide, "내용 수정", 5.15, 4.01, 1.46, 0.23, size=10.4, color=WARNING, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "문제 셀·표 정보", 5.15, 4.31, 1.46, 0.17, size=7.8, color=MUTED, align=PP_ALIGN.CENTER)

    # Thin two-segment return path; "재검수" is an arrow annotation, not a separate step.
    add_thin_connector(slide, 4.94, 4.18, 3.53, 4.18, arrow=False, name="재검수 복귀_가로")
    add_thin_connector(slide, 3.53, 4.18, 3.53, 3.22, name="재검수 복귀_세로")
    add_text(slide, "재검수", 3.68, 3.70, 0.76, 0.20, size=9.0, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "수정한 내용을 현재 기준으로 다시 검수", 3.43, 4.72, 3.34, 0.20, size=8.4, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide7_aligned_upload(slide) -> None:
    clear_body(slide)
    steps = [
        ("01", "파일 읽기", "연보 내용을 파싱"),
        ("02", "정보 저장", "제목·번호·표·기본 정보를 데이터베이스에 저장"),
        ("03", "연보 등록", "연보는 독립된 게시물이며 수정 및 삭제 가능"),
    ]
    columns = [0.62, 3.67, 6.72]
    for x, (num, title, detail) in zip(columns, steps):
        add_text(slide, num, x, 1.58, 0.38, 0.20, size=9.2, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.51, 1.52, 2.14, 0.22, size=10.6, color=TEXT, bold=True)
        add_text(slide, detail, x + 0.51, 1.80, 2.28, 0.34, size=8.1, color=MUTED)
        add_rect(slide, x, 2.15, 2.75, 0.014, fill=DIVIDER, line=None)

    add_image_frame(slide, 0.48, 2.34, 3.26, 2.47, title="실제 연보의 통계 예시", subtitle="연보 안의 통계표 이미지", name="이미지 삽입_연보 통계 예시")
    add_image_frame(slide, 6.26, 2.34, 3.26, 2.47, title="플랫폼에 등록된 통계", subtitle="제목·번호·표 내용이 재구성된 화면", name="이미지 삽입_플랫폼 등록 통계")
    add_rect(slide, 4.34, 2.78, 1.32, 1.52, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "업로드", 4.53, 3.02, 0.94, 0.28, size=12.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "HWPX   ·   MD", 4.46, 3.54, 1.08, 0.22, size=8.8, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
    add_thin_connector(slide, 3.76, 3.56, 4.27, 3.56, name="업로드 전")
    add_thin_connector(slide, 5.69, 3.56, 6.20, 3.56, name="업로드 후")


def add_validation_item(slide, x: float, y: float, w: float, title: str, detail: str, *, review: bool) -> None:
    add_text(slide, title, x, y, w - 1.00, 0.24, size=10.7, color=TEXT, bold=True)
    add_text(slide, detail, x, y + 0.34, w, 0.25, size=8.5, color=MUTED)
    add_status_badge(slide, "확인 필요" if review else "오류 의심", x + w - 0.94, y - 0.01, warning=review)
    add_rect(slide, x, y + 0.80, w, 0.014, fill=DIVIDER, line=None)


def slide8_aligned_checks(slide) -> None:
    clear_body(slide)
    add_section_title(slide, 0.48, 1.55, 5.72, "4종 규칙·통계 기반 검수")
    add_text(slide, "표 구조와 계산 관계를 미리 구성한 기준으로 확인", 0.48, 1.98, 5.45, 0.20, size=8.5, color=MUTED)
    add_validation_item(slide, 0.58, 2.40, 2.42, "합계", "계·합계·소계를 다시 합산", review=False)
    add_validation_item(slide, 3.28, 2.40, 2.42, "비율", "비율·구성비·증감 값을 재계산", review=False)
    add_validation_item(slide, 0.58, 3.60, 2.42, "메타데이터", "단위·기준일·출처 확인", review=True)
    add_validation_item(slide, 3.28, 3.60, 2.42, "이상치", "다른 값과 차이가 큰 수치 후보", review=True)

    add_rect(slide, 6.34, 1.47, 3.18, 3.43, fill=PRIMARY_PALE, line=None)
    add_section_title(slide, 6.58, 1.55, 2.70, "2종 LLM 기반 검수")
    add_text(slide, "공식 용어·이전 결과를 먼저 적용하고\n남은 후보만 확인", 6.52, 1.98, 2.76, 0.42, size=7.9, color=MUTED)
    add_validation_item(slide, 6.60, 2.67, 2.46, "오탈자", "철자·문자 깨짐·표기 불일치", review=False)
    add_validation_item(slide, 6.60, 3.76, 2.46, "번역", "한글·영문 의미 대응 확인", review=True)


def slide9_aligned_profiles(slide) -> None:
    clear_body(slide)
    stages = [
        (0.48, "01", "공통 기준", "모든 표의 필수 정보 확인\n숫자 열의 특이값 후보 생성", ("메타", "이상치")),
        (2.78, "02", "표 유형별 기준", "지역별·연도별 등 구조 파악\n전국계와 지역 값 관계 적용", ("합계",)),
        (5.08, "03", "표별 자동 분석", "계·합계·소계 항목 탐색\n분자·분모·증감 관계 탐색", ("합계", "비율")),
        (7.38, "04", "통계별 예외 기준", "담당자가 확인한 규칙 반영\n고정 셀·분할표 예외 처리", ("오탈자", "번역", "합계", "비율")),
    ]
    add_rect(slide, 0.72, 2.06, 8.12, 0.018, fill="B9D0E1", line=None)
    for idx, (x, num, title, detail, tags) in enumerate(stages):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.02, 1.83, 0.46, 0.46, fill=PRIMARY_2 if idx == 2 else PRIMARY, line=None, radius_name=f"검수 기준_{num}")
        add_text(slide, num, x + 0.02, 1.95, 0.46, 0.18, size=8.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x, 2.50, 2.00, 0.25, size=10.7, color=TEXT, bold=True)
        add_text(slide, detail, x, 2.93, 2.00, 0.54, size=8.4, color=MUTED)
        chip_w = 0.44 if len(tags) == 4 else 0.70
        chip_x = x
        for tag in tags:
            add_rect(slide, chip_x, 3.87, chip_w, 0.28, fill=PRIMARY_LIGHT, line=None)
            add_text(slide, tag, chip_x, 3.94, chip_w, 0.17, size=7.3 if chip_w < 0.6 else 7.9, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
            chip_x += chip_w + 0.05


def add_recheck_rule(slide, y: float, number: str, title: str, detail: str, *, highlighted: bool) -> None:
    if highlighted:
        add_rect(slide, 5.05, y, 4.47, 0.56, fill=WARNING_BG, line=None)
        add_rect(slide, 5.05, y, 0.05, 0.56, fill=WARNING, line=None)
        number_color = WARNING
    else:
        add_rect(slide, 5.05, y + 0.55, 4.47, 0.014, fill=DIVIDER, line=None)
        number_color = PRIMARY_2
    add_text(slide, number, 5.20, y + 0.15, 0.38, 0.20, size=9.1, color=number_color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, 5.70, y + 0.08, 1.34, 0.22, size=10.5, color=TEXT, bold=True)
    add_text(slide, detail, 5.70, y + 0.33, 3.50, 0.18, size=8.3, color=MUTED)


def slide14_aligned_edit(slide) -> None:
    clear_body(slide)
    add_rect(slide, 0.62, 1.49, 8.76, 0.62, fill=PRIMARY_PALE, line=None)
    flow = [("수정 저장", PRIMARY), ("재검수 필요", WARNING), ("검수 실행", PRIMARY_2), ("새 결과 갱신", PRIMARY)]
    flow_x = [0.94, 3.05, 5.16, 7.27]
    for idx, ((label, color), x) in enumerate(zip(flow, flow_x)):
        add_text(slide, label, x, 1.70, 1.34, 0.20, size=8.9, color=color, bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_thin_connector(slide, x + 1.43, 1.80, flow_x[idx + 1] - 0.09, 1.80, color=WARNING if idx == 0 else PRIMARY_2, name=f"재검수 흐름_{idx}")

    add_section_title(slide, 0.48, 2.38, 4.06, "수정 보호")
    col_x = [0.48, 1.95, 3.42]
    add_rect(slide, 1.04, 3.23, 2.94, 0.018, fill="B9D0E1", line=None)
    protection = [("01", "비밀번호 확인", "통계별 수정 권한 확인"), ("02", "편집 화면", "표 정보·셀 값·행/열 수정"), ("03", "수정 저장", "변경 내용 반영")]
    for x, (num, title, detail) in zip(col_x, protection):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.38, 3.01, 0.44, 0.44, fill=PRIMARY_2, line=None, radius_name=f"수정 보호_{num}")
        add_text(slide, num, x + 0.38, 3.12, 0.44, 0.18, size=7.8, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x, 3.58, 1.20, 0.22, size=9.4, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x, 3.88, 1.20, 0.36, size=7.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "저장 즉시 검수 자동 실행 X  ·  담당자가 재검수 실행", 0.64, 4.53, 3.70, 0.21, size=8.2, color=WARNING, bold=True, align=PP_ALIGN.CENTER)

    add_section_title(slide, 5.05, 2.38, 4.47, "재검수 원칙")
    add_recheck_rule(slide, 2.90, "01", "값 변경", "기존 계산 기준으로 현재 값 재검수", highlighted=True)
    add_recheck_rule(slide, 3.58, "02", "구조 변경", "새로운 표 구조에 맞춰 기준 재생성", highlighted=True)
    add_recheck_rule(slide, 4.26, "03", "변경 없음", "이전 검수 결과 재사용", highlighted=False)


def slide16_aligned_effects(slide) -> None:
    clear_body(slide)
    groups = [
        (0.48, "통계연보 담당자", "전체 진행 상황 관리", ("연보별 검수 상태 한눈에 확인", "오류 통계와 수정 현황 공유", "발간 전 최종 결과 정리"), PRIMARY),
        (3.60, "각 통계 담당자", "내 통계의 문제를 바로 수정", ("문제 셀과 근거를 빠르게 확인", "화면에서 직접 수정·저장", "변경된 표 중심으로 재검수"), PRIMARY_2),
        (6.72, "검수 진행자", "반복 계산 대신 판단에 집중", ("합계·비율 자동 재계산", "확인 우선순위와 근거 제시", "검수 리포트로 결과 공유"), PRIMARY),
    ]
    for idx, (x, title, subtitle, bullets, color) in enumerate(groups):
        add_text(slide, title, x, 1.80, 2.80, 0.27, size=11.2, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, subtitle, x, 2.20, 2.80, 0.22, size=8.7, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_rect(slide, x + 0.12, 2.59, 2.56, 0.016, fill=DIVIDER, line=None)
        add_bullets(slide, bullets, x + 0.10, 2.93, 2.60, 1.40, size=8.7, color=TEXT, bullet_color=color, gap=11)
        if idx < 2:
            add_rect(slide, x + 2.94, 1.76, 0.014, 2.68, fill=DIVIDER, line=None)
    add_rect(slide, 0.88, 4.68, 8.24, 0.36, fill=PRIMARY, line=None)
    add_text(slide, "검수 시간 단축  ·  오류 누락 위험 감소  ·  공유·수정·재검수 흐름 단순화", 1.18, 4.78, 7.64, 0.19, size=9.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide17_aligned_future(slide) -> None:
    clear_body(slide)
    groups = [
        (0.48, "01", "권한·보안", ("기관 계정과 통합 로그인", "담당자별 역할·수정 권한", "세션 만료·접속 제한")),
        (3.60, "02", "수정 이력 기록", ("셀 단위 변경 전·후 기록", "수정자·시간·사유 확인", "되돌리기·변경 알림")),
        (6.72, "03", "내부망·내부 LLM", ("내부망 환경에 서비스 도입", "외부 LLM을 내부 모델로 전환", "민감 정보와 사용 이력 관리")),
    ]
    add_rect(slide, 1.88, 2.05, 6.00, 0.018, fill="B9D0E1", line=None)
    for idx, (x, num, title, bullets) in enumerate(groups):
        color = PRIMARY_2 if idx == 1 else PRIMARY
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 1.16, 1.81, 0.48, 0.48, fill=color, line=None, radius_name=f"향후 보완_{num}")
        add_text(slide, num, x + 1.16, 1.93, 0.48, 0.18, size=8.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x, 2.52, 2.80, 0.25, size=11.1, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_rect(slide, x + 0.12, 2.96, 2.56, 0.016, fill=DIVIDER, line=None)
        add_bullets(slide, bullets, x + 0.10, 3.28, 2.60, 1.36, size=8.7, color=TEXT, bullet_color=PRIMARY_2, gap=11)


def build(source: Path, output: Path) -> None:
    prs = Presentation(source)
    normalize_common_grid(prs)
    slide6_thin_flow(prs.slides[5])
    slide7_aligned_upload(prs.slides[6])
    slide8_aligned_checks(prs.slides[7])
    slide9_aligned_profiles(prs.slides[8])
    slide14_aligned_edit(prs.slides[13])
    slide16_aligned_effects(prs.slides[15])
    slide17_aligned_future(prs.slides[16])
    clear_notes(prs)
    force_font_family(prs, FONT_NAME)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()

"""Adjust layout of the user's edited final-report deck without rewriting its content."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from build_yearbook_validation_presentation import (
    MUTED,
    TEXT,
    WHITE,
    add_arrow,
    add_down_arrow,
    add_pill,
    add_rect,
    add_shape,
    add_text,
    force_font_family,
)
from refine_yearbook_validation_final_report import (
    DIVIDER,
    ERROR,
    FONT_NAME,
    PRIMARY,
    PRIMARY_2,
    PRIMARY_LIGHT,
    PRIMARY_PALE,
    WARNING,
    WARNING_BG,
    add_image_frame,
    add_open_item,
    add_section_title,
)


def clear_body(slide) -> None:
    """Remove only content shapes, preserving the user's header, footer and page number."""
    for shape in list(slide.shapes):
        top = shape.top / 914400
        if 1.43 <= top < 5.24:
            element = shape._element
            element.getparent().remove(element)


def slide6_flow(slide) -> None:
    clear_body(slide)
    nodes = [
        (1.06, "01", "연보 등록", "파일·연도 선택"),
        (3.30, "02", "검수 실행", "검수 항목 선택"),
        (5.54, "03", "결과 확인", "문제 위치·근거"),
        (7.78, "04", "연보 완성", "오류 없는 결과\n표·검수 리포트 내려받기"),
    ]
    add_rect(slide, 1.28, 2.12, 6.72, 0.025, fill="B9D0E1", line=None)
    for idx, (x, num, title, detail) in enumerate(nodes):
        color = PRIMARY_2 if idx == 2 else PRIMARY
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 1.89, 0.46, 0.46, fill=color, line=None, radius_name=f"이용 흐름_{num}")
        add_text(slide, num, x, 2.01, 0.46, 0.18, size=8.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x - 0.43, 2.56, 1.32, 0.24, size=10.7, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x - 0.56, 2.87, 1.58, 0.38, size=8.3, color=MUTED, align=PP_ALIGN.CENTER)

    add_text(slide, "오류 없음?  →  예", 6.70, 1.81, 1.20, 0.20, size=8.0, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 5.66, 3.22, 0.24, 0.34, color=ERROR)
    add_text(slide, "아니오  ·  오류 있음", 5.02, 3.54, 1.50, 0.20, size=8.4, color=ERROR, bold=True, align=PP_ALIGN.CENTER)

    # No separate re-validation node: edit, then return directly to validation.
    add_rect(slide, 4.94, 3.86, 1.88, 0.72, fill=WARNING_BG, line=None)
    add_text(slide, "내용 수정", 5.15, 4.01, 1.46, 0.23, size=10.4, color=WARNING, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "문제 셀·표 정보", 5.15, 4.31, 1.46, 0.17, size=7.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.LEFT_ARROW, 4.18, 4.06, 0.54, 0.28, fill=PRIMARY_2, line=None, radius_name="수정에서 검수 실행으로")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.UP_ARROW, 3.49, 3.11, 0.38, 1.10, fill=PRIMARY_2, line=None, radius_name="검수 실행 복귀")
    add_text(slide, "재검수", 3.73, 3.70, 0.72, 0.20, size=9.0, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "수정한 내용을 현재 기준으로 다시 검수", 3.42, 4.71, 3.35, 0.20, size=8.4, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide7_upload(slide) -> None:
    clear_body(slide)

    # Processing steps appear before the visual example.
    steps = [
        ("01", "파일 읽기", "연보 내용을 파싱"),
        ("02", "정보 저장", "제목·번호·표·기본 정보를 데이터베이스에 저장"),
        ("03", "연보 등록", "연보는 독립된 게시물이며 수정 및 삭제 가능"),
    ]
    for idx, (num, title, detail) in enumerate(steps):
        x = 0.62 + idx * 3.05
        add_open_item(slide, x, 1.52, 2.75, num, title, detail, color=PRIMARY_2)

    add_image_frame(slide, 0.48, 2.33, 3.26, 2.48, title="실제 연보의 통계 예시", subtitle="연보 안의 통계표 이미지", name="이미지 삽입_연보 통계 예시")
    add_image_frame(slide, 6.26, 2.33, 3.26, 2.48, title="플랫폼에 등록된 통계", subtitle="제목·번호·표 내용이 재구성된 화면", name="이미지 삽입_플랫폼 등록 통계")
    add_arrow(slide, 3.90, 3.35, 0.34, 0.34, color=PRIMARY_2)
    add_rect(slide, 4.34, 2.76, 1.32, 1.54, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "업로드", 4.53, 3.00, 0.94, 0.28, size=12.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "HWPX   ·   MD", 4.46, 3.52, 1.08, 0.22, size=8.8, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 5.78, 3.35, 0.34, 0.34, color=PRIMARY_2)


def slide12_language(slide) -> None:
    clear_body(slide)
    add_rect(slide, 0.62, 1.47, 8.76, 0.43, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "후보 수집", 0.88, 1.59, 1.08, 0.18, size=9.4, color=PRIMARY_2, bold=True)
    add_text(slide, "표 제목 · 헤더 · 한글/영문 셀에서 검수할 문구를 수집", 2.08, 1.57, 6.86, 0.20, size=9.0, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 4.84, 1.97, 0.30, 0.28, color=PRIMARY_2)

    filters = [
        (1.28, "01", "공식 용어·번역 DB", "기관명  행정안전부\n지역명  세종특별자치시·연기면"),
        (5.20, "02", "반복 표현 DB", "계 · 합계 · 소계\n연보 전체에서 자주 쓰는 표현"),
    ]
    for x, num, title, detail in filters:
        add_text(slide, num, x, 2.36, 0.42, 0.20, size=9.0, color=PRIMARY_2, bold=True)
        add_text(slide, title, x + 0.48, 2.31, 2.72, 0.25, size=10.4, color=TEXT, bold=True)
        add_rect(slide, x, 2.72, 3.48, 0.016, fill=PRIMARY_2, line=None)
        add_text(slide, detail, x + 0.08, 2.92, 3.32, 0.46, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        add_down_arrow(slide, x + 1.60, 3.43, 0.26, 0.28, color=PRIMARY_2)

    # One integrated result area contains both the LLM review and its effects.
    add_rect(slide, 2.06, 3.82, 5.88, 1.02, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "남은 후보만 LLM 검수", 2.42, 3.96, 5.16, 0.24, size=11.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 2.40, 4.31, 5.20, 0.014, fill="C7D9E7", line=None)
    add_pill(slide, "토큰 사용량 감소", 2.66, 4.45, 1.62, 0.28, fill=WARNING_BG, color=WARNING, size=8.2, line=None)
    add_pill(slide, "공식 표현 일관성 확보", 5.12, 4.45, 2.18, 0.28, fill=WHITE, color=PRIMARY_2, size=8.2, line=None)


def add_recheck_rule(slide, y: float, number: str, title: str, detail: str, *, highlighted: bool) -> None:
    if highlighted:
        add_rect(slide, 5.05, y, 4.47, 0.56, fill=WARNING_BG, line=None)
        add_rect(slide, 5.05, y, 0.05, 0.56, fill=WARNING, line=None)
        number_color = WARNING
    else:
        add_rect(slide, 5.05, y + 0.55, 4.47, 0.014, fill=DIVIDER, line=None)
        number_color = PRIMARY_2
    add_text(slide, number, 5.20, y + 0.15, 0.38, 0.20, size=9.1, color=number_color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, 5.70, y + 0.08, 1.34, 0.22, size=10.5, color=TEXT, bold=True)
    add_text(slide, detail, 5.70, y + 0.33, 3.50, 0.18, size=8.3, color=MUTED)


def slide14_edit_recheck(slide) -> None:
    clear_body(slide)

    # The re-validation flow is the first element read after the subtitle.
    add_rect(slide, 0.62, 1.49, 8.76, 0.62, fill=PRIMARY_PALE, line=None)
    flow = [("수정 저장", PRIMARY), ("재검수 필요", WARNING), ("검수 실행", PRIMARY_2), ("새 결과 갱신", PRIMARY)]
    for idx, (label, color) in enumerate(flow):
        x = 0.92 + idx * 2.12
        add_text(slide, label, x, 1.70, 1.30, 0.20, size=8.9, color=color, bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_arrow(slide, x + 1.52, 1.66, 0.26, 0.26, color=WARNING if idx == 0 else PRIMARY_2)

    add_section_title(slide, 0.48, 2.38, 4.06, "수정 보호")
    protection = [
        ("01", "비밀번호 확인", "통계별 수정 권한 확인"),
        ("02", "편집 화면", "표 정보·셀 값·행/열 수정"),
        ("03", "수정 저장", "변경 내용 반영"),
    ]
    add_rect(slide, 0.86, 3.23, 3.15, 0.020, fill="B9D0E1", line=None)
    for idx, (num, title, detail) in enumerate(protection):
        x = 0.60 + idx * 1.34
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 3.01, 0.44, 0.44, fill=PRIMARY_2, line=None, radius_name=f"수정 보호_{num}")
        add_text(slide, num, x, 3.12, 0.44, 0.18, size=7.8, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x - 0.24, 3.58, 0.94, 0.22, size=9.4, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x - 0.36, 3.88, 1.18, 0.36, size=7.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "저장 즉시 검수 자동 실행 X  ·  담당자가 재검수 실행", 0.64, 4.53, 3.70, 0.21, size=8.2, color=WARNING, bold=True, align=PP_ALIGN.CENTER)

    add_section_title(slide, 5.05, 2.38, 4.47, "재검수 원칙")
    add_recheck_rule(slide, 2.90, "01", "값 변경", "기존 계산 기준으로 현재 값 재검수", highlighted=True)
    add_recheck_rule(slide, 3.58, "02", "구조 변경", "새로운 표 구조에 맞춰 기준 재생성", highlighted=True)
    add_recheck_rule(slide, 4.26, "03", "변경 없음", "이전 검수 결과 재사용", highlighted=False)


def clear_notes(prs: Presentation) -> None:
    for slide in prs.slides:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            frame.clear()
            frame.text = ""


def build(source: Path, output: Path) -> None:
    prs = Presentation(source)
    slide6_flow(prs.slides[5])
    slide7_upload(prs.slides[6])
    slide12_language(prs.slides[11])
    slide14_edit_recheck(prs.slides[13])
    clear_notes(prs)
    force_font_family(prs, FONT_NAME)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()

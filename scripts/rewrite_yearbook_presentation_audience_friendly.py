"""Rewrite the yearbook validation deck for first-time, non-technical users."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_yearbook_validation_presentation import force_font_family


FONT_NAME = "Noto Sans KR"
TEXT = "172331"
MUTED = "667085"


SLIDE_TEXT: dict[int, dict[str, str]] = {
    1: {
        "Google Shape;65;p15": "서비스 소개",
    },
    2: {
        "Google Shape;93;p16": "서비스가 필요한 이유",
        "Google Shape;94;p16": "왜 검수 보조가 필요한가",
        "Google Shape;96;p16": "수백 개 표를 사람이 반복 확인하면 시간과 오류 누락 위험이 커집니다.",
        "Google Shape;76;p16": "수백 개 표를 사람이 반복 확인",
        "Google Shape;77;p16": "- 표마다 값을 다시 대조해 검수 시간이 늘어남",
        "Google Shape;80;p16": "계산·문구 오류를 한눈에 찾기 어려움",
        "Google Shape;85;p16": "- 합계·비율 오류와 오탈자·번역 문제를 놓치기 쉬움",
        "Google Shape;84;p16": "해마다 비슷한 표를 다시 검수",
        "Google Shape;81;p16": "- 이전 검수 기준과 결과를 다음 연도에 활용하기 어려움",
        "Google Shape;88;p16": "시스템이 필요한 이유",
        "Google Shape;90;p16": "반복 계산과 기본 확인은 시스템이 돕고,\v담당자는 확인이 필요한 표와 문구에 집중합니다.",
    },
    3: {
        "Google Shape;101;p17": "문제 이해",
        "Google Shape;102;p17": "해마다 반복되는 검수",
        "Google Shape;105;p17": "계산 오류는 놓치기 쉽고, 표 구조가 같아도 셀 값은 매년 달라집니다.",
        "Google Shape;110;p17": "오류를 찾기 어려운 표",
        "Google Shape;114;p17": "합계 불일치나 오탈자가 보이는 화면 권장",
        "Google Shape;116;p17": "합계 불일치  ·  오탈자  ·  번역 오류",
        "Google Shape;119;p17": "같은 구조, 달라지는 값",
        "Google Shape;130;p17": "표 구조·항목은 같음",
        "Google Shape;132;p17": "올해 값만 변경",
    },
    4: {
        "Google Shape;138;p18": "서비스 소개",
        "Google Shape;139;p18": "서비스가 바꾸는 검수 방식",
        "Google Shape;142;p18": "반복 확인은 시스템이 돕고, 담당자는 최종 판단에 집중합니다.",
        "Google Shape;147;p18": "기존 방식",
        "Google Shape;150;p18": "검수 보조 시스템 사용",
        "Google Shape;154;p18": "수백 개 표를 직접 대조",
        "Google Shape;157;p18": "검수할 항목을 선택해 일괄 확인",
        "Google Shape;161;p18": "계산·문구 오류 누락 위험",
        "Google Shape;164;p18": "계산 오류와 문구 후보를 함께 제시",
        "Google Shape;168;p18": "해마다 같은 기준을 다시 확인",
        "Google Shape;171;p18": "같은 구조의 표는 이전 기준 활용",
        "Google Shape;175;p18": "검수 결과 공유 어려움",
        "Google Shape;178;p18": "화면과 보고서로 결과 공유",
        "Google Shape;181;p18": "반복 대조 업무를 줄이고, 확인·수정·재검수를 하나의 흐름으로 연결",
    },
    5: {
        "Google Shape;187;p19": "통계연보 검수 보조 시스템",
        "Google Shape;190;p19": "연보 등록부터 오류 확인, 수정, 결과 내려받기까지 한 화면에서 지원합니다.",
        "Google Shape;197;p19": "검수할 연보 선택",
        "Google Shape;198;p19": "대상 연도와 통계표 검색",
        "Google Shape;202;p19": "검수 결과 한눈에 확인",
        "Google Shape;203;p19": "정상·확인 필요·오류 의심 건수",
        "Google Shape;207;p19": "문제 위치와 근거 확인",
        "Google Shape;208;p19": "현재값과 다시 계산한 값을 함께 확인",
    },
    6: {
        "Google Shape;213;p20": "처음 검수 흐름",
        "Google Shape;214;p20": "처음 검수는 이렇게 진행합니다",
        "Google Shape;217;p20": "연보를 등록한 뒤 항목을 선택하면 결과 확인과 재검수까지 이어집니다.",
        "Google Shape;222;p20": "1  연보 등록",
        "Google Shape;223;p20": "파일과 대상 연도 선택",
        "Google Shape;226;p20": "2  내용 확인",
        "Google Shape;227;p20": "표 번호·제목·셀 확인",
        "Google Shape;230;p20": "3  항목 선택·실행",
        "Google Shape;231;p20": "합계·비율·문구 등 선택",
        "Google Shape;234;p20": "4  결과 확인",
        "Google Shape;235;p20": "상태·문제 위치·근거 확인",
        "Google Shape;238;p20": "5  수정·재검수",
        "Google Shape;239;p20": "문제 수정 후 다시 확인",
        "Google Shape;242;p20": "6  결과 활용",
        "Google Shape;243;p20": "표·연보·보고서 저장",
        "Google Shape;246;p20": "처음 검수 핵심",
        "Google Shape;247;p20": "연보 등록 → 내용 확인 → 항목 선택 → 결과 확인 순서로 진행하고, 필요할 때 수정·재검수를 이어갑니다.",
    },
    7: {
        "Google Shape;252;p21": "1단계 · 연보 등록",
        "Google Shape;253;p21": "검수할 연보 등록",
        "Google Shape;256;p21": "새 연보를 올리거나 기존 연보에 새 연도 자료를 추가합니다.",
        "Google Shape;263;p21": "완성된 연보 파일 등록",
        "Google Shape;266;p21": "함께 입력",
        "Google Shape;267;p21": "대상 연도 · 연보명 · 원본 파일",
        "Google Shape;268;p21": "등록 결과",
        "Google Shape;270;p21": "연보 초안",
        "Google Shape;271;p21": "화면에서 지원하는 파일 형식",
        "Google Shape;276;p21": "기존 연보에 새 연도 자료 추가",
        "Google Shape;278;p21": "표 정보 · 항목 · 데이터 파일",
        "Google Shape;279;p21": "함께 입력",
        "Google Shape;280;p21": "대상 연도 · 기준 연보 · Excel 파일 3개",
        "Google Shape;281;p21": "등록 결과",
        "Google Shape;283;p21": "새 연도 초안",
        "Google Shape;287;p21": "파일 유형 확인",
        "Google Shape;289;p21": "표 번호·제목·중복 확인",
        "Google Shape;291;p21": "표·셀·기본 정보 저장",
        "Google Shape;293;p21": "내용 확인 후 연보 확정",
        "Google Shape;294;p21": "표 구조를 판단하기 어려운 경우 ‘확인 필요’로 표시하고 담당자가 확인합니다.",
    },
    8: {
        "Google Shape;299;p22": "2단계 · 검수 항목 선택",
        "Google Shape;300;p22": "무엇을 검수할지 선택",
        "Google Shape;303;p22": "계산·표 정보·문구 중 필요한 항목을 선택해 한 번에 실행합니다.",
        "Google Shape;308;p22": "계산·표 정보 확인",
        "Google Shape;311;p22": "더해야 할 값이 맞는지 다시 계산",
        "Google Shape;314;p22": "비율·증감 값을 다시 계산해 비교",
        "Google Shape;316;p22": "표 정보",
        "Google Shape;317;p22": "단위·기준일·출처 누락 확인",
        "Google Shape;319;p22": "특이값",
        "Google Shape;320;p22": "다른 값과 차이가 큰 숫자 후보",
        "Google Shape;323;p22": "오탈자·번역 확인",
        "Google Shape;326;p22": "깨진 문자·잘못된 표기 후보",
        "Google Shape;329;p22": "한글·영문 표현의 대응 확인",
        "Google Shape;331;p22": "공식 자료  →  이전 검수 결과  →  필요한 경우 AI 확인",
        "Google Shape;337;p22": "검수 결과에서 문제 없음",
        "Google Shape;342;p22": "담당자가 내용을 보고 판단",
        "Google Shape;347;p22": "계산 불일치 등 먼저 확인",
    },
    9: {
        "Google Shape;352;p23": "3단계 · 검수 기준 준비",
        "Google Shape;353;p23": "표마다 맞는 검수 기준 준비",
        "Google Shape;356;p23": "시스템이 표 구조를 읽고, 어떤 셀을 더하고 무엇으로 나눌지 기준을 마련합니다.",
        "Google Shape;361;p23": "표 구조 확인",
        "Google Shape;362;p23": "헤더 수 · 행/열 수 · 항목명 · 데이터 배치를 분석해 표마다 검수 기준을 준비",
        "Google Shape;366;p23": "공통 기준",
        "Google Shape;367;p23": "• 모든 표에 기본 확인 적용\v• 단위·기준일·출처 확인\v• 숫자 열의 특이값 확인 가능성 판단",
        "Google Shape;371;p23": "표 유형별 기준",
        "Google Shape;372;p23": "• 지역별·연도별 등 표 종류 파악\v• 전국 합계와 지역 값의 관계 확인\v• 헤더와 항목 배치를 기본 기준으로 사용",
        "Google Shape;376;p23": "표별 자동 분석",
        "Google Shape;377;p23": "• 계·합계·소계 항목 찾기\v• 분자·분모·증감 관계 찾기\v• 현재 값과 비교해 기준의 신뢰도 확인",
        "Google Shape;381;p23": "통계별 예외 기준",
        "Google Shape;382;p23": "• 담당자가 확인한 기존 기준 반영\v• 잘못 찾은 관계는 제외·보정\v• 나뉜 표·고정 셀 등 예외 처리",
        "Google Shape;384;p23": "준비 결과",
        "Google Shape;385;p23": "표마다 합계·비율 계산 기준을 저장  |  같은 구조는 재사용  |  구조가 바뀐 표만 다시 분석",
    },
    10: {
        "Google Shape;390;p24": "4단계 · 합계 확인",
        "Google Shape;391;p24": "합계가 맞는지 확인",
        "Google Shape;394;p24": "계·합계·소계 셀을 구성 항목에서 다시 더해 표시된 값과 비교합니다.",
        "Google Shape;399;p24": "어떤 합계를 확인하나",
        "Google Shape;403;p24": "행 합계",
        "Google Shape;404;p24": "같은 행의 세부 열 더하기",
        "Google Shape;408;p24": "열 합계·소계",
        "Google Shape;409;p24": "세부 행·지역 값 더하기",
        "Google Shape;413;p24": "표에 적힌 계산식",
        "Google Shape;414;p24": "a+b·a-b 관계 계산",
        "Google Shape;418;p24": "나뉜 표·반복 구역",
        "Google Shape;419;p24": "연결된 셀 값 합산",
        "Google Shape;421;p24": "확인 기준",
        "Google Shape;422;p24": "평균 열 제외 · 허용 범위 비교",
        "Google Shape;425;p24": "합계 확인 화면 이미지 삽입",
        "Google Shape;426;p24": "오류 셀·현재값·다시 계산한 값·근거 셀이 보이는 화면 권장",
        "Google Shape;428;p24": "화면에서 볼 내용",
        "Google Shape;429;p24": "오류 위치  ·  현재값  ·  다시 계산한 값  ·  판정 근거",
    },
    11: {
        "Google Shape;434;p25": "5단계 · 비율·증감 확인",
        "Google Shape;435;p25": "비율과 증감 값이 맞는지 확인",
        "Google Shape;438;p25": "분자·분모와 전년·당해 값을 다시 계산해 표시된 값과 비교합니다.",
        "Google Shape;443;p25": "어떤 값을 확인하나",
        "Google Shape;447;p25": "비율",
        "Google Shape;448;p25": "분자 ÷ 분모 × 100",
        "Google Shape;452;p25": "전체 대비 구성비",
        "Google Shape;453;p25": "항목 값 ÷ 전체 합계 × 100",
        "Google Shape;457;p25": "건당·인당 값",
        "Google Shape;458;p25": "건수·인원 기준 재계산",
        "Google Shape;462;p25": "증감 값",
        "Google Shape;463;p25": "당해-전년·증감률 재계산",
        "Google Shape;465;p25": "확인 기준",
        "Google Shape;466;p25": "총계는 다시 계산 · 분모 0 제외",
        "Google Shape;469;p25": "비율·증감 확인 화면 이미지 삽입",
        "Google Shape;470;p25": "분자·분모·전년값·당해값과 판정 결과가 보이는 화면 권장",
        "Google Shape;472;p25": "화면에서 볼 내용",
        "Google Shape;473;p25": "오류 위치  ·  현재값  ·  다시 계산한 값  ·  판정 근거",
    },
    12: {
        "Google Shape;478;p26": "6단계 · 오탈자·번역 확인",
        "Google Shape;479;p26": "오탈자·번역 후보 찾기",
        "Google Shape;482;p26": "표 제목과 한글·영문 표현을 살펴 담당자가 확인할 문구만 추립니다.",
        "Google Shape;488;p26": "검수할 문구 모으기",
        "Google Shape;489;p26": "표 제목 · 대제목 · 하위 표 제목 · 헤더 · 한글/영문이 있는 셀",
        "Google Shape;494;p26": "확인 대상만 추리기",
        "Google Shape;495;p26": "숫자만 있는 값은 제외하고, 긴 문장은 확인하기 좋은 길이로 나눔",
        "Google Shape;498;p26": "• 철자·깨진 문자·중복 문장부호 확인\v• 숫자 구분기호·단위 표기 확인",
        "Google Shape;500;p26": "• 한글과 영문이 함께 있는 표현의 대응 확인\v• 한국어만 있는 셀을 영문 누락으로 판단하지 않음",
    },
    13: {
        "Google Shape;509;p27": "6단계 · 오탈자·번역 확인",
        "Google Shape;510;p27": "오탈자·번역을 확인하는 순서",
        "Google Shape;513;p27": "검증된 표현을 먼저 적용하고, 판단이 필요한 문구만 AI가 보조합니다.",
        "Google Shape;519;p27": "공식 명칭·기본 용어 확인",
        "Google Shape;520;p27": "정확히 일치하는 표현을 먼저 처리",
        "Google Shape;525;p27": "이전 검수 결과 재사용",
        "Google Shape;526;p27": "같은 위치와 문맥의 확인 결과 활용",
        "Google Shape;531;p27": "남은 문구만 AI 확인",
        "Google Shape;532;p27": "표 제목·행·열 문맥을 함께 제공",
        "Google Shape;537;p27": "담당자 확인 후 확정",
        "Google Shape;538;p27": "현재 원문과 제안 내용을 함께 확인",
        "Google Shape;540;p27": "표현 기준 준비",
        "Google Shape;541;p27": "검수할 때 다음 자료를 모아 출처와 상태를 함께 관리",
        "Google Shape;551;p27": "이전 결과를 다시 사용하는 조건",
        "Google Shape;553;p27": "검수 항목",
        "Google Shape;555;p27": "원문·위치",
        "Google Shape;557;p27": "표 제목·단위",
        "Google Shape;559;p27": "행·열 항목",
        "Google Shape;561;p27": "주변 문맥",
        "Google Shape;563;p27": "사용한 AI",
        "Google Shape;565;p27": "적용 기준",
        "Google Shape;566;p27": "같은 문구라도 위치와 문맥이 다르면 새로 확인합니다.",
        "Google Shape;568;p27": "AI를 사용하지 못하거나 응답이 없으면 ‘재검수 필요’로 남기며, 원문을 자동으로 바꾸지 않습니다.",
    },
    14: {
        "Google Shape;573;p28": "7단계 · 표 정보·특이값 확인",
        "Google Shape;574;p28": "표 정보와 특이값 확인",
        "Google Shape;577;p28": "필수 정보가 빠졌는지, 다른 값과 차이가 큰 숫자가 있는지 알려줍니다.",
        "Google Shape;582;p28": "표 정보 확인",
        "Google Shape;586;p28": "단위 누락 또는 기존 기준과의 차이 확인",
        "Google Shape;590;p28": "기준일 입력 여부 확인",
        "Google Shape;594;p28": "출처 입력 여부 확인",
        "Google Shape;596;p28": "현재 자동 확인 범위: 단위·기준일·출처의 누락과 단위 기준",
        "Google Shape;599;p28": "특이값 확인",
        "Google Shape;600;p28": "같은 숫자 열의 값들을 비교",
        "Google Shape;612;p28": "일반 범위",
        "Google Shape;615;p28": "차이 큰 값",
        "Google Shape;616;p28": "• 합계·소계는 비교 대상에서 제외\v• 값이 5개 이상인 숫자 열에서 차이가 큰 값 제시",
        "Google Shape;618;p28": "특이값은 오류 확정이 아니라 담당자가 확인할 후보",
        "Google Shape;619;p28": "담당자는 출처·단위·업무 맥락을 함께 보고 최종 판단합니다.",
    },
    15: {
        "Google Shape;624;p29": "8단계 · 수정·재검수",
        "Google Shape;625;p29": "결과를 수정하고 다시 검수",
        "Google Shape;628;p29": "문제 위치를 확인해 수정하고, 바뀐 표만 다시 확인합니다.",
        "Google Shape;633;p29": "수정할 수 있는 내용",
        "Google Shape;636;p29": "통계 코드·제목·분류",
        "Google Shape;639;p29": "헤더·행·열 추가·삭제·순서",
        "Google Shape;642;p29": "한글·영문 셀 값·각주",
        "Google Shape;644;p29": "표 정보",
        "Google Shape;645;p29": "단위·기준일·출처·담당 정보",
        "Google Shape;648;p29": "통계별 수정 보호",
        "Google Shape;650;p29": "비밀번호 확인",
        "Google Shape;653;p29": "수정 화면 열기",
        "Google Shape;656;p29": "변경 내용 저장",
        "Google Shape;657;p29": "통계별 비밀번호를 확인한 뒤 수정 화면을 사용할 수 있습니다.",
        "Google Shape;658;p29": "※ 기관 계정·역할별 권한 관리는 향후 보완 예정",
        "Google Shape;659;p29": "수정 후 다시 확인하는 순서",
        "Google Shape;661;p29": "결과 확인",
        "Google Shape;664;p29": "문제 내용 수정",
        "Google Shape;667;p29": "‘재검수 필요’\v표시",
        "Google Shape;670;p29": "재검수\v실행",
        "Google Shape;673;p29": "새 결과로\v갱신",
        "Google Shape;675;p29": "값 변경: 기존 계산 기준으로 다시 확인   |   구조 변경: 검수 기준을 새로 준비   |   동일 표: 이전 결과 활용",
    },
    16: {
        "Google Shape;680;p30": "9단계 · 결과 내려받기",
        "Google Shape;681;p30": "검수 결과 내려받기",
        "Google Shape;684;p30": "검수 목적에 맞춰 표·연보·검수 보고서 형태로 저장합니다.",
        "Google Shape;691;p30": "선택한 표 검토·공유",
        "Google Shape;693;p30": "• 선택한 통계표와 하위 표\v• 한글·영문 헤더와 셀 내용\v• 단위·기준일·출처 등 표 정보",
        "Google Shape;698;p30": "전체 연보 취합·분석",
        "Google Shape;700;p30": "• 전체 또는 선택한 통계표\v• 통계표를 찾기 위한 목차\v• 통계표·하위 표별 개별 시트",
        "Google Shape;705;p30": "오류 확인·결과 보고",
        "Google Shape;707;p30": "• 확인할 문제와 판정 근거\v• 문제 위치가 표시된 표\v• 표 정보 확인 자료",
        "Google Shape;709;p30": "필요한 범위와 검수 상태를 선택해 업무 목적에 맞는 결과물을 만듭니다.",
    },
    17: {
        "Google Shape;714;p31": "기대효과",
        "Google Shape;715;p31": "서비스를 사용하면 달라지는 점",
        "Google Shape;718;p31": "반복 계산은 시스템이 돕고, 담당자는 판단과 최종 확인에 집중합니다.",
        "Google Shape;724;p31": "반복 작업 감소",
        "Google Shape;725;p31": "합계·비율·증감 값을 같은 방식으로 반복 확인",
        "Google Shape;729;p31": "검수 기준 통일",
        "Google Shape;730;p31": "확인한 계산 기준과 용어를 다음 연도에 활용",
        "Google Shape;734;p31": "문제 위치를 쉽게 추적",
        "Google Shape;735;p31": "문제 셀·근거 셀·현재값·다시 계산한 값을 함께 확인",
        "Google Shape;739;p31": "수정 결과를 최신 상태로 관리",
        "Google Shape;740;p31": "수정 전 결과와 구분하고 바뀐 표를 중심으로 재검수",
        "Google Shape;742;p31": "담당자의 역할",
        "Google Shape;743;p31": "오류 후보 확인\v→ 원문 수정\v→ 최종 승인",
        "Google Shape;745;p31": "시스템은 반복 확인과\v우선 검토 대상 제시를 지원합니다.",
        "Google Shape;746;p31": "※ 실제 운영 자료를 쌓은 뒤 시간 절감과 정확도 개선 효과를 측정합니다.",
    },
    18: {
        "Google Shape;751;p32": "향후 개선",
        "Google Shape;752;p32": "운영 전에 더 보완할 점",
        "Google Shape;755;p32": "안정적인 기관 운영을 위해 권한·이력·검수 범위를 단계적으로 강화합니다.",
        "Google Shape;760;p32": "공통 운영 기준을 먼저 마련하고, 담당자 의견을 반영해 단계적으로 확대",
    },
}


SLIDE18_TABLE = [
    ["개선 영역", "현재 제공 범위", "추가로 보완할 내용", "우선"],
    ["권한·보안", "통계별 비밀번호로 수정 보호", "기관 계정 연계, 역할별 권한, 세션 만료와 접속 제한", "높음"],
    ["검수 기준 관리", "결과와 판정 근거 표시", "검수 기준의 담당자 확인·승인, 적용 범위와 변경 이력 관리", "높음"],
    ["확인 범위 확대", "단위·기준일·출처와 숫자 특이값", "날짜 형식·출처 품질·연도별 변화 등 추가 확인", "중간"],
    ["수정·재검수", "수정 후 담당자가 재검수 실행", "변경 전후 이력, 되돌리기, 한 번에 재검수와 알림", "중간"],
    ["운영·결과물", "대규모 처리와 발간 출력 지원", "진행 상태·실패 사유 안내, 공식 발간 형식 지원", "중간"],
]


def first_run(text_frame):
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            return run
    return None


def copy_font(source_run, target_run) -> None:
    if source_run is None:
        return
    for attr in ("name", "size", "bold", "italic", "underline"):
        value = getattr(source_run.font, attr)
        if value is not None:
            setattr(target_run.font, attr, value)
    try:
        if source_run.font.color.type is not None and source_run.font.color.rgb is not None:
            target_run.font.color.rgb = source_run.font.color.rgb
    except (AttributeError, TypeError, ValueError):
        pass


def set_text_frame_text(text_frame, value: str) -> None:
    source_run = first_run(text_frame)
    source_paragraph = text_frame.paragraphs[0]
    alignment = source_paragraph.alignment
    level = source_paragraph.level
    space_before = source_paragraph.space_before
    space_after = source_paragraph.space_after

    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.level = level
    paragraph.space_before = space_before
    paragraph.space_after = space_after
    # Paragraph.text converts vertical tabs to DrawingML soft line breaks.
    # Assigning them directly to a Run escapes them as `_x000B_` text.
    paragraph.text = value
    for run in paragraph.runs:
        copy_font(source_run, run)


def add_cover_subtitle(slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.72), Inches(3.54), Inches(7.3), Inches(0.34))
    box.name = "서비스 한줄 설명"
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = "통계연보 등록부터 오류 확인·수정·결과 활용까지 한 번에"
    run.font.name = FONT_NAME
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("B7CCE6")


def set_page_number(slide, number: int) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.left > Inches(9.0) and shape.top > Inches(5.0):
            set_text_frame_text(shape.text_frame, str(number))
            return


def rewrite(source: Path, output: Path) -> None:
    prs = Presentation(source)
    if len(prs.slides) != 18:
        raise RuntimeError(f"예상한 18쪽이 아니라 {len(prs.slides)}쪽입니다.")

    for slide_number, slide in enumerate(prs.slides, start=1):
        replacements = SLIDE_TEXT.get(slide_number, {})
        for shape in slide.shapes:
            if shape.name in replacements and shape.has_text_frame:
                set_text_frame_text(shape.text_frame, replacements[shape.name])
        set_page_number(slide, slide_number)

    add_cover_subtitle(prs.slides[0])

    table = next(shape.table for shape in prs.slides[17].shapes if shape.has_table)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            set_text_frame_text(cell.text_frame, SLIDE18_TABLE[row_index][column_index])

    for slide in prs.slides:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            frame.clear()
            frame.text = ""

    force_font_family(prs, FONT_NAME)
    prs.core_properties.title = "통계연보 검수 보조 시스템"
    prs.core_properties.subject = "처음 사용하는 담당자를 위한 서비스 소개 및 검수 흐름"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    rewrite(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()

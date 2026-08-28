"""Build a 17-slide, speaker-ready yearbook validation presentation."""

from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_yearbook_validation_presentation import (
    AMBER,
    AMBER_BG,
    BLUE,
    BLUE_2,
    BORDER,
    C,
    CYAN,
    LIGHT,
    MUTED,
    MUTED_2,
    NAVY,
    RED,
    RED_BG,
    SKY_2,
    TEAL,
    TEAL_BG,
    TEXT,
    WHITE,
    add_arrow,
    add_bullets,
    add_circle,
    add_down_arrow,
    add_header,
    add_pill,
    add_rect,
    add_rich_text,
    add_shape,
    add_text,
    clear_slide,
    force_font_family,
)


FONT_NAME = "Noto Sans KR"
FOOTER = "청년인턴 완료보고  |  행정안전부 데이터정보화담당관"

# 제한된 발표 팔레트: 기본 정보는 파란색 계열, 주의·오류만 노랑/빨강.
PRIMARY = "153E70"
PRIMARY_2 = "2F78B7"
PRIMARY_LIGHT = "EEF5FA"
PRIMARY_PALE = "F6F9FC"
WARNING = "C5861A"
WARNING_BG = "FFF4DC"
ERROR = "C94655"
ERROR_BG = "FDECEF"
DIVIDER = "D8E1E8"


def find_logo_blobs(prs: Presentation) -> tuple[bytes, bytes]:
    cover_blob = None
    normal_blob = None
    for shape in prs.slides[0].shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            cover_blob = shape.image.blob
            break
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            blob = shape.image.blob
            if normal_blob is None or len(blob) > len(normal_blob):
                normal_blob = blob
    if cover_blob is None or normal_blob is None:
        raise RuntimeError("행정안전부 로고 이미지를 찾지 못했습니다.")
    return cover_blob, normal_blob


def set_background(slide, color: str = WHITE) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = C(color)


def add_open_heading(slide, x: float, y: float, w: float, title: str, *, accent: str = BLUE, size: float = 11.3) -> None:
    add_rect(slide, x, y, 0.045, 0.35, fill=accent, line=None)
    add_text(slide, title, x + 0.15, y + 0.04, w - 0.15, 0.25, size=size, color=TEXT, bold=True, valign=MSO_ANCHOR.MIDDLE)


def add_image_frame(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    title: str,
    subtitle: str = "",
    accent: str = BLUE,
    name: str,
) -> None:
    frame = add_shape(
        slide,
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill="FAFBFC",
        line="CAD5DF",
        line_width=1.0,
        radius_name=name,
    )
    add_text(slide, "+", x + w / 2 - 0.25, y + h / 2 - 0.43, 0.50, 0.45, size=23, color=accent, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, title, x + 0.25, y + h / 2 + 0.05, w - 0.50, 0.26, size=10.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        add_text(slide, subtitle, x + 0.25, y + h / 2 + 0.37, w - 0.50, 0.24, size=8.3, color=MUTED, align=PP_ALIGN.CENTER)


def add_numbered_row(
    slide,
    x: float,
    y: float,
    w: float,
    number: str,
    title: str,
    detail: str,
    *,
    accent: str = BLUE,
    fill: str = LIGHT,
    title_size: float = 10.1,
    detail_size: float = 8.5,
) -> None:
    add_rect(slide, x, y, w, 0.54, fill=fill, line=None)
    add_rect(slide, x, y, 0.045, 0.54, fill=accent, line=None)
    add_text(slide, number, x + 0.13, y + 0.15, 0.35, 0.20, size=8.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.57, y + 0.08, w - 0.72, 0.20, size=title_size, color=TEXT, bold=True)
    add_text(slide, detail, x + 0.57, y + 0.31, w - 0.72, 0.16, size=detail_size, color=MUTED)


def add_status_pill(slide, text: str, x: float, y: float, *, review: bool) -> None:
    add_pill(
        slide,
        text,
        x,
        y,
        0.92 if review else 0.88,
        0.26,
        fill=AMBER_BG if review else RED_BG,
        color=AMBER if review else RED,
        size=7.8,
        line=None,
    )


def add_small_badge(slide, text: str, x: float, y: float, w: float, *, fill: str = SKY_2, color: str = BLUE) -> None:
    add_pill(slide, text, x, y, w, 0.27, fill=fill, color=color, size=7.8, line=None)


def add_takeaway(slide, text: str, *, y: float = 4.86, fill: str = NAVY, color: str = WHITE) -> None:
    add_rect(slide, 1.00, y, 8.00, 0.26, fill=fill, line=None)
    add_text(slide, text, 1.25, y + 0.05, 7.50, 0.17, size=8.4, color=color, bold=True, align=PP_ALIGN.CENTER)


def add_slide_header(slide, section: str, title: str, subtitle: str, page: int, logo_blob: bytes) -> None:
    set_background(slide)
    add_header(slide, section, subtitle, page, logo_blob, headline=title)


def build_cover(slide, cover_logo_blob: bytes) -> None:
    clear_slide(slide)
    set_background(slide, NAVY)
    add_rect(slide, 0, 0, 10, 5.625, fill=NAVY, line=None)
    slide.shapes.add_picture(BytesIO(cover_logo_blob), Inches(8.40), Inches(0.12), width=Inches(1.45))
    add_text(slide, "통계연보 발간 지원", 0.72, 1.28, 4.0, 0.26, size=9.2, color=CYAN, bold=True)
    add_text(slide, "통계연보 검수\n보조 시스템", 0.72, 1.64, 6.9, 1.45, size=31.5, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "반복 검수는 줄이고, 오류 확인과 수정은 한 흐름으로", 0.74, 3.36, 6.8, 0.32, size=12.3, color="B7CCE6")
    add_text(slide, "검수·수정·재검수·결과 공유를 하나의 플랫폼에서 지원", 0.74, 3.78, 6.8, 0.28, size=9.0, color="8FB2D8")
    add_text(slide, "01", 7.00, 1.40, 2.1, 2.2, size=92, color="153E70", bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, f"{FOOTER}  |  2026. 8.", 0.72, 5.02, 6.7, 0.18, size=7.3, color="B7CCE6")
    add_text(slide, "1", 9.18, 5.04, 0.30, 0.20, size=7.5, color="B7CCE6", align=PP_ALIGN.RIGHT)


def slide_need_for_review(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "통계연보 발간과 검수",
        "정확한 통계 정보를 제공하기 위해 발간 전 검수는 반드시 필요합니다.",
        "통계연보 발간 전, 왜 검수가 필요한가",
        2,
        logo_blob,
    )
    add_image_frame(
        slide,
        0.55,
        1.98,
        5.75,
        2.88,
        title="통계연보 예시 이미지 삽입",
        subtitle="표지 또는 여러 통계표가 보이는 화면 권장",
        name="이미지 삽입_통계연보 예시",
    )
    add_open_heading(slide, 6.62, 2.00, 2.75, "통계연보 발간 과정")
    add_numbered_row(slide, 6.62, 2.50, 2.80, "01", "여러 데이터 취합", "부서·기관 자료를 한 권으로 정리", accent=NAVY)
    add_numbered_row(slide, 6.62, 3.13, 2.80, "02", "발간 전 검수", "정확한 정보 제공을 위한 마지막 확인", accent=BLUE)
    add_numbered_row(slide, 6.62, 3.76, 2.80, "03", "여러 역할이 참여", "담당자·주무관·청년인턴 협업", accent=TEAL)
    add_small_badge(slide, "데이터 취합", 6.75, 4.52, 0.82)
    add_text(slide, "→", 7.60, 4.56, 0.28, 0.18, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_small_badge(slide, "검수", 7.91, 4.52, 0.62)
    add_text(slide, "→", 8.56, 4.56, 0.28, 0.18, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_small_badge(slide, "발간", 8.87, 4.52, 0.62)


def slide_as_is_to_be(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "검수 업무의 어려움",
        "300개 이상의 표를 반복 확인하고 오류를 공유·수정하는 과정은 시간이 많이 듭니다.",
        "반복 검수에서 협업 플랫폼으로",
        3,
        logo_blob,
    )
    add_open_heading(slide, 0.55, 2.00, 4.05, "AS-IS  사람이 반복 확인", accent=NAVY, size=12)
    add_open_heading(slide, 5.40, 2.00, 4.05, "TO-BE  서비스가 반복 업무 지원", accent=BLUE, size=12)
    as_is = [
        ("300개+", "표를 하나씩 반복 확인"),
        ("누락 위험", "계산 오류·오탈자를 놓치기 쉬움"),
        ("복잡한 공유", "오류 전달·수정 반영·재확인이 분산"),
    ]
    to_be = [
        ("자동 계산", "합계·비율을 구성값에서 다시 계산"),
        ("근거 확인", "문제 셀과 판정 근거를 한눈에 확인"),
        ("한 곳에서 협업", "수정·재검수·결과 공유까지 연결"),
    ]
    for i, ((title_a, detail_a), (title_b, detail_b)) in enumerate(zip(as_is, to_be), start=1):
        y = 2.55 + (i - 1) * 0.68
        add_numbered_row(slide, 0.62, y, 3.92, f"0{i}", title_a, detail_a, accent=NAVY)
        add_numbered_row(slide, 5.48, y, 3.92, f"0{i}", title_b, detail_b, accent=BLUE, fill=SKY_2)
    add_arrow(slide, 4.72, 3.19, 0.46, 0.44, color=BLUE)
    add_takeaway(slide, "반복 계산은 서비스가 처리하고, 담당자는 판단·수정·공유에 집중")


def slide_main_page(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "서비스 화면 ①",
        "연보를 등록하고 검수를 실행한 뒤, 오류가 있는 통계와 전체 현황을 확인합니다.",
        "메인 페이지에서 연보와 검수 현황을 한눈에",
        4,
        logo_blob,
    )
    add_image_frame(
        slide,
        0.55,
        1.96,
        6.10,
        3.02,
        title="메인 페이지 이미지 삽입",
        subtitle="연보 목록·검수 실행·상태 요약이 보이는 화면 권장",
        name="이미지 삽입_메인 페이지",
    )
    add_open_heading(slide, 6.92, 1.99, 2.45, "메인 페이지에서 하는 일")
    add_numbered_row(slide, 6.92, 2.50, 2.50, "01", "연보 등록·검수 시작", "대상 연보 선택 후 검수 실행", accent=NAVY)
    add_numbered_row(slide, 6.92, 3.14, 2.50, "02", "오류 통계 자동 분류", "확인이 필요한 통계를 우선 표시", accent=BLUE)
    add_numbered_row(slide, 6.92, 3.78, 2.50, "03", "검수 현황 요약", "정상·확인 필요·오류 의심 집계", accent=TEAL)
    add_small_badge(slide, "연보 선택", 7.02, 4.54, 0.72)
    add_text(slide, "→", 7.78, 4.58, 0.28, 0.18, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_small_badge(slide, "검수 실행", 8.10, 4.54, 0.78)
    add_text(slide, "→", 8.91, 4.58, 0.28, 0.18, size=11, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_small_badge(slide, "결과 요약", 9.19, 4.54, 0.78)


def slide_detail_page(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "서비스 화면 ②",
        "통계 상세 화면에서 문제 위치와 근거를 확인하고 바로 수정할 수 있습니다.",
        "문제 셀을 확인하고 바로 수정",
        5,
        logo_blob,
    )
    add_image_frame(
        slide,
        0.55,
        1.96,
        6.10,
        3.02,
        title="통계 상세 페이지 이미지 삽입",
        subtitle="하이라이트 셀·검수 내역·수정 화면이 보이는 화면 권장",
        name="이미지 삽입_통계 상세 페이지",
    )
    add_open_heading(slide, 6.92, 1.99, 2.45, "상세 화면에서 확인")
    add_numbered_row(slide, 6.92, 2.50, 2.50, "01", "문제 셀 하이라이트", "어떤 위치에 문제가 있는지 표시", accent=RED, fill=RED_BG)
    add_numbered_row(slide, 6.92, 3.14, 2.50, "02", "검수 내역·근거", "현재값·기대값·근거 셀 확인", accent=BLUE)
    add_numbered_row(slide, 6.92, 3.78, 2.50, "03", "내용 수정", "표 정보·셀 값·구조를 화면에서 편집", accent=TEAL, fill=TEAL_BG)
    add_takeaway(slide, "문제 위치 확인 → 근거 검토 → 내용 수정 → 재검수")


def add_flow_box(slide, x: float, y: float, w: float, title: str, detail: str, *, accent: str, fill: str = LIGHT) -> None:
    add_rect(slide, x, y, w, 0.82, fill=fill, line=None)
    add_rect(slide, x, y, 0.045, 0.82, fill=accent, line=None)
    add_text(slide, title, x + 0.18, y + 0.13, w - 0.34, 0.22, size=10.3, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, detail, x + 0.18, y + 0.43, w - 0.34, 0.22, size=8.3, color=MUTED, align=PP_ALIGN.CENTER)


def slide_service_flow(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "서비스 이용 흐름",
        "검수 결과에 문제가 있으면 수정·재검수를 반복하고, 오류가 없으면 결과를 내려받습니다.",
        "연보 등록부터 재검수까지 반복되는 흐름",
        6,
        logo_blob,
    )
    add_flow_box(slide, 0.55, 2.06, 1.45, "1  연보 등록", "파일·연도 선택", accent=NAVY)
    add_arrow(slide, 2.06, 2.28, 0.30, 0.34, color=CYAN)
    add_flow_box(slide, 2.42, 2.06, 1.45, "2  검수 실행", "검수 항목 선택", accent=BLUE, fill=SKY_2)
    add_arrow(slide, 3.93, 2.28, 0.30, 0.34, color=CYAN)
    add_flow_box(slide, 4.29, 2.06, 1.45, "3  결과 확인", "문제 위치·근거", accent=BLUE)
    add_arrow(slide, 5.80, 2.28, 0.30, 0.34, color=CYAN)
    diamond = add_shape(slide, MSO_AUTO_SHAPE_TYPE.DIAMOND, 6.18, 1.98, 1.25, 1.02, fill=WHITE, line=BLUE, line_width=1.2, radius_name="판단_오류 없음")
    add_text(slide, "오류 없음?", 6.38, 2.28, 0.85, 0.26, size=9.4, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "예", 7.43, 2.10, 0.28, 0.18, size=8.0, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 7.60, 2.28, 0.30, 0.34, color=TEAL)
    add_flow_box(slide, 7.96, 2.06, 1.48, "4  연보 완성", "오류 없는 결과", accent=TEAL, fill=TEAL_BG)
    add_down_arrow(slide, 8.54, 3.02, 0.32, 0.34, color=TEAL)
    add_pill(slide, "표·검수 리포트 내려받기", 7.65, 3.46, 2.10, 0.34, fill=TEAL_BG, color=TEAL, size=8.6, line=None)

    add_text(slide, "아니오", 6.64, 3.03, 0.45, 0.18, size=8.0, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 6.64, 3.22, 0.32, 0.34, color=RED)
    add_flow_box(slide, 5.72, 3.66, 1.55, "내용 수정", "문제 셀·표 정보", accent=AMBER, fill=AMBER_BG)
    left_arrow = add_shape(slide, MSO_AUTO_SHAPE_TYPE.LEFT_ARROW, 5.28, 3.88, 0.32, 0.34, fill=CYAN, line=None, radius_name="재검수 방향")
    add_flow_box(slide, 3.62, 3.66, 1.55, "재검수", "변경된 표 다시 확인", accent=BLUE, fill=SKY_2)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.LEFT_ARROW, 3.17, 3.88, 0.32, 0.34, fill=CYAN, line=None, radius_name="검수 반복 방향")
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.UP_ARROW, 2.68, 3.05, 0.38, 1.17, fill=CYAN, line=None, radius_name="검수 실행 복귀")
    add_text(slide, "검수 실행으로 다시 돌아가 반복", 1.12, 4.59, 6.10, 0.22, size=8.8, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


def slide_upload(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "1단계 · 연보 등록",
        "연보 파일을 읽어 통계별 제목·번호·내용을 저장하고 플랫폼 화면으로 재구성합니다.",
        "연보 파일이 플랫폼의 통계 게시물로",
        7,
        logo_blob,
    )
    add_image_frame(slide, 0.55, 2.02, 3.00, 2.28, title="실제 연보의 통계 예시", subtitle="연보 안의 통계표 이미지", name="이미지 삽입_연보 통계 예시")
    add_arrow(slide, 3.74, 2.90, 0.38, 0.42, color=CYAN)
    add_rect(slide, 4.22, 2.42, 1.56, 1.35, fill=SKY_2, line=None)
    add_text(slide, "업로드", 4.50, 2.65, 1.00, 0.28, size=12.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_pill(slide, "HWPX", 4.39, 3.09, 0.55, 0.27, fill=WHITE, color=BLUE, size=8.0, line=None)
    add_pill(slide, "MD", 5.05, 3.09, 0.55, 0.27, fill=WHITE, color=BLUE, size=8.0, line=None)
    add_text(slide, "지원 확장자", 4.49, 3.45, 1.00, 0.18, size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
    add_arrow(slide, 5.90, 2.90, 0.38, 0.42, color=CYAN)
    add_image_frame(slide, 6.40, 2.02, 3.00, 2.28, title="플랫폼에 등록된 통계", subtitle="제목·번호·표 내용이 재구성된 화면", name="이미지 삽입_플랫폼 등록 통계")

    steps = [
        ("01", "파일 읽기", "연보 내용을 통계별로 분리"),
        ("02", "정보 저장", "제목·번호·표·기본 정보를 데이터베이스에 저장"),
        ("03", "게시물 구성", "연보 1건을 독립된 게시물로 생성"),
    ]
    for i, (num, title, detail) in enumerate(steps):
        x = 0.82 + i * 3.02
        add_numbered_row(slide, x, 4.37, 2.72, num, title, detail, accent=(NAVY, BLUE, TEAL)[i], fill=WHITE)
    add_takeaway(slide, "등록된 연보는 독립 게시물로 관리하며 수정·삭제 가능", y=4.98, fill=SKY_2, color=NAVY)


def add_validation_card(slide, x: float, y: float, w: float, title: str, detail: str, *, accent: str, review: bool) -> None:
    add_rect(slide, x, y, w, 0.78, fill=LIGHT, line=None)
    add_rect(slide, x, y, 0.045, 0.78, fill=accent, line=None)
    add_text(slide, title, x + 0.16, y + 0.12, w - 1.17, 0.22, size=10.2, color=TEXT, bold=True)
    add_text(slide, detail, x + 0.16, y + 0.42, w - 0.34, 0.19, size=8.2, color=MUTED)
    add_status_pill(slide, "확인 필요" if review else "오류 의심", x + w - 1.02, y + 0.11, review=review)


def slide_validation_types(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "2단계 · 검수 실행",
        "현재 6가지 검수를 지원하며, 처리 방식과 결과 상태가 항목마다 다릅니다.",
        "6가지 검수, 두 가지 처리 방식",
        8,
        logo_blob,
    )
    add_pill(slide, "현재 지원  6종", 8.08, 1.54, 1.30, 0.30, fill=NAVY, color=WHITE, size=8.4, line=None)
    add_open_heading(slide, 0.55, 2.00, 5.60, "규칙·통계 기반 검수", accent=BLUE, size=12)
    add_text(slide, "표 구조와 계산 관계를 미리 구성한 기준으로 확인", 0.72, 2.34, 4.95, 0.20, size=8.4, color=MUTED)
    add_validation_card(slide, 0.62, 2.68, 2.65, "합계", "계·합계·소계를 다시 합산", accent=BLUE, review=False)
    add_validation_card(slide, 3.40, 2.68, 2.65, "비율", "비율·구성비·증감 값을 재계산", accent=BLUE, review=False)
    add_validation_card(slide, 0.62, 3.59, 2.65, "메타데이터", "단위·기준일·출처 확인", accent=AMBER, review=True)
    add_validation_card(slide, 3.40, 3.59, 2.65, "이상치", "다른 값과 차이가 큰 수치 후보", accent=AMBER, review=True)

    add_open_heading(slide, 6.37, 2.00, 3.05, "후보 생성 + 선택적 LLM", accent=TEAL, size=12)
    add_text(slide, "공식 용어·이전 결과를 먼저 적용하고 남은 후보만 확인", 6.54, 2.34, 2.75, 0.30, size=8.2, color=MUTED)
    add_validation_card(slide, 6.44, 2.82, 2.96, "오탈자", "철자·문자 깨짐·표기 불일치", accent=RED, review=False)
    add_validation_card(slide, 6.44, 3.73, 2.96, "번역", "한글·영문 의미 대응 확인", accent=TEAL, review=True)
    add_takeaway(slide, "오류 의심  합계·비율·오탈자   |   확인 필요  메타데이터·이상치·번역", y=4.78)


def add_criterion_card(
    slide,
    x: float,
    title: str,
    bullets: tuple[str, str],
    tags: tuple[str, ...],
    *,
    accent: str,
) -> None:
    add_open_heading(slide, x, 2.03, 2.02, title, accent=accent, size=10.5)
    add_text(slide, f"• {bullets[0]}\n• {bullets[1]}", x + 0.12, 2.49, 1.85, 0.55, size=8.1, color=TEXT)
    tag_x = x + 0.10
    for tag in tags:
        tag_w = 0.63 if len(tag) <= 3 else 0.78
        add_small_badge(slide, tag, tag_x, 3.12, tag_w, fill=SKY_2 if accent != TEAL else TEAL_BG, color=accent)
        tag_x += tag_w + 0.08


def slide_rule_generation(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "검수 기준 생성 방식",
        "표 구조와 통계 특성을 단계적으로 분석해 각 표에 적용할 검수를 결정합니다.",
        "어떤 표에 어떤 검수를 적용할까",
        9,
        logo_blob,
    )
    add_criterion_card(slide, 0.55, "① 공통 기준", ("모든 표의 필수 정보 확인", "숫자 열의 특이값 후보 생성"), ("메타", "이상치"), accent=NAVY)
    add_criterion_card(slide, 2.82, "② 표 유형별 기준", ("지역별·연도별 등 구조 파악", "전국계와 지역 값 관계 적용"), ("합계",), accent=BLUE)
    add_criterion_card(slide, 5.09, "③ 표별 자동 분석", ("계·합계·소계 항목 탐색", "분자·분모·증감 관계 탐색"), ("합계", "비율"), accent=BLUE)
    add_criterion_card(slide, 7.36, "④ 통계별 예외 기준", ("담당자가 확인한 규칙 반영", "고정 셀·분할표 예외 처리"), ("합계", "비율"), accent=TEAL)

    for x in (2.41, 4.68, 6.95):
        add_arrow(slide, x, 2.72, 0.24, 0.30, color=CYAN)
    add_rect(slide, 0.72, 3.62, 6.08, 0.54, fill=SKY_2, line=None)
    add_text(slide, "표 구조 기준", 0.93, 3.79, 1.04, 0.19, size=9.3, color=BLUE, bold=True)
    add_text(slide, "공통 + 유형 + 표 자동 분석 + 예외 기준", 2.05, 3.77, 3.44, 0.22, size=9.2, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 5.55, 3.77, 0.35, 0.22, size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "합계 · 비율 · 메타데이터 · 이상치", 5.93, 3.77, 2.55, 0.22, size=8.8, color=NAVY, bold=True)

    add_rect(slide, 0.72, 4.28, 8.58, 0.54, fill=TEAL_BG, line=None)
    add_text(slide, "언어 표현 기준", 0.93, 4.45, 1.04, 0.19, size=9.3, color=TEAL, bold=True)
    add_text(slide, "공식 용어 + 반복 표현 + 이전 검수 결과 + 남은 후보", 2.05, 4.43, 4.25, 0.22, size=8.9, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 6.38, 4.43, 0.35, 0.22, size=13, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "오탈자 · 번역", 6.84, 4.43, 1.70, 0.22, size=9.1, color=TEAL, bold=True)


def add_calculation_slide(
    slide,
    logo_blob: bytes,
    *,
    page: int,
    section: str,
    title: str,
    subtitle: str,
    items: tuple[tuple[str, str, str], ...],
    formula: str,
    frame_title: str,
    frame_subtitle: str,
    accent: str,
) -> None:
    clear_slide(slide)
    add_slide_header(slide, section, subtitle, title, page, logo_blob)
    add_open_heading(slide, 0.55, 2.00, 3.00, "검수 대상과 계산 방식", accent=accent, size=11.8)
    for idx, (item_title, detail, example) in enumerate(items, start=1):
        y = 2.46 + (idx - 1) * 0.71
        add_numbered_row(slide, 0.62, y, 3.10, f"0{idx}", item_title, detail, accent=accent, fill=LIGHT)
        add_pill(slide, example, 2.05, y + 0.11, 1.48, 0.26, fill=WHITE, color=accent, size=7.3, line=None)
    add_rect(slide, 0.62, 4.65, 3.10, 0.37, fill=SKY_2 if accent == BLUE else TEAL_BG, line=None)
    add_text(slide, formula, 0.84, 4.75, 2.67, 0.18, size=8.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_image_frame(slide, 4.02, 2.00, 5.40, 2.66, title=frame_title, subtitle=frame_subtitle, accent=accent, name=f"이미지 삽입_{frame_title}")
    add_rect(slide, 4.02, 4.79, 5.40, 0.29, fill=SKY_2 if accent == BLUE else TEAL_BG, line=None)
    add_text(slide, "확인 화면  오류 위치 · 현재값 · 재계산값 · 판정 근거", 4.28, 4.86, 4.88, 0.18, size=8.5, color=accent, bold=True, align=PP_ALIGN.CENTER)


def slide_sum(slide, logo_blob: bytes) -> None:
    add_calculation_slide(
        slide,
        logo_blob,
        page=10,
        section="합계 검수",
        title="구성 항목을 다시 더해 합계를 확인",
        subtitle="계·합계·소계와 표에 표시된 산식을 구성값에서 다시 계산합니다.",
        items=(
            ("지역별 합계", "세부 지역 값을 더해 전국계 확인", "전국 = 시·도 합"),
            ("행·열 합계", "같은 행의 열 또는 세부 행 합산", "계 = 남 + 여"),
            ("명시 산식·분할표", "a+b·a-b와 연결된 표 관계 계산", "c = a + b"),
        ),
        formula="구성값 합산  →  현재 합계와 비교  →  오류 의심",
        frame_title="합계 검수 예시 이미지 삽입",
        frame_subtitle="지역별 합계 또는 산식 오류가 보이는 화면 권장",
        accent=BLUE,
    )


def slide_ratio(slide, logo_blob: bytes) -> None:
    add_calculation_slide(
        slide,
        logo_blob,
        page=11,
        section="비율·증감 검수",
        title="분자·분모와 전년값으로 비율을 재계산",
        subtitle="비율·구성비·건당 값과 증감액·증감률을 다시 계산합니다.",
        items=(
            ("비율·구성비", "항목 값과 전체 값을 분자·분모로 연결", "지역 ÷ 전체 × 100"),
            ("건당·인당 값", "건수·인원을 기준으로 원단위 계산", "건수 ÷ 인원"),
            ("증감액·증감률", "전년과 당해 값을 비교해 변화 계산", "(당해-전년)÷전년"),
        ),
        formula="총계는 비율 합산 X  →  합계 분자 ÷ 합계 분모",
        frame_title="비율·증감 검수 예시 이미지 삽입",
        frame_subtitle="지역별 비율 또는 증감 산식 오류가 보이는 화면 권장",
        accent=TEAL,
    )


def slide_language(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "오탈자·번역 검수",
        "모든 텍스트를 바로 LLM에 보내지 않고, 후보를 줄인 뒤 남은 항목만 검토합니다.",
        "공식 용어와 반복 표현을 먼저, LLM은 마지막에",
        12,
        logo_blob,
    )
    add_rect(slide, 0.72, 1.98, 8.58, 0.46, fill=SKY_2, line=None)
    add_text(slide, "후보 수집", 0.92, 2.11, 1.02, 0.20, size=10.0, color=BLUE, bold=True)
    add_text(slide, "표 제목 · 헤더 · 한글/영문 셀에서 검수할 문구를 수집", 2.06, 2.10, 6.85, 0.22, size=9.1, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 4.82, 2.51, 0.32, 0.30, color=CYAN)

    filters = [
        (0.72, "공식 용어·번역 DB", "기관명  행정안전부\n지역명  세종특별자치시·연기면", BLUE),
        (3.65, "반복 표현 DB", "계 · 합계 · 소계\n연보 전체에서 자주 쓰는 표현", NAVY),
        (6.58, "이전 검수 결과", "같은 위치·문맥에서\n담당자가 확인한 표현", TEAL),
    ]
    for x, title, detail, accent in filters:
        add_rect(slide, x, 2.84, 2.50, 1.02, fill=LIGHT, line=None)
        add_rect(slide, x, 2.84, 0.045, 1.02, fill=accent, line=None)
        add_text(slide, title, x + 0.17, 2.98, 2.16, 0.22, size=10.0, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.18, 3.32, 2.14, 0.36, size=8.1, color=MUTED, align=PP_ALIGN.CENTER)

    for x in (1.82, 4.75, 7.68):
        add_down_arrow(slide, x, 3.97, 0.28, 0.28, color=CYAN)
    add_rect(slide, 2.72, 4.28, 4.56, 0.46, fill=TEAL_BG, line=None)
    add_text(slide, "남은 후보만 LLM 문맥 검토", 3.02, 4.40, 3.96, 0.22, size=11.0, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_small_badge(slide, "토큰 사용량 감소", 0.92, 4.51, 1.23, fill=AMBER_BG, color=AMBER)
    add_small_badge(slide, "공식 표현 일관성", 7.73, 4.51, 1.23, fill=TEAL_BG, color=TEAL)
    add_takeaway(slide, "담당자 확인 전에는 원문을 자동으로 바꾸지 않음", y=4.89)


def slide_metadata_outlier(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "메타데이터·이상치 검수",
        "필수 정보의 누락과 통계적으로 차이가 큰 값을 ‘확인 필요’로 제시합니다.",
        "빠진 표 정보와 눈에 띄는 수치를 확인",
        13,
        logo_blob,
    )
    add_open_heading(slide, 0.55, 2.00, 4.10, "메타데이터 검수", accent=BLUE, size=12)
    meta_items = [
        ("단위", "누락 여부 + 기존 기준 단위와 일치"),
        ("기준일", "입력 여부 확인"),
        ("출처", "입력 여부 확인"),
    ]
    for idx, (title, detail) in enumerate(meta_items, start=1):
        y = 2.52 + (idx - 1) * 0.61
        add_numbered_row(slide, 0.62, y, 3.96, f"0{idx}", title, detail, accent=BLUE)
    add_status_pill(slide, "확인 필요", 3.50, 4.47, review=True)
    add_text(slide, "자동 확인 범위  단위 · 기준일 · 출처", 0.84, 4.52, 2.40, 0.18, size=8.3, color=BLUE, bold=True)

    add_open_heading(slide, 5.05, 2.00, 4.35, "이상치 검수", accent=TEAL, size=12)
    add_text(slide, "같은 숫자 열에서 값이 크게 벗어난 후보", 5.25, 2.43, 3.78, 0.22, size=9.2, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    values = [(5.58, "9", TEAL), (6.20, "10", TEAL), (6.82, "10", TEAL), (7.44, "11", TEAL), (8.55, "43", RED)]
    add_rect(slide, 5.44, 3.23, 3.58, 0.02, fill=BORDER, line=None)
    for x, value, color in values:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 3.10 if value != "43" else 2.84, 0.20, 0.20, fill=color, line=None, radius_name=f"이상치 값_{value}")
        add_text(slide, value, x - 0.08, 3.41, 0.36, 0.18, size=8.0, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_pill(slide, "확인 필요", 8.18, 2.55, 0.92, 0.28, fill=RED_BG, color=RED, size=7.8, line=None)
    add_text(slide, "• 값이 5개 이상인 숫자 열\n• 합계·소계 제외\n• 표당 차이가 큰 값 최대 1건 제시", 5.35, 3.75, 3.55, 0.60, size=8.2, color=TEXT)
    add_rect(slide, 5.31, 4.48, 3.78, 0.30, fill=TEAL_BG, line=None)
    add_text(slide, "오류 확정이 아니라 담당자가 확인할 후보", 5.55, 4.56, 3.30, 0.18, size=8.4, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "출처·단위·업무 맥락을 함께 보고 담당자가 최종 판단", y=4.91)


def slide_edit_recheck(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "수정·재검수",
        "통계별 수정 권한을 확인한 뒤 내용을 고치고, 변경 범위에 맞춰 다시 검수합니다.",
        "수정은 보호하고, 바뀐 표만 다시 확인",
        14,
        logo_blob,
    )
    add_open_heading(slide, 0.55, 2.00, 4.12, "수정 보호", accent=NAVY, size=12)
    protection = [
        ("비밀번호 확인", "통계별 수정 권한 확인"),
        ("편집 화면", "표 정보·셀 값·행/열 수정"),
        ("수정 저장", "변경 내용 반영"),
    ]
    for idx, (title, detail) in enumerate(protection):
        x = 0.62 + idx * 1.37
        add_rect(slide, x, 2.54, 1.13, 0.84, fill=LIGHT if idx < 2 else TEAL_BG, line=None)
        add_text(slide, title, x + 0.10, 2.68, 0.93, 0.22, size=9.1, color=TEAL if idx == 2 else NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.10, 2.99, 0.93, 0.20, size=7.4, color=MUTED, align=PP_ALIGN.CENTER)
        if idx < 2:
            add_arrow(slide, x + 1.18, 2.78, 0.18, 0.28, color=CYAN)
    add_text(slide, "저장 즉시 검수 자동 실행 X  ·  담당자가 재검수 실행", 0.83, 3.57, 3.50, 0.22, size=8.4, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

    add_open_heading(slide, 5.03, 2.00, 4.35, "재검수 원칙", accent=BLUE, size=12)
    principles = [
        ("값 변경", "기존 계산 기준으로 현재 값 재검수"),
        ("구조 변경", "새로운 표 구조에 맞춰 기준 재생성"),
        ("변경 없음", "이전 검수 결과 재사용"),
    ]
    for idx, (title, detail) in enumerate(principles, start=1):
        add_numbered_row(slide, 5.12, 2.49 + (idx - 1) * 0.60, 4.08, f"0{idx}", title, detail, accent=BLUE)

    add_open_heading(slide, 0.55, 4.22, 8.80, "수정 후 재검수 흐름", accent=TEAL, size=10.5)
    flow = ["수정 저장", "재검수 필요", "재검수 실행", "새 결과 갱신"]
    for idx, title in enumerate(flow):
        x = 0.72 + idx * 2.20
        fill = TEAL_BG if idx == 3 else (AMBER_BG if idx == 1 else LIGHT)
        color = TEAL if idx == 3 else (AMBER if idx == 1 else NAVY)
        add_rect(slide, x, 4.63, 1.62, 0.40, fill=fill, line=None)
        add_text(slide, title, x + 0.12, 4.74, 1.38, 0.19, size=8.8, color=color, bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_arrow(slide, x + 1.72, 4.69, 0.28, 0.28, color=CYAN)


def add_output_panel(slide, x: float, title: str, subtitle: str, bullets: tuple[str, ...], *, accent: str, number: str) -> None:
    add_open_heading(slide, x, 2.03, 4.25, f"{number}  {title}", accent=accent, size=12)
    add_text(slide, subtitle, x + 0.18, 2.48, 3.90, 0.24, size=10.0, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, x + 0.18, 2.86, 3.86, 0.02, fill=BORDER, line=None)
    for idx, bullet in enumerate(bullets):
        y = 3.16 + idx * 0.47
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.24, y + 0.02, 0.12, 0.12, fill=accent, line=None, radius_name=f"산출물 항목_{title}_{idx}")
        add_text(slide, bullet, x + 0.48, y, 3.33, 0.20, size=9.0, color=TEXT)


def slide_download(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "검수 결과 활용",
        "필요한 통계표와 검수 결과를 업무 목적에 맞춰 내려받습니다.",
        "표와 검수 리포트로 결과 공유",
        15,
        logo_blob,
    )
    add_output_panel(
        slide,
        0.65,
        "표",
        "선택한 통계표를 검토·공유",
        ("선택한 통계표와 하위 표", "한글·영문 헤더와 셀 내용", "단위·기준일·출처 등 표 정보", "담당자 수정·검토용 자료"),
        accent=BLUE,
        number="01",
    )
    add_output_panel(
        slide,
        5.10,
        "검수 리포트",
        "오류 확인 내역과 판정 근거 보고",
        ("오류 의심·확인 필요 항목", "문제 셀과 근거 셀 위치", "현재값·재계산값·판정 사유", "표 정보 확인 결과"),
        accent=TEAL,
        number="02",
    )
    add_takeaway(slide, "필요한 범위와 검수 상태를 선택해 담당자 공유·수정·결과 보고에 활용")


def add_stakeholder(slide, x: float, title: str, subtitle: str, bullets: tuple[str, ...], *, accent: str, icon: str) -> None:
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 1.08, 2.02, 0.48, 0.48, fill=accent, line=None, radius_name=f"관점 아이콘_{title}")
    add_text(slide, icon, x + 1.08, 2.05, 0.48, 0.34, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, title, x, 2.63, 2.64, 0.24, size=11.2, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, subtitle, x, 2.94, 2.64, 0.20, size=8.4, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, x + 0.18, 3.26, 2.28, 0.02, fill=BORDER, line=None)
    for idx, bullet in enumerate(bullets):
        add_text(slide, f"• {bullet}", x + 0.16, 3.51 + idx * 0.40, 2.33, 0.20, size=8.7, color=TEXT)


def slide_effects(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "기대효과",
        "발간 담당자·통계 담당자·검수 진행자가 같은 화면과 기준으로 협업합니다.",
        "역할별 부담은 줄이고, 검수 효율은 높이고",
        16,
        logo_blob,
    )
    add_stakeholder(slide, 0.65, "통계연보 담당자", "전체 진행 상황 관리", ("연보별 검수 상태 한눈에 확인", "오류 통계와 수정 현황 공유", "발간 전 최종 결과 정리"), accent=NAVY, icon="연")
    add_stakeholder(slide, 3.68, "각 통계 담당자", "내 통계의 문제를 바로 수정", ("문제 셀과 근거를 빠르게 확인", "화면에서 직접 수정·저장", "변경된 표 중심으로 재검수"), accent=TEAL, icon="통")
    add_stakeholder(slide, 6.72, "검수 진행자", "반복 계산 대신 판단에 집중", ("합계·비율 자동 재계산", "확인 우선순위와 근거 제시", "검수 리포트로 결과 공유"), accent=BLUE, icon="검")
    add_rect(slide, 0.88, 4.76, 8.24, 0.34, fill=NAVY, line=None)
    add_text(slide, "검수 시간 단축  ·  오류 누락 위험 감소  ·  공유·수정·재검수 흐름 단순화", 1.18, 4.85, 7.64, 0.19, size=9.3, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_future_panel(slide, x: float, title: str, bullets: tuple[str, ...], *, accent: str, number: str) -> None:
    add_open_heading(slide, x, 2.05, 2.72, f"{number}  {title}", accent=accent, size=11.0)
    add_rect(slide, x + 0.16, 2.58, 2.40, 1.62, fill=LIGHT, line=None)
    for idx, bullet in enumerate(bullets):
        add_text(slide, f"• {bullet}", x + 0.34, 2.88 + idx * 0.42, 2.04, 0.22, size=8.8, color=TEXT)


def slide_future(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "향후 보완점",
        "기관 운영을 위해 권한·이력·AI 환경을 단계적으로 보완해야 합니다.",
        "실제 운영을 위한 다음 과제",
        17,
        logo_blob,
    )
    add_future_panel(slide, 0.62, "권한·보안", ("기관 계정과 통합 로그인", "담당자별 역할·수정 권한", "세션 만료·접속 제한"), accent=NAVY, number="01")
    add_future_panel(slide, 3.64, "수정 이력 기록", ("셀 단위 변경 전·후 기록", "수정자·시간·사유 확인", "되돌리기·변경 알림"), accent=BLUE, number="02")
    add_future_panel(slide, 6.66, "내부망·내부 LLM", ("내부망 환경에 서비스 도입", "외부 LLM을 내부 모델로 전환", "민감 정보와 사용 이력 관리"), accent=TEAL, number="03")
    add_rect(slide, 1.10, 4.48, 7.80, 0.40, fill=SKY_2, line=None)
    add_text(slide, "1  권한 체계 정비", 1.42, 4.59, 1.55, 0.20, size=8.8, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 3.04, 4.58, 0.28, 0.20, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "2  변경 이력 관리", 3.38, 4.59, 1.65, 0.20, size=8.8, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 5.11, 4.58, 0.28, 0.20, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "3  내부망·내부 모델 전환", 5.48, 4.59, 2.28, 0.20, size=8.8, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "공통 운영 기준부터 마련하고, 담당자 의견을 반영해 단계적으로 확대", y=4.96)


# ---------------------------------------------------------------------------
# Final-report visual system
# ---------------------------------------------------------------------------


def add_slide_header(slide, section: str, title: str, subtitle: str, page: int, logo_blob: bytes) -> None:
    """Match the reference deck's title, rule, subtitle, footer and page positions."""
    set_background(slide, WHITE)
    add_text(slide, section, 0.48, 0.24, 3.91, 0.19, size=9.0, color=PRIMARY_2, bold=True, name="상단 분류")
    add_text(slide, title, 0.48, 0.51, 7.70, 0.41, size=20.0, color=TEXT, bold=True, valign=MSO_ANCHOR.MIDDLE, name="슬라이드 제목")
    add_rect(slide, 0.48, 0.98, 9.03, 0.012, fill=DIVIDER, line=None, name="상단 구분선")
    add_text(slide, subtitle, 0.48, 1.12, 8.82, 0.30, size=11.0, color=MUTED, name="슬라이드 부제목")
    slide.shapes.add_picture(BytesIO(logo_blob), Inches(8.30), Inches(0.11), width=Inches(1.26))
    add_text(slide, FOOTER, 0.48, 5.33, 3.55, 0.17, size=9.0, color=MUTED_2, name="하단 문구")
    add_text(slide, str(page), 9.23, 5.34, 0.28, 0.16, size=8.0, color=MUTED, align=PP_ALIGN.RIGHT, name="페이지 번호")


def add_section_title(slide, x: float, y: float, w: float, title: str, *, number: str | None = None, color: str = PRIMARY) -> None:
    label = f"{number}  {title}" if number else title
    add_text(slide, label, x, y, w, 0.25, size=12.0, color=TEXT, bold=True)
    add_rect(slide, x, y + 0.34, w, 0.016, fill=DIVIDER, line=None)
    add_rect(slide, x, y + 0.34, min(0.72, w), 0.025, fill=color, line=None)


def add_image_frame(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    title: str,
    subtitle: str = "",
    accent: str = PRIMARY_2,
    name: str,
) -> None:
    add_rect(slide, x, y, w, h, fill="FBFCFD", line="C9D5DF", line_width=0.9, name=name)
    add_rect(slide, x, y, 0.055, h, fill=accent, line=None)
    add_text(slide, "+", x + w / 2 - 0.23, y + h / 2 - 0.42, 0.46, 0.42, size=22, color=accent, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, title, x + 0.25, y + h / 2 + 0.03, w - 0.50, 0.26, size=10.5, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        add_text(slide, subtitle, x + 0.25, y + h / 2 + 0.35, w - 0.50, 0.23, size=8.4, color=MUTED, align=PP_ALIGN.CENTER)


def add_open_item(
    slide,
    x: float,
    y: float,
    w: float,
    number: str,
    title: str,
    detail: str,
    *,
    color: str = PRIMARY_2,
    formula: str | None = None,
    formula_w: float = 1.35,
) -> None:
    add_text(slide, number, x, y + 0.05, 0.38, 0.20, size=9.2, color=color, bold=True, align=PP_ALIGN.CENTER)
    body_w = w - 0.53 - (formula_w + 0.12 if formula else 0)
    add_text(slide, title, x + 0.51, y, body_w, 0.22, size=10.6, color=TEXT, bold=True)
    add_text(slide, detail, x + 0.51, y + 0.27, body_w, 0.20, size=8.6, color=MUTED)
    if formula:
        add_text(slide, formula, x + w - formula_w, y + 0.08, formula_w, 0.26, size=8.0, color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, x, y + 0.57, w, 0.012, fill=DIVIDER, line=None)


def add_status_badge(slide, text: str, x: float, y: float, *, warning: bool) -> None:
    add_pill(
        slide,
        text,
        x,
        y,
        0.94 if warning else 0.90,
        0.26,
        fill=WARNING_BG if warning else ERROR_BG,
        color=WARNING if warning else ERROR,
        size=7.8,
        line=None,
    )


def build_cover(slide, cover_logo_blob: bytes) -> None:
    clear_slide(slide)
    set_background(slide, NAVY)
    add_rect(slide, 0, 0, 10, 5.625, fill=NAVY, line=None)
    slide.shapes.add_picture(BytesIO(cover_logo_blob), Inches(8.37), Inches(0.12), width=Inches(1.45))
    add_rect(slide, 0.72, 1.22, 0.04, 0.28, fill=PRIMARY_2, line=None)
    add_text(slide, "통계연보 발간 지원  |  26-1 청년인턴 완료보고", 0.86, 1.24, 5.6, 0.24, size=9.3, color="87B7DD", bold=True)
    add_text(slide, "통계연보 검수 보조 시스템", 0.72, 1.70, 8.00, 0.68, size=30.0, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "반복 검수는 줄이고, 오류 확인과 수정은 한 흐름으로", 0.74, 2.68, 7.35, 0.32, size=12.2, color="B7CCE6")
    add_rect(slide, 0.72, 3.24, 6.80, 0.015, fill="4C75A0", line=None)
    add_text(slide, "검수·수정·재검수·결과 공유를 하나의 플랫폼에서 지원", 0.74, 3.48, 7.10, 0.26, size=9.5, color="8FB2D8")
    add_text(slide, f"{FOOTER}  |  2026. 8.", 0.72, 5.02, 6.7, 0.18, size=7.5, color="B7CCE6")
    add_text(slide, "1", 9.23, 5.04, 0.28, 0.16, size=8.0, color="B7CCE6", align=PP_ALIGN.RIGHT)


def slide_need_for_review(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "통계연보 발간과 검수",
        "통계연보 발간 전 검수",
        "정확한 통계 정보를 제공하기 위해 발간 전 검수는 반드시 필요합니다.",
        2,
        logo_blob,
    )
    add_image_frame(slide, 0.48, 1.55, 5.62, 3.42, title="통계연보 예시 이미지 삽입", subtitle="표지 또는 여러 통계표가 보이는 화면 권장", name="이미지 삽입_통계연보 예시")
    add_section_title(slide, 6.52, 1.58, 2.95, "통계연보 발간 과정")
    steps = [
        ("01", "여러 데이터 취합", "부서·기관 자료를 한 권으로 정리"),
        ("02", "발간 전 검수", "정확한 정보 제공을 위한 마지막 확인"),
        ("03", "여러 역할이 참여", "담당자·주무관·청년인턴 협업"),
    ]
    add_rect(slide, 6.72, 2.12, 0.018, 2.12, fill="BCD2E3", line=None)
    for idx, (num, title, detail) in enumerate(steps):
        y = 2.02 + idx * 0.78
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, 6.58, y, 0.30, 0.30, fill=PRIMARY if idx != 1 else PRIMARY_2, line=None, radius_name=f"발간 과정_{num}")
        add_text(slide, num[-1], 6.58, y + 0.04, 0.30, 0.19, size=8.3, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, 7.06, y - 0.01, 2.22, 0.22, size=10.4, color=TEXT, bold=True)
        add_text(slide, detail, 7.06, y + 0.27, 2.25, 0.19, size=8.4, color=MUTED)
    add_rect(slide, 6.52, 4.58, 2.95, 0.34, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "데이터 취합  →  검수  →  발간", 6.72, 4.67, 2.55, 0.18, size=8.6, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_as_is_to_be(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "검수 업무의 어려움",
        "반복 검수의 한계와 개선 방향",
        "300개 이상의 표를 반복 확인하고 오류를 공유·수정하는 과정은 시간이 많이 듭니다.",
        3,
        logo_blob,
    )
    add_rect(slide, 5.12, 1.46, 4.40, 3.55, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "AS-IS", 0.52, 1.58, 0.78, 0.22, size=10.0, color=ERROR, bold=True)
    add_text(slide, "사람이 반복 확인", 1.35, 1.55, 3.25, 0.26, size=12.0, color=TEXT, bold=True)
    add_text(slide, "TO-BE", 5.42, 1.58, 0.82, 0.22, size=10.0, color=PRIMARY_2, bold=True)
    add_text(slide, "서비스가 반복 업무 지원", 6.30, 1.55, 2.95, 0.26, size=12.0, color=TEXT, bold=True)
    as_is = [("300개+", "표를 하나씩 반복 확인"), ("누락 위험", "계산 오류·오탈자를 놓치기 쉬움"), ("복잡한 공유", "오류 전달·수정 반영·재확인이 분산")]
    to_be = [("자동 계산", "합계·비율을 구성값에서 다시 계산"), ("근거 확인", "문제 셀과 판정 근거를 한눈에 확인"), ("한 곳에서 협업", "수정·재검수·결과 공유까지 연결")]
    for i, ((ta, da), (tb, db)) in enumerate(zip(as_is, to_be), start=1):
        y = 2.05 + (i - 1) * 0.82
        add_open_item(slide, 0.52, y, 4.10, f"0{i}", ta, da, color=ERROR)
        add_open_item(slide, 5.42, y, 3.78, f"0{i}", tb, db, color=PRIMARY_2)
    add_arrow(slide, 4.68, 3.03, 0.30, 0.34, color=PRIMARY_2)
    add_text(slide, "반복 계산은 서비스가 처리하고, 담당자는 판단·수정·공유에 집중", 2.04, 4.78, 5.92, 0.22, size=9.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_main_page(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "서비스 화면 ①",
        "연보·검수 현황 한눈에",
        "연보를 등록하고 검수를 실행한 뒤, 오류가 있는 통계와 전체 현황을 확인합니다.",
        4,
        logo_blob,
    )
    add_image_frame(slide, 0.48, 1.53, 6.45, 3.48, title="메인 페이지 이미지 삽입", subtitle="연보 목록·검수 실행·상태 요약이 보이는 화면 권장", name="이미지 삽입_메인 페이지")
    add_section_title(slide, 7.24, 1.56, 2.25, "메인 페이지에서 하는 일")
    items = [("01", "연보 등록·검수 시작", "대상 연보 선택 후 검수 실행"), ("02", "오류 통계 자동 분류", "확인이 필요한 통계를 우선 표시"), ("03", "검수 현황 요약", "정상·확인 필요·오류 의심 집계")]
    for idx, (num, title, detail) in enumerate(items):
        add_open_item(slide, 7.24, 2.10 + idx * 0.78, 2.25, num, title, detail, color=PRIMARY_2)
    add_text(slide, "연보 선택  →  검수 실행  →  결과 요약", 7.20, 4.72, 2.32, 0.22, size=8.3, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_detail_page(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "서비스 화면 ②",
        "문제 위치 확인부터 수정까지",
        "통계 상세 화면에서 문제 위치와 근거를 확인하고 바로 수정할 수 있습니다.",
        5,
        logo_blob,
    )
    add_image_frame(slide, 0.48, 1.53, 6.45, 3.48, title="통계 상세 페이지 이미지 삽입", subtitle="하이라이트 셀·검수 내역·수정 화면이 보이는 화면 권장", name="이미지 삽입_통계 상세 페이지")
    add_section_title(slide, 7.24, 1.56, 2.25, "상세 화면에서 확인")
    items = [("01", "문제 셀 하이라이트", "어떤 위치에 문제가 있는지 표시", ERROR), ("02", "검수 내역·근거", "현재값·기대값·근거 셀 확인", PRIMARY_2), ("03", "내용 수정", "표 정보·셀 값·구조를 화면에서 편집", PRIMARY_2)]
    for idx, (num, title, detail, color) in enumerate(items):
        add_open_item(slide, 7.24, 2.10 + idx * 0.78, 2.25, num, title, detail, color=color)
    add_text(slide, "문제 위치 확인  →  근거 검토\n→  내용 수정  →  재검수", 6.92, 4.65, 2.60, 0.42, size=7.8, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_service_flow(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "서비스 이용 흐름",
        "연보 등록부터 재검수까지",
        "검수 결과에 문제가 있으면 수정·재검수를 반복하고, 오류가 없으면 결과를 내려받습니다.",
        6,
        logo_blob,
    )
    nodes = [(1.06, "01", "연보 등록", "파일·연도 선택"), (3.30, "02", "검수 실행", "검수 항목 선택"), (5.54, "03", "결과 확인", "문제 위치·근거"), (7.78, "04", "연보 완성", "오류 없는 결과\n표·검수 리포트 내려받기")]
    add_rect(slide, 1.28, 2.13, 6.72, 0.025, fill="B9D0E1", line=None)
    for idx, (x, num, title, detail) in enumerate(nodes):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 1.90, 0.46, 0.46, fill=PRIMARY if idx != 2 else PRIMARY_2, line=None, radius_name=f"이용 흐름_{num}")
        add_text(slide, num, x, 2.02, 0.46, 0.18, size=8.2, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x - 0.43, 2.57, 1.32, 0.24, size=10.7, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x - 0.56, 2.88, 1.58, 0.38, size=8.3, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "오류 없음?  →  예", 6.70, 1.82, 1.20, 0.20, size=8.0, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 5.66, 3.22, 0.24, 0.34, color=ERROR)
    add_text(slide, "아니오  ·  오류 있음", 5.05, 3.55, 1.44, 0.20, size=8.4, color=ERROR, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 2.44, 3.87, 4.05, 0.82, fill=WARNING_BG, line=None)
    add_text(slide, "내용 수정", 5.08, 4.08, 1.06, 0.23, size=10.2, color=WARNING, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "문제 셀·표 정보", 5.03, 4.36, 1.16, 0.17, size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.LEFT_ARROW, 4.42, 4.08, 0.40, 0.26, fill=WARNING, line=None, radius_name="수정에서 재검수")
    add_text(slide, "재검수", 3.18, 4.08, 1.02, 0.23, size=10.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "변경된 표 다시 확인", 3.06, 4.36, 1.26, 0.17, size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.UP_ARROW, 3.49, 3.10, 0.36, 0.82, fill=PRIMARY_2, line=None, radius_name="검수 실행으로 복귀")
    add_text(slide, "검수 실행으로 다시 돌아가 반복", 2.62, 4.78, 3.72, 0.18, size=8.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_upload(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "1단계 · 연보 등록",
        "연보 등록과 통계표 재구성",
        "연보 파일을 읽어 통계별 제목·번호·내용을 저장하고 플랫폼 화면으로 재구성합니다.",
        7,
        logo_blob,
    )
    add_image_frame(slide, 0.48, 1.55, 3.26, 2.55, title="실제 연보의 통계 예시", subtitle="연보 안의 통계표 이미지", name="이미지 삽입_연보 통계 예시")
    add_image_frame(slide, 6.26, 1.55, 3.26, 2.55, title="플랫폼에 등록된 통계", subtitle="제목·번호·표 내용이 재구성된 화면", name="이미지 삽입_플랫폼 등록 통계")
    add_arrow(slide, 3.90, 2.62, 0.34, 0.34, color=PRIMARY_2)
    add_rect(slide, 4.34, 2.02, 1.32, 1.56, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "업로드", 4.53, 2.28, 0.94, 0.28, size=12.2, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "HWPX   ·   MD", 4.46, 2.82, 1.08, 0.22, size=8.8, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "지원 확장자", 4.50, 3.16, 1.00, 0.18, size=7.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_arrow(slide, 5.78, 2.62, 0.34, 0.34, color=PRIMARY_2)
    steps = [("01", "파일 읽기", "연보 내용을 파싱"), ("02", "정보 저장", "제목·번호·표·기본 정보를 데이터베이스에 저장"), ("03", "연보 등록", "독립된 게시물로 생성")]
    for idx, (num, title, detail) in enumerate(steps):
        x = 0.72 + idx * 3.04
        add_open_item(slide, x, 4.34, 2.76, num, title, detail, color=PRIMARY_2)
    add_text(slide, "연보는 독립된 게시물이며 수정 및 삭제 가능", 2.30, 4.98, 5.40, 0.19, size=8.8, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_validation_types(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "2단계 · 검수 실행",
        "지원 검수 6종과 처리 방식",
        "현재 6가지 검수를 지원하며, 처리 방식과 결과 상태가 항목마다 다릅니다.",
        8,
        logo_blob,
    )
    add_section_title(slide, 0.48, 1.55, 5.55, "규칙·통계 기반 검수", number="4종")
    add_text(slide, "표 구조와 계산 관계를 미리 구성한 기준으로 확인", 0.48, 1.98, 5.35, 0.20, size=8.5, color=MUTED)
    rule_items = [
        (0.58, 2.35, "합계", "계·합계·소계를 다시 합산", False),
        (3.02, 2.35, "비율", "비율·구성비·증감 값을 재계산", False),
        (0.58, 3.55, "메타데이터", "단위·기준일·출처 확인", True),
        (3.02, 3.55, "이상치", "다른 값과 차이가 큰 수치 후보", True),
    ]
    for x, y, title, detail, review in rule_items:
        add_text(slide, title, x, y, 1.55, 0.25, size=11.0, color=TEXT, bold=True)
        add_text(slide, detail, x, y + 0.34, 1.96, 0.28, size=8.5, color=MUTED)
        add_status_badge(slide, "확인 필요" if review else "오류 의심", x + 1.35, y - 0.02, warning=review)
        add_rect(slide, x, y + 0.82, 1.98, 0.014, fill=DIVIDER, line=None)

    add_rect(slide, 6.30, 1.43, 3.22, 3.58, fill=PRIMARY_LIGHT, line=None)
    add_section_title(slide, 6.58, 1.55, 2.62, "LLM 기반 검수", number="2종")
    add_text(slide, "공식 용어·이전 결과를 먼저 적용하고\n남은 후보만 확인", 6.58, 1.99, 2.56, 0.45, size=8.5, color=MUTED)
    language_items = [("오탈자", "철자·문자 깨짐·표기 불일치", False), ("번역", "한글·영문 의미 대응 확인", True)]
    for idx, (title, detail, review) in enumerate(language_items):
        y = 2.73 + idx * 1.05
        add_text(slide, title, 6.60, y, 1.40, 0.25, size=11.0, color=TEXT, bold=True)
        add_text(slide, detail, 6.60, y + 0.34, 2.42, 0.22, size=8.5, color=MUTED)
        add_status_badge(slide, "확인 필요" if review else "오류 의심", 8.20, y - 0.02, warning=review)
        add_rect(slide, 6.60, y + 0.72, 2.42, 0.014, fill="C9DCEB", line=None)
    add_rich_text(
        slide,
        [
            {"text": "오류 의심  ", "size": 8.7, "color": ERROR, "bold": True},
            {"text": "합계·비율·오탈자    ", "size": 8.7, "color": TEXT},
            {"text": "확인 필요  ", "size": 8.7, "color": WARNING, "bold": True},
            {"text": "메타데이터·이상치·번역", "size": 8.7, "color": TEXT},
        ],
        1.95,
        4.84,
        6.10,
        0.22,
        align=PP_ALIGN.CENTER,
    )


def slide_rule_generation(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "검수 기준 생성 방식",
        "표마다 달라지는 검수 기준",
        "표 구조와 통계 특성을 단계적으로 분석해 각 표에 적용할 검수를 결정합니다.",
        9,
        logo_blob,
    )
    stages = [
        (0.48, "01", "공통 기준", "모든 표의 필수 정보 확인\n숫자 열의 특이값 후보 생성", ("메타", "이상치")),
        (2.78, "02", "표 유형별 기준", "지역별·연도별 등 구조 파악\n전국계와 지역 값 관계 적용", ("합계",)),
        (5.08, "03", "표별 자동 분석", "계·합계·소계 항목 탐색\n분자·분모·증감 관계 탐색", ("합계", "비율")),
        (7.38, "04", "통계별 예외 기준", "담당자가 확인한 규칙 반영\n고정 셀·분할표 예외 처리", ("오탈자", "번역", "합계", "비율")),
    ]
    add_rect(slide, 0.74, 1.90, 7.94, 0.025, fill="B9D0E1", line=None)
    for idx, (x, num, title, detail, tags) in enumerate(stages):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.02, 1.68, 0.45, 0.45, fill=PRIMARY if idx != 2 else PRIMARY_2, line=None, radius_name=f"검수 기준_{num}")
        add_text(slide, num, x + 0.02, 1.79, 0.45, 0.18, size=8.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x, 2.26, 1.98, 0.25, size=10.7, color=TEXT, bold=True)
        add_text(slide, detail, x, 2.65, 1.98, 0.54, size=8.4, color=MUTED)
        chip_x = x
        chip_w = 0.46 if len(tags) >= 4 else 0.70
        for tag in tags:
            add_rect(slide, chip_x, 3.44, chip_w, 0.28, fill=PRIMARY_LIGHT, line=None)
            add_text(slide, tag, chip_x, 3.51, chip_w, 0.17, size=7.4 if chip_w < 0.6 else 7.9, color=PRIMARY_2, bold=True, align=PP_ALIGN.CENTER)
            chip_x += chip_w + 0.05
    add_rect(slide, 0.62, 4.03, 8.76, 0.42, fill=PRIMARY_PALE, line=None)
    add_text(slide, "표 구조 기준", 0.86, 4.15, 1.05, 0.18, size=9.0, color=PRIMARY_2, bold=True)
    add_text(slide, "공통 + 유형 + 표 자동 분석 + 예외 기준   →   합계 · 비율 · 메타데이터 · 이상치", 2.02, 4.13, 6.95, 0.20, size=8.8, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.62, 4.58, 8.76, 0.42, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "언어 표현 기준", 0.86, 4.70, 1.05, 0.18, size=9.0, color=PRIMARY_2, bold=True)
    add_text(slide, "공식 용어 + 반복 표현 + 이전 검수 결과 + 남은 후보   →   오탈자 · 번역", 2.02, 4.68, 6.95, 0.20, size=8.8, color=TEXT, bold=True, align=PP_ALIGN.CENTER)


def add_calculation_slide(
    slide,
    logo_blob: bytes,
    *,
    page: int,
    section: str,
    title: str,
    subtitle: str,
    items: tuple[tuple[str, str, str], ...],
    formula: str,
    frame_title: str,
    frame_subtitle: str,
) -> None:
    clear_slide(slide)
    add_slide_header(slide, section, title, subtitle, page, logo_blob)
    add_section_title(slide, 0.48, 1.53, 3.38, "검수 대상과 계산 방식")
    for idx, (item_title, detail, example) in enumerate(items, start=1):
        add_open_item(slide, 0.48, 2.02 + (idx - 1) * 0.82, 3.38, f"0{idx}", item_title, detail, color=PRIMARY_2, formula=example, formula_w=1.38)
    add_image_frame(slide, 4.16, 1.53, 5.36, 3.44, title=frame_title, subtitle=frame_subtitle, accent=PRIMARY_2, name=f"이미지 삽입_{frame_title}")
    add_text(slide, formula, 0.68, 4.70, 2.98, 0.24, size=8.6, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "확인 화면  ·  오류 위치  ·  현재값  ·  재계산값  ·  판정 근거", 4.60, 4.75, 4.46, 0.20, size=8.4, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_sum(slide, logo_blob: bytes) -> None:
    add_calculation_slide(
        slide,
        logo_blob,
        page=10,
        section="합계 검수",
        title="구성값으로 합계 재계산",
        subtitle="계·합계·소계와 표에 표시된 산식을 구성값에서 다시 계산합니다.",
        items=(("지역별 합계", "세부 지역 값을 더해 전국계 확인", "전국 = 시·도 합"), ("행·열 합계", "같은 행의 열 또는 세부 행 합산", "계 = 남 + 여"), ("명시 산식·분할표", "a+b·a-b와 연결된 표 관계 계산", "c = a + b")),
        formula="구성값 합산  →  현재 합계와 비교  →  오류 의심",
        frame_title="합계 검수 예시 이미지 삽입",
        frame_subtitle="지역별 합계 또는 산식 오류가 보이는 화면 권장",
    )


def slide_ratio(slide, logo_blob: bytes) -> None:
    add_calculation_slide(
        slide,
        logo_blob,
        page=11,
        section="비율·증감 검수",
        title="분자·분모와 전년값 재계산",
        subtitle="비율·구성비·건당 값과 증감액·증감률을 다시 계산합니다.",
        items=(("비율·구성비", "항목 값과 전체 값을 분자·분모로 연결", "지역 ÷ 전체 × 100"), ("건당·인당 값", "건수·인원을 기준으로 원단위 계산", "건수 ÷ 인원"), ("증감액·증감률", "전년과 당해 값을 비교해 변화 계산", "(당해-전년)÷전년")),
        formula="총계는 비율 합산 X  →  합계 분자 ÷ 합계 분모",
        frame_title="비율·증감 검수 예시 이미지 삽입",
        frame_subtitle="지역별 비율 또는 증감 산식 오류가 보이는 화면 권장",
    )


def slide_language(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "오탈자·번역 검수",
        "후보를 줄인 뒤 LLM 검토",
        "모든 텍스트를 바로 LLM에 보내지 않고, 후보를 줄인 뒤 남은 항목만 검토합니다.",
        12,
        logo_blob,
    )
    add_rect(slide, 0.62, 1.48, 8.76, 0.43, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "후보 수집", 0.88, 1.60, 1.08, 0.18, size=9.4, color=PRIMARY_2, bold=True)
    add_text(slide, "표 제목 · 헤더 · 한글/영문 셀에서 검수할 문구를 수집", 2.08, 1.58, 6.86, 0.20, size=9.0, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 4.84, 1.99, 0.30, 0.28, color=PRIMARY_2)
    filters = [
        (0.62, "01", "공식 용어·번역 DB", "기관명  행정안전부\n지역명  세종특별자치시·연기면"),
        (3.67, "02", "반복 표현 DB", "계 · 합계 · 소계\n연보 전체에서 자주 쓰는 표현"),
        (6.72, "03", "이전 검수 결과", "같은 위치·문맥에서\n담당자가 확인한 표현"),
    ]
    for x, num, title, detail in filters:
        add_text(slide, num, x, 2.40, 0.42, 0.20, size=9.0, color=PRIMARY_2, bold=True)
        add_text(slide, title, x + 0.48, 2.35, 2.18, 0.25, size=10.4, color=TEXT, bold=True)
        add_rect(slide, x, 2.77, 2.66, 0.016, fill=PRIMARY_2, line=None)
        add_text(slide, detail, x + 0.06, 2.98, 2.54, 0.48, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        add_down_arrow(slide, x + 1.20, 3.55, 0.26, 0.28, color=PRIMARY_2)
    add_rect(slide, 2.56, 3.91, 4.88, 0.52, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "남은 후보만 LLM 검수", 2.88, 4.06, 4.24, 0.22, size=11.0, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
    add_pill(slide, "토큰 사용량 감소", 0.86, 4.50, 1.32, 0.28, fill=WARNING_BG, color=WARNING, size=8.0, line=None)
    add_pill(slide, "공식 표현 일관성", 7.82, 4.50, 1.32, 0.28, fill=PRIMARY_LIGHT, color=PRIMARY_2, size=8.0, line=None)
    add_text(slide, "담당자 확인 전에는 원문을 자동으로 바꾸지 않음", 2.72, 4.72, 4.56, 0.20, size=8.6, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_metadata_outlier(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "메타데이터·이상치 검수",
        "누락 정보와 이상값 후보 확인",
        "필수 정보의 누락과 통계적으로 차이가 큰 값을 ‘확인 필요’로 제시합니다.",
        13,
        logo_blob,
    )
    add_section_title(slide, 0.48, 1.53, 4.02, "메타데이터 검수")
    meta_items = [("01", "단위", "누락 여부 + 기존 기준 단위와 일치"), ("02", "기준일", "입력 여부 확인"), ("03", "출처", "입력 여부 확인")]
    for idx, (num, title, detail) in enumerate(meta_items):
        add_open_item(slide, 0.48, 2.08 + idx * 0.78, 4.02, num, title, detail, color=PRIMARY_2)
    add_text(slide, "자동 확인 범위  ·  단위  ·  기준일  ·  출처", 0.80, 4.60, 2.95, 0.20, size=8.5, color=PRIMARY, bold=True)
    add_status_badge(slide, "확인 필요", 3.55, 4.55, warning=True)

    add_section_title(slide, 5.05, 1.53, 4.47, "이상치 검수")
    add_text(slide, "같은 숫자 열에서 값이 크게 벗어난 후보", 5.18, 2.08, 4.02, 0.22, size=9.4, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 5.42, 3.16, 3.75, 0.018, fill=DIVIDER, line=None)
    values = [(5.58, "9", PRIMARY_2, 3.03), (6.22, "10", PRIMARY_2, 3.03), (6.86, "10", PRIMARY_2, 3.03), (7.50, "11", PRIMARY_2, 3.03), (8.62, "43", ERROR, 2.70)]
    for x, value, color, y in values:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, y, 0.22, 0.22, fill=color, line=None, radius_name=f"이상치 값_{value}")
        add_text(slide, value, x - 0.08, 3.38, 0.38, 0.18, size=8.0, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_status_badge(slide, "확인 필요", 8.25, 2.31, warning=True)
    add_text(slide, "• 값이 5개 이상인 숫자 열\n• 합계·소계 제외\n• 표당 차이가 큰 값 최대 1건 제시", 5.40, 3.73, 3.70, 0.64, size=8.4, color=TEXT)
    add_text(slide, "오류 확정이 아니라 담당자가 확인할 후보", 5.55, 4.58, 3.42, 0.20, size=8.6, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_edit_recheck(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "수정·재검수",
        "보호된 수정과 선택적 재검수",
        "통계별 수정 권한을 확인한 뒤 내용을 고치고, 변경 범위에 맞춰 다시 검수합니다.",
        14,
        logo_blob,
    )
    add_section_title(slide, 0.48, 1.52, 4.06, "수정 보호")
    protection = [("01", "비밀번호 확인", "통계별 수정 권한 확인"), ("02", "편집 화면", "표 정보·셀 값·행/열 수정"), ("03", "수정 저장", "변경 내용 반영")]
    add_rect(slide, 0.86, 2.42, 3.15, 0.020, fill="B9D0E1", line=None)
    for idx, (num, title, detail) in enumerate(protection):
        x = 0.60 + idx * 1.34
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 2.20, 0.44, 0.44, fill=PRIMARY_2, line=None, radius_name=f"수정 보호_{num}")
        add_text(slide, num, x, 2.31, 0.44, 0.18, size=7.8, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x - 0.24, 2.82, 0.94, 0.22, size=9.5, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x - 0.36, 3.14, 1.18, 0.38, size=7.9, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "저장 즉시 검수 자동 실행 X  ·  담당자가 재검수 실행", 0.64, 3.66, 3.70, 0.21, size=8.3, color=WARNING, bold=True, align=PP_ALIGN.CENTER)

    add_section_title(slide, 5.05, 1.52, 4.47, "재검수 원칙")
    principles = [("01", "값 변경", "기존 계산 기준으로 현재 값 재검수"), ("02", "구조 변경", "새로운 표 구조에 맞춰 기준 재생성"), ("03", "변경 없음", "이전 검수 결과 재사용")]
    for idx, (num, title, detail) in enumerate(principles):
        add_open_item(slide, 5.05, 2.04 + idx * 0.70, 4.47, num, title, detail, color=PRIMARY_2)

    add_rect(slide, 0.62, 4.36, 8.76, 0.58, fill=PRIMARY_PALE, line=None)
    flow = [("수정 저장", PRIMARY), ("재검수 필요", WARNING), ("재검수 실행", PRIMARY_2), ("새 결과 갱신", PRIMARY)]
    for idx, (label, color) in enumerate(flow):
        x = 0.92 + idx * 2.12
        add_text(slide, label, x, 4.54, 1.30, 0.20, size=8.8, color=color, bold=True, align=PP_ALIGN.CENTER)
        if idx < 3:
            add_arrow(slide, x + 1.52, 4.51, 0.26, 0.25, color=PRIMARY_2 if idx != 0 else WARNING)


def slide_download(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "검수 결과 활용",
        "표와 검수 리포트 내려받기",
        "필요한 통계표와 검수 결과를 업무 목적에 맞춰 내려받습니다.",
        15,
        logo_blob,
    )
    add_rect(slide, 5.00, 1.52, 0.015, 3.50, fill=DIVIDER, line=None)
    add_text(slide, "01", 0.62, 1.65, 0.54, 0.24, size=11.0, color=PRIMARY_2, bold=True)
    add_text(slide, "표", 1.20, 1.56, 1.20, 0.42, size=20.0, color=TEXT, bold=True)
    add_text(slide, "선택한 통계표를 검토·공유", 0.62, 2.10, 3.90, 0.24, size=10.5, color=PRIMARY_2, bold=True)
    add_rect(slide, 0.62, 2.50, 3.88, 0.018, fill=DIVIDER, line=None)
    add_bullets(slide, ("선택한 통계표와 하위 표", "한글·영문 헤더와 셀 내용", "단위·기준일·출처 등 표 정보", "담당자 수정·검토용 자료"), 0.72, 2.83, 3.65, 1.78, size=9.0, color=TEXT, bullet_color=PRIMARY_2, gap=12)
    add_text(slide, "02", 5.46, 1.65, 0.54, 0.24, size=11.0, color=PRIMARY_2, bold=True)
    add_text(slide, "검수 리포트", 6.04, 1.56, 2.30, 0.42, size=20.0, color=TEXT, bold=True)
    add_text(slide, "오류 확인 내역과 판정 근거 보고", 5.46, 2.10, 3.90, 0.24, size=10.5, color=PRIMARY_2, bold=True)
    add_rect(slide, 5.46, 2.50, 3.88, 0.018, fill=DIVIDER, line=None)
    add_bullets(slide, ("오류 의심·확인 필요 항목", "문제 셀과 근거 셀 위치", "현재값·재계산값·판정 사유", "표 정보 확인 결과"), 5.56, 2.83, 3.65, 1.78, size=9.0, color=TEXT, bullet_color=PRIMARY_2, gap=12)
    add_text(slide, "필요한 범위와 검수 상태를 선택해 담당자 공유·수정·결과 보고에 활용", 2.02, 4.83, 5.96, 0.22, size=8.8, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def slide_effects(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "기대효과",
        "역할별 부담 감소, 협업 효율 향상",
        "발간 담당자·통계 담당자·검수 진행자가 같은 화면과 기준으로 협업합니다.",
        16,
        logo_blob,
    )
    groups = [
        (0.62, "연", "통계연보 담당자", "전체 진행 상황 관리", ("연보별 검수 상태 한눈에 확인", "오류 통계와 수정 현황 공유", "발간 전 최종 결과 정리"), NAVY),
        (3.68, "통", "각 통계 담당자", "내 통계의 문제를 바로 수정", ("문제 셀과 근거를 빠르게 확인", "화면에서 직접 수정·저장", "변경된 표 중심으로 재검수"), PRIMARY_2),
        (6.74, "검", "검수 진행자", "반복 계산 대신 판단에 집중", ("합계·비율 자동 재계산", "확인 우선순위와 근거 제시", "검수 리포트로 결과 공유"), BLUE),
    ]
    for x, icon, title, subtitle, bullets, color in groups:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.98, 1.55, 0.52, 0.52, fill=color, line=None, radius_name=f"역할_{title}")
        add_text(slide, icon, x + 0.98, 1.66, 0.52, 0.22, size=13.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x, 2.26, 2.48, 0.25, size=11.0, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, subtitle, x, 2.64, 2.48, 0.22, size=8.6, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_rect(slide, x + 0.12, 3.03, 2.24, 0.016, fill=DIVIDER, line=None)
        add_bullets(slide, bullets, x + 0.12, 3.32, 2.26, 1.18, size=8.6, color=TEXT, bullet_color=color, gap=10)
    add_rect(slide, 0.88, 4.74, 8.24, 0.34, fill=NAVY, line=None)
    add_text(slide, "검수 시간 단축  ·  오류 누락 위험 감소  ·  공유·수정·재검수 흐름 단순화", 1.18, 4.83, 7.64, 0.19, size=9.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_future(slide, logo_blob: bytes) -> None:
    clear_slide(slide)
    add_slide_header(
        slide,
        "향후 보완점",
        "안전한 기관 운영을 위한 보완",
        "기관 운영을 위해 권한·이력·AI 환경을 단계적으로 보완해야 합니다.",
        17,
        logo_blob,
    )
    roadmap = [
        (0.62, "01", "권한·보안", ("기관 계정과 통합 로그인", "담당자별 역할·수정 권한", "세션 만료·접속 제한")),
        (3.68, "02", "수정 이력 기록", ("셀 단위 변경 전·후 기록", "수정자·시간·사유 확인", "되돌리기·변경 알림")),
        (6.74, "03", "내부망·내부 LLM", ("내부망 환경에 서비스 도입", "외부 LLM을 내부 모델로 전환", "민감 정보와 사용 이력 관리")),
    ]
    add_rect(slide, 1.08, 1.90, 6.75, 0.025, fill="B9D0E1", line=None)
    for idx, (x, num, title, bullets) in enumerate(roadmap):
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.96, 1.67, 0.48, 0.48, fill=PRIMARY if idx != 1 else PRIMARY_2, line=None, radius_name=f"향후 보완_{num}")
        add_text(slide, num, x + 0.96, 1.79, 0.48, 0.18, size=8.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x, 2.40, 2.48, 0.25, size=11.2, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_rect(slide, x + 0.12, 2.83, 2.24, 0.016, fill=DIVIDER, line=None)
        add_bullets(slide, bullets, x + 0.12, 3.15, 2.26, 1.30, size=8.7, color=TEXT, bullet_color=PRIMARY_2, gap=11)
    add_rect(slide, 1.22, 4.63, 7.56, 0.37, fill=PRIMARY_LIGHT, line=None)
    add_text(slide, "권한 체계 정비  →  변경 이력 관리  →  내부망·내부 모델 전환", 1.57, 4.73, 6.86, 0.19, size=8.9, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)


def clear_notes(prs: Presentation) -> None:
    for slide in prs.slides:
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            frame.clear()
            frame.text = ""


def build(source: Path, output: Path) -> None:
    prs = Presentation(source)
    cover_logo_blob, normal_logo_blob = find_logo_blobs(prs)

    # Preserve only the first slide and rebuild the new 17-slide story.
    slide_ids = prs.slides._sldIdLst
    while len(prs.slides) > 1:
        slide_id = slide_ids[1]
        prs.part.drop_rel(slide_id.rId)
        del slide_ids[1]

    build_cover(prs.slides[0], cover_logo_blob)

    builders = [
        slide_need_for_review,
        slide_as_is_to_be,
        slide_main_page,
        slide_detail_page,
        slide_service_flow,
        slide_upload,
        slide_validation_types,
        slide_rule_generation,
        slide_sum,
        slide_ratio,
        slide_language,
        slide_metadata_outlier,
        slide_edit_recheck,
        slide_download,
        slide_effects,
        slide_future,
    ]

    blank_layout = prs.slide_layouts[10]
    for layout_shape in list(blank_layout.shapes):
        element = layout_shape._element
        element.getparent().remove(element)
    for builder in builders:
        slide = prs.slides.add_slide(blank_layout)
        slide._element.set("showMasterSp", "0")
        builder(slide, normal_logo_blob)

    clear_notes(prs)
    force_font_family(prs, FONT_NAME)
    prs.core_properties.title = "통계연보 검수 보조 시스템"
    prs.core_properties.subject = "통계연보 검수 업무 개선과 서비스 주요 기능"
    prs.core_properties.author = "행정안전부 데이터정보화담당관"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()

"""Build the editable project handover presentation.

Usage from the repository root:
    ./.venv/bin/python scripts/build_project_handover_presentation.py

The deck is intentionally built only from editable PowerPoint shapes and text.
The detailed source of truth remains docs/PROJECT_HANDOVER_GUIDE.md.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "exports" / "통계연보_검수_보조_시스템_프로젝트_인수인계_가이드_상세_최신본.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "맑은 고딕"
MONO = "Consolas"

NAVY = "0B2545"
BLUE = "1D5F9A"
BRIGHT_BLUE = "2E78B7"
SKY = "EAF3FA"
PALE_BLUE = "F5F9FC"
TEAL = "177A74"
TEAL_BG = "E8F5F3"
AMBER = "B76E00"
AMBER_BG = "FFF3D8"
RED = "B83A4B"
RED_BG = "FCECEF"
GREEN = "2C7A55"
GREEN_BG = "E9F6EF"
TEXT = "172432"
MUTED = "617181"
SOFT = "8B99A7"
BORDER = "D8E1E8"
LIGHT = "F7F9FB"
WHITE = "FFFFFF"
BLACK = "111827"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(
    run,
    *,
    size: float,
    fill: str = TEXT,
    bold: bool = False,
    font: str = FONT,
) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            r_pr.append(node)
        node.set("typeface", font)


def remove_style(shape) -> None:
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str | None = BORDER,
    radius: bool = False,
    width: float = 0.8,
):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    remove_style(shape)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(width)
    return shape


def line(slide, x1: float, y1: float, x2: float, y2: float, *, fill: str = BORDER, width: float = 1.0):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = color(fill)
    shape.line.width = Pt(width)
    return shape


def text_box(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 14,
    fill: str = TEXT,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.04,
    font: str = FONT,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    remove_style(shape)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    set_run(run, size=size, fill=fill, bold=bold, font=font)
    return shape


def rich_text(
    slide,
    runs: Sequence[dict],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.04,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    remove_style(shape)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    for spec in runs:
        run = p.add_run()
        run.text = spec["text"]
        set_run(
            run,
            size=spec.get("size", 14),
            fill=spec.get("fill", TEXT),
            bold=spec.get("bold", False),
            font=spec.get("font", FONT),
        )
    return shape


def bullets(
    slide,
    items: Iterable[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 14,
    fill: str = TEXT,
    bullet_fill: str = BLUE,
    gap: float = 6,
    leading: float = 1.10,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    remove_style(shape)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(gap)
        p.line_spacing = leading
        marker = p.add_run()
        marker.text = "•  "
        set_run(marker, size=size, fill=bullet_fill, bold=True)
        body = p.add_run()
        body.text = item
        set_run(body, size=size, fill=fill)
    return shape


def code_block(
    slide,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 10.5,
    title: str | None = None,
):
    rect(slide, x, y, w, h, fill=NAVY, line=None, radius=True)
    top = y + 0.16
    if title:
        text_box(slide, title, x + 0.20, y + 0.11, w - 0.40, 0.26, size=10, fill="90D4FF", bold=True)
        top = y + 0.43
    text_box(
        slide,
        value,
        x + 0.18,
        top,
        w - 0.36,
        h - (top - y) - 0.12,
        size=size,
        fill="F5F9FC",
        font=MONO,
        margin=0,
    )


def pill(slide, value: str, x: float, y: float, w: float, *, fill: str = SKY, text_fill: str = BLUE):
    rect(slide, x, y, w, 0.34, fill=fill, line=None, radius=True)
    text_box(
        slide,
        value,
        x + 0.04,
        y + 0.04,
        w - 0.08,
        0.24,
        size=10,
        fill=text_fill,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )


def card(
    slide,
    title: str,
    body: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    accent: str = BLUE,
    fill: str = WHITE,
    title_size: float = 15,
    body_size: float = 11.5,
):
    rect(slide, x, y, w, h, fill=fill, line=BORDER, radius=True)
    rect(slide, x, y, 0.08, h, fill=accent, line=None, radius=False)
    text_box(slide, title, x + 0.25, y + 0.14, w - 0.45, 0.32, size=title_size, fill=TEXT, bold=True)
    text_box(slide, body, x + 0.25, y + 0.56, w - 0.45, h - 0.67, size=body_size, fill=MUTED)


def number_step(
    slide,
    number: int,
    title: str,
    detail: str,
    x: float,
    y: float,
    w: float,
    *,
    accent: str = BLUE,
):
    rect(slide, x, y, w, 0.68, fill=WHITE, line=BORDER, radius=True)
    rect(slide, x + 0.12, y + 0.12, 0.44, 0.44, fill=accent, line=None, radius=True)
    text_box(slide, str(number), x + 0.12, y + 0.16, 0.44, 0.26, size=11, fill=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    text_box(slide, title, x + 0.70, y + 0.11, w - 0.85, 0.24, size=12.2, fill=TEXT, bold=True)
    text_box(slide, detail, x + 0.70, y + 0.36, w - 0.85, 0.22, size=9.6, fill=MUTED)


def set_background(slide, fill: str = WHITE) -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color(fill)


def header(slide, section: str, title: str, subtitle: str, page: int) -> None:
    set_background(slide, WHITE)
    rect(slide, 0, 0, SLIDE_W, 0.12, fill=BLUE, line=None)
    text_box(slide, section.upper(), 0.62, 0.38, 3.2, 0.25, size=9.5, fill=BLUE, bold=True)
    text_box(slide, title, 0.62, 0.75, 12.0, 0.55, size=25, fill=NAVY, bold=True)
    text_box(slide, subtitle, 0.64, 1.34, 11.9, 0.35, size=11.5, fill=MUTED)
    line(slide, 0.62, 1.82, 12.71, 1.82, fill=BORDER, width=0.8)
    text_box(slide, "통계연보 검수 보조 시스템  |  프로젝트 인수인계", 0.62, 7.12, 7.7, 0.20, size=8.2, fill=SOFT)
    text_box(slide, str(page), 12.15, 7.10, 0.55, 0.20, size=8.5, fill=SOFT, bold=True, align=PP_ALIGN.RIGHT)


def section_slide(slide, number: str, title: str, subtitle: str, page: int, *, accent: str = BLUE) -> None:
    set_background(slide, NAVY)
    rect(slide, 0, 0, 0.18, SLIDE_H, fill=accent, line=None)
    text_box(slide, number, 0.82, 0.75, 2.2, 1.0, size=54, fill=accent, bold=True)
    text_box(slide, title, 0.82, 2.18, 11.2, 0.85, size=34, fill=WHITE, bold=True)
    text_box(slide, subtitle, 0.84, 3.30, 10.9, 0.65, size=16, fill="B9CEE1")
    line(slide, 0.84, 4.35, 11.95, 4.35, fill="244566", width=1.0)
    text_box(slide, "상세 절차는 docs/PROJECT_HANDOVER_GUIDE.md", 0.84, 5.78, 6.8, 0.35, size=11.5, fill="8FB2D1")
    text_box(slide, str(page), 12.12, 6.95, 0.55, 0.22, size=9, fill="8FB2D1", align=PP_ALIGN.RIGHT)


def table(
    slide,
    rows: Sequence[Sequence[str]],
    widths: Sequence[float],
    x: float,
    y: float,
    *,
    row_h: float = 0.50,
    header_fill: str = NAVY,
    font_size: float = 10.5,
):
    total_w = sum(widths)
    for r_index, row in enumerate(rows):
        current_x = x
        current_y = y + r_index * row_h
        for c_index, value in enumerate(row):
            cell_fill = header_fill if r_index == 0 else (WHITE if r_index % 2 else PALE_BLUE)
            text_fill = WHITE if r_index == 0 else TEXT
            rect(slide, current_x, current_y, widths[c_index], row_h, fill=cell_fill, line=BORDER, width=0.6)
            text_box(
                slide,
                value,
                current_x + 0.08,
                current_y + 0.06,
                widths[c_index] - 0.16,
                row_h - 0.12,
                size=font_size if r_index else font_size - 0.2,
                fill=text_fill,
                bold=r_index == 0,
                valign=MSO_ANCHOR.MIDDLE,
                margin=0,
            )
            current_x += widths[c_index]
    return total_w, row_h * len(rows)


def build_cover(slide) -> None:
    set_background(slide, NAVY)
    rect(slide, 0, 0, 0.20, SLIDE_H, fill=BRIGHT_BLUE, line=None)
    pill(slide, "PROJECT HANDOVER · 2026.08", 0.82, 0.78, 2.55, fill="173C61", text_fill="8ED2FF")
    text_box(slide, "통계연보 검수 보조 시스템", 0.82, 1.65, 10.9, 0.48, size=16, fill="A8C4DC", bold=True)
    text_box(slide, "프로젝트 인수인계\n실행·수정·운영 가이드", 0.82, 2.26, 10.7, 1.55, size=37, fill=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "macOS에서 Windows로 — 처음 받는 사람도 그대로 따라 하는 실무 절차", 0.85, 4.18, 11.2, 0.42, size=15, fill="90B9D8")
    rect(slide, 0.84, 5.25, 11.6, 0.82, fill="122F4D", line="274E71", radius=True)
    rich_text(
        slide,
        [
            {"text": "시작점  ", "size": 12, "fill": "8ED2FF", "bold": True},
            {"text": "설치 → 실행 → 사용 → 수정 → 테스트 → 배포 → 복구", "size": 14, "fill": WHITE, "bold": True},
        ],
        1.12, 5.49, 10.9, 0.30, valign=MSO_ANCHOR.MIDDLE,
    )
    text_box(slide, "기준일 2026-08-28  |  상세본: docs/PROJECT_HANDOVER_GUIDE.md", 0.84, 6.85, 9.4, 0.25, size=9.5, fill="789AB6")
    text_box(slide, "1", 12.12, 6.85, 0.55, 0.25, size=9, fill="789AB6", align=PP_ALIGN.RIGHT)


def slide_part_install(slide, page: int) -> None:
    section_slide(slide, "01", "설치 단계", "도구 설치부터 빈 SQLite 기동과 정상 확인까지", page, accent=BRIGHT_BLUE)


def slide_part_features(slide, page: int) -> None:
    section_slide(slide, "02", "기능 설명", "연보 등록·검수·원장·수정·다운로드의 실제 사용 흐름", page, accent=TEAL)


def slide_part_maintenance(slide, page: int) -> None:
    section_slide(slide, "03", "유지보수와 프로젝트 구조", "코드 변경·DB·테스트·배포·보안·복구와 최종 인계", page, accent=AMBER)


def slide_outcome(slide, page: int) -> None:
    header(slide, "ORIENTATION", "이 자료를 끝까지 보면 할 수 있는 것", "소스 전달을 넘어 실제 운영 가능한 상태까지 만든다.", page)
    items = [
        ("01", "설치", "macOS와 Windows에서 같은 개발 환경을 만든다.", BLUE, SKY),
        ("02", "실행", "프론트·백엔드·DB를 올리고 정상 여부를 판단한다.", TEAL, TEAL_BG),
        ("03", "수정", "기능별 코드 위치를 찾고 안전한 순서로 변경한다.", AMBER, AMBER_BG),
        ("04", "검증", "테스트 기준선과 알려진 실패를 구분한다.", RED, RED_BG),
        ("05", "운영", "배포·백업·복구·계정 인계를 수행한다.", GREEN, GREEN_BG),
    ]
    for i, (num, title, body, accent, bg) in enumerate(items):
        x = 0.65 + (i % 3) * 4.15
        y = 2.18 + (i // 3) * 1.72
        w = 3.83 if i < 3 else 5.90
        if i >= 3:
            x = 0.65 + (i - 3) * 6.15
        rect(slide, x, y, w, 1.38, fill=bg, line=None, radius=True)
        text_box(slide, num, x + 0.20, y + 0.18, 0.55, 0.30, size=12, fill=accent, bold=True)
        text_box(slide, title, x + 0.82, y + 0.16, w - 1.05, 0.33, size=17, fill=TEXT, bold=True)
        text_box(slide, body, x + 0.22, y + 0.64, w - 0.44, 0.52, size=11.2, fill=MUTED)


def slide_five_rules(slide, page: int) -> None:
    header(slide, "ORIENTATION", "시작 전에 기억할 다섯 가지", "문제의 대부분은 프로세스·데이터·환경을 섞을 때 생긴다.", page)
    rules = [
        ("1", "두 서버", "React :5173 + FastAPI :8000을 각각 실행"),
        ("2", "별도 인계", "Git 밖의 DB·uploads·비밀값은 따로 전달"),
        ("3", "안전한 시작", "처음은 빈 테스트 DB + LLM OFF"),
        ("4", "루트 실행", "백엔드는 저장소 루트에서 실행"),
        ("5", "먼저 백업", "실제 DB는 최초 실행 전 반드시 복사"),
    ]
    for i, (num, title, body) in enumerate(rules):
        y = 2.08 + i * 0.90
        rect(slide, 0.72, y, 11.88, 0.72, fill=WHITE if i % 2 == 0 else PALE_BLUE, line=BORDER, radius=True)
        rect(slide, 0.88, y + 0.13, 0.46, 0.46, fill=BLUE if i < 4 else RED, line=None, radius=True)
        text_box(slide, num, 0.88, y + 0.17, 0.46, 0.28, size=12, fill=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0)
        text_box(slide, title, 1.58, y + 0.14, 1.65, 0.28, size=14, fill=NAVY, bold=True)
        text_box(slide, body, 3.22, y + 0.15, 8.95, 0.28, size=12.2, fill=MUTED)


def slide_snapshot(slide, page: int) -> None:
    header(slide, "HANDOVER RISK", "현재 작업 트리: 그냥 clone하면 같은 상태가 아니다", "2026-08-28 점검 스냅샷 — 인계 직전에 다시 측정해야 한다.", page)
    stats = [
        ("45", "git status 변경 항목", AMBER, AMBER_BG),
        ("11", "Git 미추적 항목", RED, RED_BG),
        ("4.1GB", "로컬 주 SQLite", BLUE, SKY),
        ("1.3GB", "로컬 업로드 원본", TEAL, TEAL_BG),
    ]
    for i, (value, label, accent, bg) in enumerate(stats):
        x = 0.68 + i * 3.05
        rect(slide, x, 2.13, 2.78, 1.42, fill=bg, line=None, radius=True)
        text_box(slide, value, x + 0.20, 2.35, 2.38, 0.45, size=24, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, label, x + 0.18, 2.91, 2.42, 0.31, size=10.5, fill=TEXT, bold=True, align=PP_ALIGN.CENTER)
    card(slide, "왜 위험한가", "SQLite와 backend/uploads는 .gitignore 대상이다. 소스만 전달하면 기존 연보·검수 run·재시도 가능한 원본이 사라진다.", 0.70, 4.02, 5.90, 1.55, accent=RED, fill=RED_BG)
    card(slide, "인계 직전 조치", "git status를 정리하고, 최종 커밋·DB 백업 시각·uploads 백업 시각·각 SHA-256을 한 장의 인계 기록으로 남긴다.", 6.75, 4.02, 5.88, 1.55, accent=GREEN, fill=GREEN_BG)
    text_box(slide, "기준 브랜치 master  |  마지막 커밋 6292d55 이후 로컬 변경 포함", 0.76, 6.12, 11.75, 0.35, size=11.5, fill=MUTED, align=PP_ALIGN.CENTER)


def slide_receive(slide, page: int) -> None:
    header(slide, "HANDOVER", "받아야 할 것은 여섯 묶음", "코드·데이터·비밀·계정·기준·문서를 분리하면 누락과 유출을 동시에 줄인다.", page)
    labels = [
        ("01", "SOURCE", "최종 commit / branch", BLUE),
        ("02", "DATABASE", "SQLite 또는 PostgreSQL", TEAL),
        ("03", "UPLOADS", "원본 + 같은 시점 백업", BRIGHT_BLUE),
        ("04", "SECRETS", "비밀관리 채널", RED),
        ("05", "ACCOUNTS", "GitHub·Vercel·Railway·R2", AMBER),
        ("06", "BASELINE", "승인된 검수 기준·대표 파일", GREEN),
    ]
    for i, (num, title, body, accent) in enumerate(labels):
        x = 0.70 + (i % 3) * 4.12
        y = 2.10 + (i // 3) * 1.72
        rect(slide, x, y, 3.82, 1.38, fill=WHITE, line=BORDER, radius=True)
        text_box(slide, num, x + 0.20, y + 0.17, 0.55, 0.28, size=10, fill=accent, bold=True)
        text_box(slide, title, x + 0.82, y + 0.15, 2.62, 0.32, size=14.5, fill=NAVY, bold=True)
        text_box(slide, body, x + 0.20, y + 0.68, 3.38, 0.40, size=10.5, fill=MUTED)
    rect(slide, 0.70, 5.74, 12.0, 0.68, fill=NAVY, line=None, radius=True)
    text_box(slide, "금지: .env/API 키를 DB·uploads 압축 안에 함께 넣어 일반 메신저로 전송", 0.98, 5.95, 11.43, 0.25, size=12, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_product(slide, page: int) -> None:
    header(slide, "PROJECT", "이 시스템이 하는 일", "연보 업로드부터 검수·수정·다음 연도 생성까지 한 흐름으로 잇는다.", page)
    features = [
        ("업로드", "HWPX · Markdown · Excel", BLUE),
        ("구조화", "표 · 셀 · 메타정보", TEAL),
        ("자동 검수", "합계 · 비율 · 증감 · 언어", AMBER),
        ("결과 확인", "상태 · 근거 · 하이라이트", RED),
        ("웹 수정", "인증 후 수정 · 재검수", BRIGHT_BLUE),
        ("다음 연도", "원장 3종 + 기준 연보", GREEN),
    ]
    for i, (title, body, accent) in enumerate(features):
        x = 0.67 + (i % 3) * 4.18
        y = 2.10 + (i // 3) * 1.70
        card(slide, title, body, x, y, 3.87, 1.35, accent=accent, fill=PALE_BLUE if i % 2 else WHITE, title_size=15.5, body_size=11.4)
    text_box(slide, "핵심 가치: 반복 계산은 시스템이, 최종 판단은 담당자가", 0.75, 5.83, 11.80, 0.45, size=17, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)


def slide_user_features(slide, page: int) -> None:
    header(slide, "USE", "현재 화면에서 찾을 수 있는 주요 기능", "버튼 이름은 2026-08-28 현재 구현을 기준으로 한다.", page)
    features = [
        ("연보 등록", "전체 연보 / 증분 업로드", BLUE, SKY),
        ("통계 연보", "검색·상태 필터·표 상세", TEAL, TEAL_BG),
        ("검수 진행", "선택 검수·재사용·중단", AMBER, AMBER_BG),
        ("연보 비교", "추가·삭제·변경 표/셀", RED, RED_BG),
        ("통계 원장", "통계명·ID·부서·시점 조회", BRIGHT_BLUE, SKY),
        ("표 수정", "메타·셀·복제·삭제 후 재검수", GREEN, GREEN_BG),
        ("다운로드", "표 CSV/XLSX·연보·검수 ZIP", BLUE, PALE_BLUE),
        ("수정 권한", "통계별 비밀번호 변경", RED, RED_BG),
    ]
    for i, (title, body, accent, bg) in enumerate(features):
        x = 0.70 + (i % 4) * 3.06
        y = 2.12 + (i // 4) * 1.70
        rect(slide, x, y, 2.78, 1.38, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.18, y + 0.20, 2.42, 0.32, size=14.5, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.18, y + 0.72, 2.42, 0.38, size=9.8, fill=TEXT, align=PP_ALIGN.CENTER)
    rect(slide, 0.72, 5.83, 11.86, 0.58, fill=NAVY, line=None, radius=True)
    text_box(slide, "화면의 증감률 결과는 비율 검수, 증감액 결과는 합계 검수 범주로 표시", 1.00, 6.01, 11.30, 0.24, size=11.5, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_architecture(slide, page: int) -> None:
    header(slide, "ARCHITECTURE", "다섯 구성요소와 연결", "문제가 생기면 화살표 순서대로 어느 경계에서 끊겼는지 본다.", page)
    boxes = [
        (0.58, 2.56, 2.10, 1.18, "브라우저", "사용자 화면", SKY, BLUE),
        (3.08, 2.56, 2.12, 1.18, "React + Vite", "frontend :5173", SKY, BLUE),
        (5.62, 2.56, 2.12, 1.18, "FastAPI", "backend :8000", TEAL_BG, TEAL),
        (8.16, 2.08, 2.15, 1.18, "DB", "SQLite / PostgreSQL", AMBER_BG, AMBER),
        (8.16, 3.56, 2.15, 1.18, "원본", "local / S3·R2", GREEN_BG, GREEN),
        (10.70, 2.56, 2.05, 1.18, "LLM", "선택적 외부 호출", RED_BG, RED),
    ]
    for x, y, w, h, title, body, bg, accent in boxes:
        rect(slide, x, y, w, h, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.12, y + 0.18, w - 0.24, 0.30, size=15, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.12, y + 0.62, w - 0.24, 0.24, size=10, fill=TEXT, align=PP_ALIGN.CENTER)
    for x1, y1, x2, y2 in [(2.68, 3.15, 3.08, 3.15), (5.20, 3.15, 5.62, 3.15), (7.74, 2.85, 8.16, 2.67), (7.74, 3.45, 8.16, 4.07), (10.31, 2.67, 10.70, 3.15)]:
        line(slide, x1, y1, x2, y2, fill=BLUE, width=2.0)
    rect(slide, 0.72, 5.22, 12.0, 0.82, fill=LIGHT, line=BORDER, radius=True)
    text_box(slide, "운영 worker는 FastAPI와 같은 PostgreSQL·S3/R2를 사용해 queued 작업을 처리", 0.98, 5.47, 11.45, 0.28, size=12.5, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)


def slide_flow(slide, page: int) -> None:
    header(slide, "ARCHITECTURE", "업로드와 검수는 작업 상태로 이어진다", "실패해도 단계·오류가 DB에 남고, 조건이 맞으면 재시도할 수 있다.", page)
    stages = [
        ("01", "원본 저장", "크기·SHA-256", BLUE),
        ("02", "파싱", "표·셀·메타", TEAL),
        ("03", "DB 저장", "연보 snapshot", GREEN),
        ("04", "계산 검수", "규칙·profile", AMBER),
        ("05", "언어 검수", "사전·캐시·LLM", RED),
        ("06", "완료", "화면·export", BRIGHT_BLUE),
    ]
    for i, (num, title, body, accent) in enumerate(stages):
        x = 0.55 + i * 2.10
        rect(slide, x, 2.48, 1.72, 1.52, fill=WHITE, line=BORDER, radius=True)
        text_box(slide, num, x + 0.12, 2.66, 0.42, 0.24, size=9, fill=accent, bold=True)
        text_box(slide, title, x + 0.12, 3.02, 1.48, 0.30, size=12.5, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.12, 3.48, 1.48, 0.25, size=9.2, fill=MUTED, align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            text_box(slide, "→", x + 1.73, 3.02, 0.36, 0.32, size=17, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)
    card(slide, "inprocess", "로컬 기본. API 프로세스가 백그라운드 작업 실행", 0.72, 4.68, 5.75, 1.08, accent=BLUE, fill=SKY, title_size=14, body_size=10.5)
    card(slide, "worker", "운영 분리. persisted queue + lease + heartbeat", 6.85, 4.68, 5.75, 1.08, accent=TEAL, fill=TEAL_BG, title_size=14, body_size=10.5)


def slide_tree(slide, page: int) -> None:
    header(slide, "CODE MAP", "폴더를 보면 책임이 보인다", "처음부터 모든 파일을 읽지 말고 변경하려는 기능의 경계부터 좁힌다.", page)
    left = """backend/app/
├─ main.py        시작·CORS·health
├─ api/           HTTP 경계
├─ application/   유스케이스 연결
├─ services/      업무 규칙
├─ ingest/        HWPX·MD·Excel 파서
├─ validation/    검수 엔진·언어
├─ db/            schema·migration
├─ storage/       local·S3/R2
├─ workflows/     job·worker
└─ yearbook_data/ 원장·재현"""
    right = """frontend/src/
├─ App.tsx        최상위 화면
├─ api/           백엔드 호출
├─ components/    화면 조각
├─ hooks/         상태·비동기 흐름
├─ types.ts       API 타입 계약
├─ utils/         표시·다운로드
└─ styles/        전역 디자인

tests/
├─ backend/tests  unittest
└─ *.test.tsx     Vitest"""
    code_block(slide, left, 0.70, 2.10, 5.85, 3.92, size=11.2, title="BACKEND")
    code_block(slide, right, 6.77, 2.10, 5.85, 3.92, size=11.2, title="FRONTEND")
    text_box(slide, "API 계약을 바꾸면 backend models/routes와 frontend types/api를 함께 확인", 0.78, 6.30, 11.72, 0.30, size=11.5, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_data_map(slide, page: int) -> None:
    header(slide, "CODE MAP", "DB 테이블은 기능 묶음으로 이해한다", "전체 정의는 backend/app/db/schema.py · 직접 DELETE 전 외래키를 확인한다.", page)
    groups = [
        ("연보 본문", "annual_reports\nstat_tables · cells", BLUE, SKY),
        ("검수", "profiles · runs · checks\nvalidation_jobs", TEAL, TEAL_BG),
        ("업로드 작업", "report_processing_jobs\nstages", AMBER, AMBER_BG),
        ("원천 원장", "statistical_datasets\nitems · values · revisions", GREEN, GREEN_BG),
        ("증분 연보", "report_incremental_*\nrecipes · policies", BRIGHT_BLUE, SKY),
        ("웹 수정", "passwords · sessions\nmanual_edits", RED, RED_BG),
        ("비교", "report_update_*\nreport_comparison_*", AMBER, AMBER_BG),
        ("언어·LLM", "glossary · candidates\ncache · usage_events", TEAL, TEAL_BG),
    ]
    for i, (title, body, accent, bg) in enumerate(groups):
        x = 0.70 + (i % 4) * 3.06
        y = 2.08 + (i // 4) * 1.72
        rect(slide, x, y, 2.78, 1.42, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.17, y + 0.18, 2.44, 0.30, size=13.5, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.17, y + 0.68, 2.44, 0.46, size=9.5, fill=TEXT, align=PP_ALIGN.CENTER, font=MONO)
    rect(slide, 0.72, 5.82, 11.86, 0.62, fill=RED_BG, line=None, radius=True)
    text_box(slide, "annual_reports 삭제는 연결된 표·셀·검수 기록을 cascade 삭제할 수 있음", 0.98, 6.01, 11.34, 0.26, size=11.7, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_tool_install(slide, page: int) -> None:
    header(slide, "SETUP", "처음 설치하는 세 프로그램", "검색 광고가 아니라 공식 사이트의 설치 파일만 사용한다.", page)
    tools = [
        ("1 · Git", "git-scm.com/install\nWindows: 64-bit, 기본 PATH 옵션\nmacOS: xcode-select --install", BLUE, SKY),
        ("2 · Python", "python.org · 3.12 계열\nWindows: Add python.exe to PATH\nmacOS: universal2 .pkg", TEAL, TEAL_BG),
        ("3 · Node.js", "nodejs.org/en/download\n24 LTS Installer\nnpm은 함께 설치됨", GREEN, GREEN_BG),
        ("4 · 다시 열기", "터미널/PowerShell을 닫고 새로 열기\ngit · python/py · node · npm\n네 명령 모두 버전 확인", AMBER, AMBER_BG),
    ]
    for i, (title, body, accent, bg) in enumerate(tools):
        x = 0.72 + (i % 2) * 6.08
        y = 2.08 + (i // 2) * 1.80
        rect(slide, x, y, 5.76, 1.50, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.22, y + 0.18, 1.45, 0.32, size=14.5, fill=accent, bold=True)
        text_box(slide, body, x + 1.73, y + 0.16, 3.75, 1.04, size=10.3, fill=TEXT, font=MONO)
    rect(slide, 0.73, 5.91, 11.82, 0.58, fill=RED_BG, line=None, radius=True)
    text_box(slide, "‘찾을 수 없습니다’가 나오면 다음 단계로 넘어가지 말고 설치와 PATH를 먼저 해결", 0.99, 6.09, 11.30, 0.25, size=11.6, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_prerequisites(slide, page: int) -> None:
    header(slide, "SETUP", "고정해야 할 도구 버전", "macOS와 Windows의 차이보다 런타임 버전 차이가 더 큰 문제를 만든다.", page)
    rows = [
        ("도구", "권장", "프로젝트 기준"),
        ("Python", "3.12.x", "Dockerfile이 python:3.12-slim"),
        ("Node.js", "24 LTS", "현재 LTS · Vite 8 범위 충족"),
        ("npm", "Node 포함 버전", "package-lock v3 + npm ci"),
        ("Git", "최신 안정", "Windows longpaths / LF 주의"),
        ("Editor", "VS Code 권장", "Python + TypeScript + PowerShell"),
    ]
    table(slide, rows, [2.0, 2.55, 6.95], 0.92, 2.16, row_h=0.60, font_size=11.3)
    text_box(slide, "공식 설치: git-scm.com · python.org(3.12) · nodejs.org(24 LTS)", 1.15, 5.76, 11.05, 0.23, size=10.5, fill=MUTED, align=PP_ALIGN.CENTER)
    rect(slide, 0.92, 6.06, 11.50, 0.56, fill=AMBER_BG, line=None, radius=True)
    text_box(slide, "다른 OS의 .venv와 node_modules를 복사하지 말고 새 PC에서 다시 만든다.", 1.15, 6.23, 11.05, 0.23, size=11.8, fill=AMBER, bold=True, align=PP_ALIGN.CENTER)


def slide_mac(slide, page: int) -> None:
    header(slide, "SETUP · macOS", "macOS 첫 설치 — 저장소 루트에서 시작", "빈 테스트 DB와 LLM OFF가 안전한 첫 실행 기준이다.", page)
    code_block(slide, """git clone https://github.com/wtdnrr/statyearbook.git
cd statyearbook
cp .env.example .env
python3.12 -m venv .venv
./.venv/bin/python -m pip install --upgrade pip
./.venv/bin/python -m pip install -r backend/requirements.txt""", 0.68, 2.10, 6.05, 3.75, size=10.8, title="1 · CODE + BACKEND")
    code_block(slide, """cd frontend
npm ci
npm run build
cd ..

# Terminal 1
./.venv/bin/python -m uvicorn app.main:app \\
  --app-dir backend --reload

# Terminal 2
cd frontend && npm run dev""", 6.92, 2.10, 5.72, 3.75, size=10.4, title="2 · FRONTEND + RUN")
    text_box(slide, "확인  http://127.0.0.1:5173  ·  /health  ·  /ready  ·  /docs", 0.75, 6.22, 11.78, 0.35, size=12.5, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)


def slide_windows(slide, page: int) -> None:
    header(slide, "SETUP · WINDOWS", "Windows 첫 설치 — 짧은 로컬 경로를 사용", "권장 위치 C:\\work\\statyearbook · OneDrive/네트워크 드라이브는 피한다.", page)
    code_block(slide, """Set-Location C:\\work
git -c core.longpaths=true -c core.autocrlf=false clone `
  https://github.com/wtdnrr/statyearbook.git
Set-Location .\\statyearbook
Copy-Item .env.example .env
py -3.12 -m venv .venv
& .\\.venv\\Scripts\\python.exe -m pip install `
  -r backend\\requirements.txt""", 0.68, 2.10, 6.08, 3.80, size=9.7, title="1 · CODE + BACKEND")
    code_block(slide, """Set-Location frontend
npm ci
npm run build
Set-Location ..

# PowerShell 1
& .\\.venv\\Scripts\\python.exe -m uvicorn `
  app.main:app --app-dir backend --reload

# PowerShell 2
Set-Location frontend
npm run dev""", 6.95, 2.10, 5.68, 3.80, size=9.7, title="2 · FRONTEND + RUN")
    text_box(slide, "Activate.ps1이 막혀도 python.exe 직접 실행 방식은 정상 동작", 0.75, 6.23, 11.78, 0.33, size=12, fill=TEAL, bold=True, align=PP_ALIGN.CENTER)


def slide_windows_rules(slide, page: int) -> None:
    header(slide, "SETUP · WINDOWS", "Windows에서 꼭 지킬 네 가지", "OS 차이로 생기는 가짜 오류를 먼저 제거한다.", page)
    rules = [
        ("경로", "C:\\work\\statyearbook", "짧은 ASCII 경로 · OneDrive 밖", BLUE, SKY),
        ("Python", ".venv\\Scripts\\python.exe", "활성화 없이 직접 실행", TEAL, TEAL_BG),
        ("Git", "longpaths=true", "core.autocrlf=false", AMBER, AMBER_BG),
        ("데이터", "로컬 디스크", "대용량 SQLite를 공유 드라이브에서 열지 않음", RED, RED_BG),
    ]
    for i, (title, key, detail, accent, bg) in enumerate(rules):
        x = 0.72 + (i % 2) * 6.08
        y = 2.18 + (i // 2) * 1.75
        rect(slide, x, y, 5.76, 1.43, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.22, y + 0.18, 1.12, 0.32, size=13, fill=accent, bold=True)
        text_box(slide, key, x + 1.35, y + 0.16, 4.08, 0.35, size=14, fill=NAVY, bold=True, font=MONO)
        text_box(slide, detail, x + 0.22, y + 0.75, 5.30, 0.34, size=10.8, fill=MUTED)
    rect(slide, 0.72, 5.95, 11.84, 0.56, fill=NAVY, line=None, radius=True)
    text_box(slide, "다른 OS에서 만든 .venv · node_modules · dist를 복사하지 않는다", 0.98, 6.12, 11.30, 0.24, size=11.8, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_env_model(slide, page: int) -> None:
    header(slide, "CONFIG", "환경변수는 두 군데, 읽는 시점도 다르다", "값을 바꿨는데 반영되지 않으면 파일 위치와 재시작/재배포 여부를 본다.", page)
    card(slide, "ROOT / .env", "FastAPI가 시작할 때 읽음\nDB · CORS · storage · worker · LLM\n변경 후 백엔드 재시작", 0.75, 2.18, 5.75, 2.28, accent=BLUE, fill=SKY, title_size=18, body_size=13)
    card(slide, "frontend / .env.local", "Vite가 개발/빌드 시 읽음\nVITE_API_BASE_URL 하나\n변경 후 Vite 재시작 또는 재배포", 6.83, 2.18, 5.75, 2.28, accent=TEAL, fill=TEAL_BG, title_size=18, body_size=13)
    text_box(slide, "OS 환경변수", 1.08, 5.20, 2.0, 0.30, size=13.5, fill=RED, bold=True)
    text_box(slide, "이미 설정되어 있으면 .env보다 우선 — 테스트 실패의 숨은 원인이 될 수 있음", 3.18, 5.20, 8.80, 0.35, size=12.2, fill=MUTED)
    text_box(slide, "프론트 API URL 뒤에는 /api를 붙이지 않는다", 0.83, 6.04, 11.65, 0.32, size=14, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)


def slide_env_db(slide, page: int) -> None:
    header(slide, "CONFIG", "처음 실행용 안전 설정", "운영 연결 문자열과 키는 빈 값으로 두고 기능 플래그를 끈다.", page)
    code_block(slide, """DATABASE_BACKEND=sqlite
DATABASE_URL=
ANNUAL_STATISTICS_DB_PATH=backend/app/db/annual_statistics_test.sqlite
UPLOAD_DIR=backend/uploads
MAX_UPLOAD_BYTES=104857600
ALLOWED_ORIGINS=http://localhost:5173,http://127.0.0.1:5173

SOURCE_STORAGE_BACKEND=local
JOB_EXECUTOR_BACKEND=inprocess""", 0.73, 2.12, 6.25, 3.72, size=10.7, title="LOCAL · .env")
    card(slide, "DB 선택 규칙", "auto + DATABASE_URL 있음 → postgres\nauto + URL 없음 → sqlite\n처음은 명시적으로 sqlite", 7.26, 2.12, 5.35, 1.60, accent=BLUE, fill=SKY, title_size=15, body_size=11.2)
    card(slide, "상대 경로 규칙", "항상 저장소 루트에서 실행해야 DB와 uploads가 같은 위치를 가리킨다.", 7.26, 3.92, 5.35, 1.42, accent=AMBER, fill=AMBER_BG, title_size=15, body_size=11.2)
    text_box(slide, "운영 DB 최초 연결 전 백업", 7.55, 5.72, 4.80, 0.38, size=15.5, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_env_jobs(slide, page: int) -> None:
    header(slide, "CONFIG", "작업 실행과 원본 저장은 한 쌍", "API와 worker가 분리되면 로컬 파일시스템도 분리된다고 생각한다.", page)
    rows = [
        ("환경", "JOB_EXECUTOR_BACKEND", "SOURCE_STORAGE_BACKEND", "추가 프로세스"),
        ("로컬 개발", "inprocess", "local", "없음"),
        ("단일 API", "inprocess", "local + 영속 볼륨 또는 s3", "없음"),
        ("API + worker", "worker", "s3 / R2 / 공유 볼륨", "worker 필수"),
    ]
    table(slide, rows, [2.1, 3.0, 3.6, 2.6], 1.00, 2.18, row_h=0.64, font_size=10.5)
    card(slide, "queued에서 멈추면", "worker 실행 · 같은 DATABASE_URL · 같은 storage · 같은 Git commit을 비교", 0.98, 5.14, 5.55, 1.18, accent=RED, fill=RED_BG, title_size=14, body_size=10.5)
    card(slide, "running에서 멈추면", "lease 만료 · 원본 접근 · 외부 LLM timeout · worker 로그를 확인", 6.79, 5.14, 5.55, 1.18, accent=AMBER, fill=AMBER_BG, title_size=14, body_size=10.5)


def slide_env_llm(slide, page: int) -> None:
    header(slide, "CONFIG", "LLM은 선택 기능 — 기본값은 OFF", "키 존재, 기능 플래그, 화면 옵션의 세 조건을 모두 확인한다.", page)
    code_block(slide, """LLM_PROVIDER=auto
LLM_REVIEW_ENABLED=0
PROFILE_LLM_ENABLED=0
AUTO_IMPORT_INCLUDE_LLM=0

BIZROUTER_API_KEY=
OPENAI_API_KEY=
LLM_REVIEW_BATCH_SIZE=10
LLM_REVIEW_CONCURRENCY=3
LLM_REVIEW_LIMIT=
LLM_REVIEW_TIMEOUT=90""", 0.72, 2.10, 5.70, 3.92, size=10.6, title="SAFE DEFAULT")
    number_step(slide, 1, "정책 확인", "외부 전송 가능 데이터인지", 6.72, 2.12, 5.85, accent=RED)
    number_step(slide, 2, "비용 승인", "모델·후보 수·월 한도", 6.72, 2.96, 5.85, accent=AMBER)
    number_step(slide, 3, "작게 시험", "LLM_REVIEW_LIMIT 설정", 6.72, 3.80, 5.85, accent=BLUE)
    number_step(slide, 4, "지표 확인", "token과 예상 비용", 6.72, 4.64, 5.85, accent=TEAL)
    text_box(slide, "키 값을 로그·이슈·PPT·화면 캡처에 넣지 않는다", 6.90, 5.80, 5.50, 0.35, size=12.5, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_run(slide, page: int) -> None:
    header(slide, "RUN", "실행 순서와 정상 판정", "프론트가 열리는 것과 시스템이 준비된 것은 다르다.", page)
    steps = [
        (1, "백엔드 시작", "Uvicorn :8000", BLUE),
        (2, "/health", "process = ok", TEAL),
        (3, "/ready", "DB = sqlite/postgres", GREEN),
        (4, "프론트 시작", "Vite :5173", BRIGHT_BLUE),
        (5, "브라우저", "Network/CORS 확인", AMBER),
    ]
    for i, (num, title, detail, accent) in enumerate(steps):
        number_step(slide, num, title, detail, 0.75, 2.10 + i * 0.82, 5.63, accent=accent)
    code_block(slide, """# macOS
curl http://127.0.0.1:8000/health
curl http://127.0.0.1:8000/ready

# Windows PowerShell
Invoke-RestMethod http://127.0.0.1:8000/health
Invoke-RestMethod http://127.0.0.1:8000/ready""", 6.72, 2.12, 5.80, 2.55, size=10.5, title="HEALTH CHECK")
    card(slide, "빈 화면", "새 SQLite라면 연보가 없는 것이 정상\n업로드 또는 인계 DB 연결", 6.72, 4.92, 2.75, 1.22, accent=BLUE, fill=SKY, title_size=13, body_size=9.8)
    card(slide, "ready 실패", "DB 경로·URL·권한·migration·디스크 확인", 9.72, 4.92, 2.80, 1.22, accent=RED, fill=RED_BG, title_size=13, body_size=9.8)


def slide_upload(slide, page: int) -> None:
    header(slide, "USE", "완성 연보 업로드 — 첫 성공 시나리오", "작은 테스트 파일로 전체 경로를 먼저 확인한다.", page)
    for i, (title, detail) in enumerate([
        ("파일 선택", "화면: .hwpx · .md / 100MB 이하"),
        ("정보 확인", "연도 · 제목 · 새 연보 모드"),
        ("LLM 끄기", "첫 실행은 외부 호출 없이"),
        ("상태 확인", "queued → running → completed"),
        ("결과 확인", "생성 연보 자동 선택 · 표 상세"),
    ]):
        number_step(slide, i + 1, title, detail, 0.72, 2.06 + i * 0.82, 5.60, accent=BLUE if i < 3 else TEAL)
    code_block(slide, """POST /api/imports
GET  /api/imports/{job_id}
POST /api/imports/{job_id}/retry

작업 기록에는
· source file / SHA-256
· current stage / progress
· error message / retry 상태가 남음""", 6.68, 2.08, 5.86, 2.62, size=10.7, title="API FLOW")
    card(slide, "실패 시 보관", "작업 ID · 파일명/해시 · 오류 메시지 · 백엔드/worker 로그 · 사용 DB", 6.68, 4.96, 5.86, 1.20, accent=RED, fill=RED_BG, title_size=14, body_size=10.5)


def slide_validation(slide, page: int) -> None:
    header(slide, "USE", "검수 유형과 상태 읽는 법", "자동 판정은 담당자의 확인 순서를 정해주는 보조 결과다.", page)
    types = ["합계(증감액)", "비율(증감률)", "오탈자", "번역", "메타정보", "이상치", "파란색 표기"]
    for i, label in enumerate(types):
        x = 0.72 + (i % 4) * 2.40
        y = 2.16 + (i // 4) * 0.60
        pill(slide, label, x, y, 2.10, fill=SKY if i < 2 else TEAL_BG, text_fill=BLUE if i < 2 else TEAL)
    states = [
        ("정상", "규칙상 문제 없음", GREEN, GREEN_BG),
        ("확인 필요", "사람이 원문/원장 확인", AMBER, AMBER_BG),
        ("오류 의심", "근거 확인 후 수정/예외", RED, RED_BG),
        ("대기", "언어 후보 판정 전", BLUE, SKY),
    ]
    for i, (title, body, accent, bg) in enumerate(states):
        x = 0.73 + i * 3.03
        rect(slide, x, 3.74, 2.75, 1.38, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.16, 3.95, 2.43, 0.30, size=14, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.16, 4.49, 2.43, 0.30, size=9.8, fill=TEXT, align=PP_ALIGN.CENTER)
    rect(slide, 0.73, 5.48, 11.85, 0.86, fill=NAVY, line=None, radius=True)
    text_box(slide, "진행률·단계 표시  |  중단하고 저장  |  계산 profile LLM과 언어 LLM을 따로 선택", 1.02, 5.64, 11.28, 0.24, size=11.3, fill="9ED7FF", bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "수정한 표는 기존 검수 결과가 오래된 상태가 될 수 있으므로 반드시 다시 검수", 1.02, 5.96, 11.28, 0.24, size=11.3, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_ledger(slide, page: int) -> None:
    header(slide, "USE", "원장 3종으로 다음 연도 연보 만들기", "기준 연보의 모양과 새 원장의 값을 연결한다.", page)
    inputs = [("표정보", "통계 메타"), ("표항목", "행·열 구조"), ("표데이터", "기준시점 값")]
    for i, (title, body) in enumerate(inputs):
        x = 0.73 + i * 2.15
        rect(slide, x, 2.31, 1.82, 1.06, fill=SKY, line=None, radius=True)
        text_box(slide, title, x + 0.12, 2.50, 1.58, 0.28, size=14, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.12, 2.93, 1.58, 0.20, size=9.5, fill=MUTED, align=PP_ALIGN.CENTER)
    text_box(slide, "+", 7.18, 2.58, 0.42, 0.35, size=20, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 7.72, 2.31, 2.02, 1.06, fill=TEAL_BG, line=None, radius=True)
    text_box(slide, "기준 연보", 7.90, 2.54, 1.66, 0.28, size=14, fill=TEAL, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "표 형식", 7.90, 2.95, 1.66, 0.20, size=9.5, fill=MUTED, align=PP_ALIGN.CENTER)
    text_box(slide, "→", 9.94, 2.61, 0.48, 0.32, size=20, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)
    rect(slide, 10.48, 2.18, 2.12, 1.34, fill=NAVY, line=None, radius=True)
    text_box(slide, "새 연도 초안", 10.66, 2.49, 1.76, 0.33, size=14.5, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "연결 검토 필요", 10.66, 2.99, 1.76, 0.22, size=9.5, fill="9EC6E5", align=PP_ALIGN.CENTER)
    items = [
        "같은 배치에서 추출한 Excel 3종인지 확인",
        "기준 연보와 대상 연도/제목을 ID까지 확인",
        "연결·신규·무시·기존 유지 대상을 사람이 검토",
        "확정 후 검수 실행 — 자동 연결을 곧바로 정답으로 보지 않음",
    ]
    bullets(slide, items, 0.84, 4.16, 11.70, 1.70, size=12.2, gap=7)


def slide_change_map(slide, page: int) -> None:
    header(slide, "CHANGE", "무엇을 바꾸려는지에 따라 시작 파일이 다르다", "API 계약 변경은 프론트와 백엔드를 함께 바꾸는 작업이다.", page)
    rows = [
        ("변경 목적", "첫 파일", "함께 확인"),
        ("화면/메뉴", "App.tsx · components", "global.css · test.tsx"),
        ("API 호출", "frontend/api/reportApi.ts", "types.ts · backend models"),
        ("API 경로", "backend/api/routes.py", "application · services"),
        ("계산 검수", "backend/validation", "validation_execution_service"),
        ("언어/LLM", "validation/language_pipeline", "core/llm_client.py"),
        ("DB 구조", "db/schema.py", "db/migrations.py · tests"),
        ("원장 생성", "yearbook_data", "ledger_yearbook_service"),
    ]
    table(slide, rows, [2.6, 4.05, 4.70], 0.98, 2.10, row_h=0.54, font_size=10.2)
    text_box(slide, "원칙: routes는 HTTP 경계, 실제 업무 규칙은 application/services", 1.02, 6.43, 11.27, 0.30, size=12, fill=BLUE, bold=True, align=PP_ALIGN.CENTER)


def slide_change_cycle(slide, page: int) -> None:
    header(slide, "CHANGE", "안전한 수정 사이클", "작은 변경과 검증 가능한 기록이 롤백 비용을 줄인다.", page)
    steps = [
        ("1", "요구사항", "기대 결과·재현", BLUE),
        ("2", "브랜치", "작업 단위 분리", TEAL),
        ("3", "기준 테스트", "수정 전 결과", GREEN),
        ("4", "작게 수정", "관련 계층만", AMBER),
        ("5", "검증", "관련 → 전체 → build", BRIGHT_BLUE),
        ("6", "인계 기록", "DB·env·rollback", RED),
    ]
    for i, (num, title, body, accent) in enumerate(steps):
        x = 0.66 + (i % 3) * 4.18
        y = 2.14 + (i // 3) * 1.76
        rect(slide, x, y, 3.87, 1.43, fill=WHITE, line=BORDER, radius=True)
        rect(slide, x + 0.16, y + 0.17, 0.50, 0.50, fill=accent, line=None, radius=True)
        text_box(slide, num, x + 0.16, y + 0.22, 0.50, 0.26, size=12, fill=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0)
        text_box(slide, title, x + 0.82, y + 0.18, 2.74, 0.30, size=14, fill=NAVY, bold=True)
        text_box(slide, body, x + 0.82, y + 0.75, 2.74, 0.27, size=10.5, fill=MUTED)
    rect(slide, 0.68, 5.91, 11.94, 0.58, fill=LIGHT, line=BORDER, radius=True)
    text_box(slide, "커밋 전  git status --short  ·  git diff --check  ·  테스트 결과 기록", 0.98, 6.08, 11.33, 0.25, size=11.5, fill=NAVY, bold=True, align=PP_ALIGN.CENTER, font=MONO)


def slide_database(slide, page: int) -> None:
    header(slide, "DATA", "DB 변경은 코드 변경보다 한 단계 더 검증", "앱 시작 시 migration이 실행될 수 있으므로 실제 DB를 백업하지 않고 시험하지 않는다.", page)
    for i, (title, detail, accent) in enumerate([
        ("1 · schema", "새 설치 구조를 schema.py에 반영", BLUE),
        ("2 · migration", "기존 DB 이동 경로와 재실행 안전성", TEAL),
        ("3 · fresh DB", "새 SQLite에서 전체 생성", GREEN),
        ("4 · DB copy", "기존 DB 복사본에서 migration", AMBER),
        ("5 · PostgreSQL", "전용 테스트 DB에서 호환성", BRIGHT_BLUE),
        ("6 · backup", "운영 적용·rollback 담당 확정", RED),
    ]):
        x = 0.72 + (i % 2) * 6.05
        y = 2.08 + (i // 2) * 1.16
        card(slide, title, detail, x, y, 5.74, 0.91, accent=accent, fill=WHITE, title_size=12.5, body_size=9.5)
    rect(slide, 0.73, 5.82, 11.82, 0.68, fill=RED_BG, line=None, radius=True)
    text_box(slide, "이전 코드를 재배포해도 이미 적용된 DB 변경은 자동으로 취소되지 않는다", 1.00, 6.03, 11.28, 0.26, size=12.2, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_backup(slide, page: int) -> None:
    header(slide, "DATA", "SQLite와 uploads는 같은 시점의 세트", "DB만 있으면 과거 원본 재처리가, uploads만 있으면 작업 연결이 어려울 수 있다.", page)
    card(slide, "1 · STOP", "백엔드·worker를 Ctrl+C로 종료\n처리 중 작업이 없는지 확인", 0.72, 2.12, 3.72, 1.48, accent=RED, fill=RED_BG, title_size=15, body_size=11)
    card(slide, "2 · COPY", "SQLite + backend/uploads\n같은 시각 기준으로 복사", 4.80, 2.12, 3.72, 1.48, accent=BLUE, fill=SKY, title_size=15, body_size=11)
    card(slide, "3 · HASH", "macOS shasum -a 256\nWindows Get-FileHash", 8.88, 2.12, 3.72, 1.48, accent=TEAL, fill=TEAL_BG, title_size=15, body_size=11)
    code_block(slide, """# macOS
cp -p annual_statistics.sqlite annual_statistics_20260828.sqlite
shasum -a 256 annual_statistics_20260828.sqlite

# Windows PowerShell
Copy-Item annual_statistics.sqlite annual_statistics_20260828.sqlite
Get-FileHash annual_statistics_20260828.sqlite -Algorithm SHA256""", 0.76, 4.08, 11.82, 1.78, size=10.5, title="BACKUP EXAMPLE")
    text_box(slide, "대용량 DB를 Git 또는 일반 메신저로 전달하지 않는다", 0.88, 6.24, 11.58, 0.28, size=12.2, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_tests(slide, page: int) -> None:
    header(slide, "VERIFY", "현재 테스트 기준선 — 사실 그대로 인계", "2026-08-28 현재 작업 트리에서 직접 실행한 결과다.", page)
    stats = [
        ("49 / 49", "Frontend tests", GREEN, GREEN_BG),
        ("SUCCESS", "TypeScript + Vite build", BLUE, SKY),
        ("397", "Backend tests run", NAVY, PALE_BLUE),
        ("6 fail · 2 skip", "Backend current baseline", RED, RED_BG),
    ]
    for i, (value, label, accent, bg) in enumerate(stats):
        x = 0.72 + i * 3.03
        rect(slide, x, 2.12, 2.76, 1.35, fill=bg, line=None, radius=True)
        text_box(slide, value, x + 0.16, 2.37, 2.44, 0.36, size=20, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, label, x + 0.16, 2.90, 2.44, 0.26, size=9.8, fill=TEXT, align=PP_ALIGN.CENTER)
    card(slide, "실패 1건", "로컬 .env의 LLM 가능 상태가 ‘LLM 없음’ 테스트에 영향 — 환경 격리 필요", 0.73, 4.00, 5.78, 1.35, accent=AMBER, fill=AMBER_BG, title_size=14, body_size=10.3)
    card(slide, "실패 5건", "로컬 DB의 2025 검수 결과/2026 제목과 승인 snapshot 불일치 — 기준선 재합의 필요", 6.78, 4.00, 5.78, 1.35, accent=RED, fill=RED_BG, title_size=14, body_size=10.3)
    text_box(slide, "금지: 통과시키기 위해 검수 baseline을 검토 없이 갱신", 0.83, 5.95, 11.70, 0.33, size=13, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_test_commands(slide, page: int) -> None:
    header(slide, "VERIFY", "수정 후 실행할 명령", "관련 테스트 → 전체 테스트 → production build 순서가 기본이다.", page)
    code_block(slide, """# macOS · repository root
PYTHONPATH=backend ./.venv/bin/python \\
  -m unittest discover -s backend/tests -q

# Frontend
cd frontend
npm run test
npm run build""", 0.72, 2.10, 5.78, 3.68, size=10.8, title="macOS")
    code_block(slide, """# Windows PowerShell · repository root
$env:PYTHONPATH = "backend"
& .\\.venv\\Scripts\\python.exe -m unittest `
  discover -s backend\\tests -q
Remove-Item Env:PYTHONPATH

Set-Location frontend
npm run test
npm run build""", 6.78, 2.10, 5.78, 3.68, size=10.4, title="WINDOWS")
    text_box(slide, "PostgreSQL 통합 테스트는 운영 DB가 아닌 전용 TEST_DATABASE_URL에서만", 0.82, 6.15, 11.68, 0.33, size=12.2, fill=AMBER, bold=True, align=PP_ALIGN.CENTER)


def slide_deploy(slide, page: int) -> None:
    header(slide, "DEPLOY", "운영 구성: Vercel + Railway + PostgreSQL + R2", "API와 worker는 같은 commit, DB, object storage를 사용한다.", page)
    components = [
        ("Vercel", "React dist\nVITE_API_BASE_URL", SKY, BLUE, 0.72, 2.66),
        ("Railway API", "railway.json\n/ready healthcheck", TEAL_BG, TEAL, 3.34, 2.66),
        ("PostgreSQL", "같은 DATABASE_URL\nbackup / restore", AMBER_BG, AMBER, 6.08, 2.02),
        ("S3 / R2", "같은 bucket\n원본 공유", GREEN_BG, GREEN, 6.08, 3.52),
        ("Railway worker", "railway.worker.json\npublic network 불필요", RED_BG, RED, 9.08, 2.66),
    ]
    for title, body, bg, accent, x, y in components:
        rect(slide, x, y, 2.36, 1.22, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.12, y + 0.18, 2.12, 0.30, size=14, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.12, y + 0.62, 2.12, 0.40, size=9.3, fill=TEXT, align=PP_ALIGN.CENTER)
    for coords in [(3.08, 3.27, 3.34, 3.27), (5.70, 2.98, 6.08, 2.61), (5.70, 3.56, 6.08, 4.12), (8.44, 2.61, 9.08, 3.16), (8.44, 4.12, 9.08, 3.60)]:
        line(slide, *coords, fill=BLUE, width=1.8)
    rect(slide, 0.74, 5.42, 11.83, 0.77, fill=NAVY, line=None, radius=True)
    text_box(slide, "Vercel URL ↔ ALLOWED_ORIGINS  |  API URL ↔ VITE_API_BASE_URL  |  뒤에 /api 금지", 1.02, 5.67, 11.28, 0.27, size=11.8, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_deploy_check(slide, page: int) -> None:
    header(slide, "DEPLOY", "배포 전·후·롤백 체크", "DB migration이 포함되면 코드 롤백과 데이터 롤백을 분리해서 판단한다.", page)
    columns = [
        ("BEFORE", ["승인 commit", "test/build 결과", "DB backup", "env/CORS", "rollback owner"], BLUE, SKY),
        ("AFTER", ["migration log", "/health + /ready", "worker poll", "작은 upload", "validation + export"], TEAL, TEAL_BG),
        ("ROLLBACK", ["추가 작업 중단", "배포·시각 기록", "이전 deploy", "DB 별도 판단", "대표 연보 재확인"], RED, RED_BG),
    ]
    for i, (title, items, accent, bg) in enumerate(columns):
        x = 0.72 + i * 4.08
        rect(slide, x, 2.13, 3.78, 3.78, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.20, 2.36, 3.38, 0.35, size=16, fill=accent, bold=True, align=PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            text_box(slide, f"{j+1:02d}", x + 0.28, 2.99 + j * 0.48, 0.38, 0.23, size=9.5, fill=accent, bold=True, align=PP_ALIGN.CENTER)
            text_box(slide, item, x + 0.78, 2.96 + j * 0.48, 2.68, 0.28, size=11.2, fill=TEXT)
    text_box(slide, "이전 코드는 이미 적용된 migration을 자동 취소하지 않는다", 0.82, 6.27, 11.70, 0.30, size=12.3, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_security(slide, page: int) -> None:
    header(slide, "SECURITY", "운영 공개 전 반드시 보완할 인증 한계", "CORS는 인증이나 방화벽이 아니다. 기관 내부 자료를 그대로 공개하지 않는다.", page)
    card(slide, "서비스 전체 로그인 없음", "읽기·업로드·삭제·검수·다운로드·운영 지표 API에 공통 사용자 인증이 없다.", 0.72, 2.10, 5.74, 1.42, accent=RED, fill=RED_BG, title_size=15, body_size=10.5)
    card(slide, "표 수정만 별도 비밀번호", "초기값은 내선 끝 4자리, 번호가 없으면 0000. 운영 전에 모두 변경한다.", 6.80, 2.10, 5.74, 1.42, accent=AMBER, fill=AMBER_BG, title_size=15, body_size=10.5)
    card(slide, "수정 세션 만료 한계", "비밀번호 변경 시 세션은 폐기되지만 현재 명시적인 세션 유효기간 검사가 없다.", 0.72, 3.84, 5.74, 1.42, accent=AMBER, fill=AMBER_BG, title_size=15, body_size=10.5)
    card(slide, "운영 보호 수단", "VPN·SSO·접근 제어 proxy/API gateway 등 조직 승인 방식으로 Railway API를 보호한다.", 6.80, 3.84, 5.74, 1.42, accent=TEAL, fill=TEAL_BG, title_size=15, body_size=10.5)
    rect(slide, 0.72, 5.82, 11.84, 0.62, fill=NAVY, line=None, radius=True)
    text_box(slide, "API 키 · 초기 비밀번호 · Bearer token · DB URL을 로그·PPT·이슈에 남기지 않는다", 1.00, 6.02, 11.28, 0.25, size=11.7, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_troubleshooting(slide, page: int) -> None:
    header(slide, "TROUBLESHOOT", "증상에서 시작하는 첫 확인", "비밀값을 붙여넣지 말고 상태·경로·작업 ID를 기록한다.", page)
    rows = [
        ("증상", "첫 확인", "대표 원인"),
        ("No module named app", "실행 위치 + --app-dir", "backend 폴더 경계"),
        ("프론트 fetch 실패", "/health → /ready", "API URL · CORS · DB"),
        ("queued 고정", "worker + 같은 DB", "worker 미실행/다른 URL"),
        ("원본 없음", "DB의 object path", "uploads 미인계/storage 불일치"),
        ("database is locked", "동시 프로세스", "SQLite 다중 쓰기/공유드라이브"),
        ("EBADENGINE", "node --version", "Vite 8 범위 밖"),
        ("LLM 안 됨", "flag → key → model", "비활성/키/timeout/rate limit"),
    ]
    table(slide, rows, [3.1, 3.75, 4.60], 0.94, 2.10, row_h=0.54, font_size=10.3)
    rect(slide, 0.95, 6.27, 11.43, 0.46, fill=LIGHT, line=None, radius=True)
    text_box(slide, "재현의 다섯 값: commit · env 종류 · DB · 원본 SHA-256 · job ID", 1.15, 6.37, 11.03, 0.24, size=10.8, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)


def slide_first_day(slide, page: int) -> None:
    header(slide, "ONBOARDING", "첫날 실습 — 운영 데이터 없이 끝까지", "네 번의 작은 성공으로 환경·기능·수정·복구를 모두 확인한다.", page)
    practices = [
        ("A", "빈 DB 실행", "/health · /ready · 빈 화면", BLUE, SKY),
        ("B", "테스트 연보", "upload · 검수 · export", TEAL, TEAL_BG),
        ("C", "작은 코드 수정", "branch · test · build · diff", AMBER, AMBER_BG),
        ("D", "백업/복원", "stop · copy · hash · restore", RED, RED_BG),
    ]
    for i, (letter, title, body, accent, bg) in enumerate(practices):
        x = 0.73 + (i % 2) * 6.05
        y = 2.13 + (i // 2) * 1.72
        rect(slide, x, y, 5.72, 1.41, fill=bg, line=None, radius=True)
        rect(slide, x + 0.20, y + 0.22, 0.62, 0.62, fill=accent, line=None, radius=True)
        text_box(slide, letter, x + 0.20, y + 0.30, 0.62, 0.30, size=14, fill=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0)
        text_box(slide, title, x + 1.08, y + 0.20, 4.30, 0.33, size=16, fill=NAVY, bold=True)
        text_box(slide, body, x + 1.08, y + 0.77, 4.30, 0.30, size=11, fill=MUTED)
    rect(slide, 0.73, 5.90, 11.78, 0.58, fill=NAVY, line=None, radius=True)
    text_box(slide, "완료 조건: 다음 담당자가 화면을 공유하며 네 실습을 직접 재현", 0.98, 6.08, 11.28, 0.25, size=12, fill=WHITE, bold=True, align=PP_ALIGN.CENTER)


def slide_final_check(slide, page: int) -> None:
    header(slide, "HANDOVER", "최종 회의에서 화면으로 확인할 것", "‘전달했다’가 아니라 ‘받은 사람이 재현했다’를 기준으로 종료한다.", page)
    groups = [
        ("CODE", ["최종 branch/commit", "git status 정리", "test/build 결과"], BLUE, SKY),
        ("DATA", ["DB + uploads 세트", "backup time + SHA-256", "검수 baseline"], TEAL, TEAL_BG),
        ("ACCESS", ["GitHub·Vercel·Railway", "Postgres·S3/R2", "LLM 비용/키 소유자"], AMBER, AMBER_BG),
        ("OPERATE", ["upload·validation 시연", "deploy·rollback 시연", "장애 연락망"], RED, RED_BG),
    ]
    for i, (title, items, accent, bg) in enumerate(groups):
        x = 0.73 + (i % 2) * 6.05
        y = 2.11 + (i // 2) * 1.78
        rect(slide, x, y, 5.74, 1.48, fill=bg, line=None, radius=True)
        text_box(slide, title, x + 0.20, y + 0.18, 1.05, 0.30, size=13, fill=accent, bold=True)
        bullets(slide, items, x + 1.34, y + 0.15, 4.14, 1.10, size=10.2, bullet_fill=accent, gap=3)
    text_box(slide, "빈칸: 운영 URL · 계정 소유자 · backup 위치 · rollback 담당자를 문서 23장에 채운다", 0.80, 6.10, 11.68, 0.34, size=11.6, fill=NAVY, bold=True, align=PP_ALIGN.CENTER)


def slide_emergency(slide, page: int) -> None:
    header(slide, "QUICK CARD", "문제가 생기면 이 순서대로", "무작정 재시작하거나 DB를 삭제하기 전에 사실을 고정한다.", page)
    steps = [
        ("1", "중단", "새 업로드·배포 잠시 멈춤", RED),
        ("2", "기록", "시각·commit·job ID·오류", AMBER),
        ("3", "경계", "health → ready → frontend", BLUE),
        ("4", "대상", "env · DB · 원본 SHA-256", TEAL),
        ("5", "복구", "백업/이전 deploy로 검증", GREEN),
    ]
    for i, (num, title, body, accent) in enumerate(steps):
        number_step(slide, int(num), title, body, 0.78, 2.07 + i * 0.83, 6.05, accent=accent)
    code_block(slide, """BACKEND
http://127.0.0.1:8000/health
http://127.0.0.1:8000/ready
http://127.0.0.1:8000/docs

FRONTEND
http://127.0.0.1:5173

DETAIL
docs/PROJECT_HANDOVER_GUIDE.md""", 7.18, 2.10, 5.15, 3.72, size=11.2, title="LOCAL ENDPOINTS")
    text_box(slide, "삭제·reset·DB 덮어쓰기는 대상과 복구 가능성을 확인한 뒤", 7.29, 6.10, 4.95, 0.34, size=11.2, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_contacts(slide, page: int) -> None:
    header(slide, "CONTACTS", "인계 회의에서 반드시 채울 빈칸", "비밀번호·API 키 값은 적지 않고 소유자와 안전한 보관 위치만 적는다.", page)
    rows = [
        ("영역", "담당자/팀", "주소·식별자 또는 보관 위치"),
        ("서비스 책임자", "", ""),
        ("GitHub", "", "repository / team"),
        ("Vercel", "", "project / production URL"),
        ("Railway + PostgreSQL", "", "project / API / worker"),
        ("S3 / R2", "", "bucket / backup"),
        ("LLM 비용·키", "", "secret manager 위치"),
        ("검수 기준 승인", "", "baseline 승인 기록"),
    ]
    table(slide, rows, [2.65, 3.45, 5.40], 0.92, 2.08, row_h=0.55, font_size=10.5)
    text_box(slide, "이 표가 비어 있으면 인수인계는 아직 끝나지 않았다", 0.90, 6.44, 11.55, 0.32, size=13, fill=RED, bold=True, align=PP_ALIGN.CENTER)


def slide_close(slide, page: int) -> None:
    set_background(slide, NAVY)
    pill(slide, "HANDOVER COMPLETE", 0.82, 0.78, 2.30, fill="173C61", text_fill="8ED2FF")
    text_box(slide, "코드가 아니라\n운영 가능한 상태를 인계합니다", 0.82, 1.74, 11.10, 1.42, size=34, fill=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "commit · env · DB · 원본 · job ID", 0.84, 3.70, 10.8, 0.45, size=19, fill="8ED2FF", bold=True)
    text_box(slide, "이 다섯 값을 기록하면 대부분의 문제를 재현할 수 있습니다.", 0.84, 4.40, 10.7, 0.40, size=14, fill="A8C4DC")
    rect(slide, 0.84, 5.38, 11.52, 0.77, fill="122F4D", line="274E71", radius=True)
    text_box(slide, "상세 절차  docs/PROJECT_HANDOVER_GUIDE.md", 1.10, 5.63, 11.00, 0.26, size=13, fill=WHITE, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    text_box(slide, str(page), 12.12, 6.86, 0.55, 0.22, size=9, fill="789AB6", align=PP_ALIGN.RIGHT)


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    builders = [
        build_cover,
        slide_outcome,
        slide_five_rules,
        slide_part_install,
        slide_tool_install,
        slide_prerequisites,
        slide_mac,
        slide_windows,
        slide_windows_rules,
        slide_env_model,
        slide_env_db,
        slide_env_jobs,
        slide_env_llm,
        slide_run,
        slide_part_features,
        slide_product,
        slide_user_features,
        slide_architecture,
        slide_flow,
        slide_upload,
        slide_validation,
        slide_ledger,
        slide_part_maintenance,
        slide_snapshot,
        slide_receive,
        slide_tree,
        slide_data_map,
        slide_change_map,
        slide_change_cycle,
        slide_database,
        slide_backup,
        slide_tests,
        slide_test_commands,
        slide_deploy,
        slide_deploy_check,
        slide_security,
        slide_troubleshooting,
        slide_first_day,
        slide_final_check,
        slide_emergency,
        slide_contacts,
        slide_close,
    ]
    for page, builder in enumerate(builders, start=1):
        slide = prs.slides.add_slide(blank)
        if builder is build_cover:
            builder(slide)
        else:
            builder(slide, page)

    prs.core_properties.title = "통계연보 검수 보조 시스템 프로젝트 인수인계 가이드"
    prs.core_properties.subject = "macOS와 Windows 개발·운영 인수인계"
    prs.core_properties.author = "Project Handover"
    prs.core_properties.keywords = "FastAPI, React, Windows, macOS, handover"
    return prs


def validate(prs: Presentation) -> None:
    errors: list[str] = []
    if len(prs.slides) != 42:
        errors.append(f"expected 42 slides, got {len(prs.slides)}")
    for slide_index, slide in enumerate(prs.slides, start=1):
        if not slide.shapes:
            errors.append(f"slide {slide_index}: no shapes")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                errors.append(f"slide {slide_index}: negative shape position")
            if shape.left + shape.width > prs.slide_width + Inches(0.01):
                errors.append(f"slide {slide_index}: shape exceeds right edge")
            if shape.top + shape.height > prs.slide_height + Inches(0.01):
                errors.append(f"slide {slide_index}: shape exceeds bottom edge")
    if errors:
        raise RuntimeError("presentation validation failed:\n" + "\n".join(errors))


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    validate(prs)
    prs.save(OUTPUT)
    print(f"saved: {OUTPUT}")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
